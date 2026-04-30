from __future__ import annotations

import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

os.environ.setdefault("GOOGLE_API_KEY", "")
os.environ.setdefault("APP_API_KEY", "")


@pytest.fixture
def tmp_data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    monkeypatch.setenv("DATA_DIR", str(data_dir))
    # Override anything a developer's .env would inject so tests stay hermetic.
    for var in (
        "LLM_PROVIDER",
        "GOOGLE_API_KEY",
        "OLLAMA_API_KEY",
        "OLLAMA_BASE_URL",
        "OLLAMA_PLAN_MODEL",
        "OLLAMA_REMEDIATION_MODEL",
        "OLLAMA_REASONING_LEVEL",
        "APP_API_KEY",
        "LIVEKIT_URL",
        "LIVEKIT_API_KEY",
        "LIVEKIT_API_SECRET",
        "LLAMA_CLOUD_API_KEY",
        "LLAMAPARSE_BASE_URL",
    ):
        monkeypatch.setenv(var, "")
    monkeypatch.setenv("LLM_PROVIDER", "gemini")
    monkeypatch.setenv("BASIC_AUTH_ENABLED", "false")
    monkeypatch.setenv("RATE_LIMIT_ENABLED", "false")
    monkeypatch.setenv("TRUST_PROXY_HEADERS", "false")
    monkeypatch.setenv("LLAMAPARSE_TIER", "cost_effective")
    monkeypatch.setenv("LLAMAPARSE_VERSION", "latest")
    monkeypatch.setenv("LLAMAPARSE_MAX_PAGES", "")
    monkeypatch.setenv("LLAMAPARSE_TARGET_PAGES", "")
    monkeypatch.setenv("LLAMAPARSE_COST_OPTIMIZER", "false")
    monkeypatch.setenv("LLAMAPARSE_DISABLE_CACHE", "false")
    monkeypatch.setenv("LLAMAPARSE_DO_NOT_CACHE", "false")
    monkeypatch.setenv("LLAMAPARSE_TIMEOUT_S", "300")
    monkeypatch.setenv("LLAMAPARSE_OCR_LANGUAGES", "en")
    from backend.app.config import get_settings

    get_settings.cache_clear()
    return data_dir


@pytest.fixture
def client(tmp_data_dir: Path):
    from backend.app.main import app

    with TestClient(app) as c:
        yield c


@pytest.fixture
def authed_client(tmp_data_dir: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("APP_API_KEY", "test-key")
    from backend.app.config import get_settings

    get_settings.cache_clear()
    from backend.app.main import app

    with TestClient(app) as c:
        c.headers.update({"X-API-Key": "test-key"})
        yield c


@pytest.fixture
def minimal_pdf_bytes() -> bytes:
    return (
        b"%PDF-1.4\n"
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 144]>>endobj\n"
        b"xref\n0 4\n0000000000 65535 f \n"
        b"0000000010 00000 n \n0000000055 00000 n \n0000000100 00000 n \n"
        b"trailer<</Size 4/Root 1 0 R>>\nstartxref\n160\n%%EOF\n"
    )


@pytest.fixture
def minimal_docx_bytes() -> bytes:
    """A one-paragraph DOCX built in-process. Avoids committing binary fixtures."""
    import io

    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_heading("Sample Memo", level=1)
    doc.add_paragraph("This is a short memo used as a DOCX fixture.")
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def gpa_calculator_xlsx_bytes() -> bytes:
    """A simple GPA-calculator-shaped workbook: four grade inputs + four
    credit inputs + a weighted-average output. Exercises the XLSX
    interactive path."""
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "GPA"
    ws["A1"] = "Course"
    ws["B1"] = "Grade (0-4)"
    ws["C1"] = "Credits"
    for i, (course, grade, credits) in enumerate(
        [
            ("English", 3.7, 3),
            ("Math", 4.0, 4),
            ("History", 3.3, 3),
            ("Science", 3.5, 4),
        ],
        start=2,
    ):
        ws[f"A{i}"] = course
        ws[f"B{i}"] = grade
        ws[f"C{i}"] = credits
    ws["A6"] = "Weighted GPA"
    ws["B6"] = "=SUMPRODUCT(B2:B5,C2:C5)/SUM(C2:C5)"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
