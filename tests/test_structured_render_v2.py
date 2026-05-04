from __future__ import annotations

import io
import json
import re
from pathlib import Path

import openpyxl
import pytest
from bs4 import BeautifulSoup
from fastapi.testclient import TestClient

from backend.app import cache
from backend.app.a11y import check_html
from backend.app.config import get_settings
from backend.app.documents import DocFormat
from backend.app.gemini.remediate import OutlineNode, RemediationResult
from backend.app.images import ExtractedImageAsset
from backend.app.layout_plan_v2 import BlockPlacementPlan, LayoutPlan, SectionLayoutPlan
from backend.app.parser_backends.llamaparse_adapter import LlamaParseDocument, LlamaParsePage
from backend.app.pdf_fetch import FetchedDoc, sha256_of
from backend.app.pipeline_version import STRUCTURED_RENDER_V2_TRANSCRIPT_PIPELINE_VERSION
from backend.app.structured_parser.types import ParsedDocument
from backend.app.structured_render_v2 import (
    LayoutCritiqueAttempt,
    LayoutCritiqueBlockPatch,
    LayoutCritiqueResponse,
    LayoutCritiqueSectionPatch,
    _apply_layout_critique,
    _build_docx_signals,
    _build_pdf_signals,
    _build_xlsx_signals,
    _choose_document_title,
    _clean_metadata_title,
    _maybe_apply_layout_agent_revision,
    _request_layout_critique,
    _signals_from_chunk,
    _split_title_and_remainder,
    _strip_json_fences,
    run_structured_render_v2,
)


def test_split_title_and_remainder_breaks_long_single_line_form_chunk() -> None:
    title, remainder = _split_title_and_remainder(
        "Registration Form   Individual Information  First Name __________"
    )
    assert title == "Registration Form"
    assert remainder == "Individual Information First Name __________"


def test_split_title_and_remainder_leaves_simple_title_alone() -> None:
    title, remainder = _split_title_and_remainder("Short academic paper")
    assert title == "Short academic paper"
    assert remainder is None


def test_static_form_label_rows_become_fields_without_fill_lines() -> None:
    emitted, next_order = _signals_from_chunk(
        ["Last Name                                                           First Name                                   MI                Date of Birth"],
        page_num=1,
        chunk_index=0,
        order=0,
        source="extract",
    )

    assert next_order == 4
    assert [signal.kind for signal in emitted] == ["form_field"] * 4
    assert [signal.field.label for signal in emitted if signal.field] == [
        "Last Name",
        "First Name",
        "MI",
        "Date of Birth",
    ]


def test_clean_metadata_title_strips_office_prefix_and_extension() -> None:
    assert _clean_metadata_title("Microsoft Word - Basic PDF Form Sample.docx") == "Basic PDF Form Sample"


def test_choose_document_title_skips_revision_banner() -> None:
    title = _choose_document_title(
        [(1, "1 Revised 02/09/2021\nEMPLOYEE REMOTE WORK ARRANGEMENT\nThis optional form ...")],
        source_hint="Advanced_PDF_Form_Sample.pdf",
        metadata_title=None,
    )
    assert title == "EMPLOYEE REMOTE WORK ARRANGEMENT"


def test_choose_document_title_skips_disclaimer_and_uses_real_heading() -> None:
    title = _choose_document_title(
        [
            (
                1,
                "Provided proper attribution is provided, Google hereby grants permission to\n"
                "reproduce the tables and figures in this paper solely for use in journalistic or\n"
                "scholarly works.\n"
                "Attention Is All You Need\n"
                "Ashish Vaswani",
            )
        ],
        source_hint="Academic_Paper_Sample.pdf",
        metadata_title=None,
    )
    assert title == "Attention Is All You Need"


def test_choose_document_title_prefers_combined_short_lines_over_single_word_metadata() -> None:
    title = _choose_document_title(
        [
            (
                1,
                "Blue-sky printing\n"
                "There are many reasons for using Prince.\n",
            )
        ],
        source_hint="brochure.pdf",
        metadata_title="Prince",
    )
    assert title == "Blue-sky printing"


def test_choose_document_title_combines_dictionary_heading_lines() -> None:
    title = _choose_document_title(
        [(1, "A Concise\nDictionaryof\nOld Icelandic\nFonts by Monokrom")],
        source_hint="dictionary.pdf",
        metadata_title=None,
    )
    assert title == "A Concise Dictionary of Old Icelandic"


def test_choose_document_title_normalizes_corrupted_invoice_heading() -> None:
    title = _choose_document_title(
        [(1, "YesLogic Pty. Ltd.\n7 / 39 Bouverie St\nInInvvoicoicee\nCustomer Name")],
        source_hint="index.pdf",
        metadata_title=None,
    )
    assert title == "Invoice"


def test_choose_document_title_skips_country_line_before_invoice_heading() -> None:
    title = _choose_document_title(
        [(1, "YesLogic Pty. Ltd.\n7 / 39 Bouverie St\nAustralia\nInInvvoicoicee\nCustomer Name")],
        source_hint="index.pdf",
        metadata_title=None,
    )
    assert title == "Invoice"


def test_choose_document_title_prefers_visible_all_caps_title_over_metadata() -> None:
    title = _choose_document_title(
        [(1, "1 Revised 02/09/2021\nEMPLOYEE REMOTE WORK ARRANGEMENT\nThis optional form ...")],
        source_hint="Advanced_PDF_Form_Sample.pdf",
        metadata_title="Remote Work Acknowledgement",
    )
    assert title == "EMPLOYEE REMOTE WORK ARRANGEMENT"


def test_signals_from_chunk_extracts_form_fields_instead_of_blob_paragraph() -> None:
    emitted, next_order = _signals_from_chunk(
        [
            "Revised 02/09/2021",
            "Remote Work Employee Information:",
            "1. Employee Name and WSUID:",
            "2. Job Title:",
            "3. Workload (FTE): ____ up to .25 (1-10 hrs) ___ up to .50 (1-20 hrs)",
            "____up to .75 (1-32 hrs) ____1.0 (40+ hrs)",
            "Remote Work Review:",
            "Remote work arrangements are to be reviewed prior to the anticipated end of remote work.",
        ],
        page_num=1,
        chunk_index=0,
        order=0,
        source="extract",
    )

    assert next_order == len(emitted)
    assert [signal.kind for signal in emitted[:4]] == ["heading", "form_field", "form_field", "form_field"]
    assert emitted[0].text == "Remote Work Employee Information"
    assert emitted[1].field is not None and emitted[1].field.label == "Employee Name and WSUID"
    assert emitted[2].field is not None and emitted[2].field.label == "Job Title"
    assert emitted[3].field is not None and emitted[3].field.field_type == "checkbox"
    assert emitted[3].field.options
    assert emitted[4].kind == "heading"
    assert emitted[5].kind == "paragraph"


def test_signals_from_chunk_extracts_figure_caption_and_skips_label_cloud() -> None:
    emitted, _ = _signals_from_chunk(
        [
            "Figure 1: Receptors in the hu-",
            "man skin: Mechanoreceptors can",
            "be free receptors or encapsulated.",
            "Hairy skinGlabrous skin",
            "Epidermis",
            "Dermis",
            "Pacinian corpuscle",
            "Cutaneous receptors",
            "Sensory information from Meissner corpuscles and rapidly adapting afferents leads to adjustment of grip force.",
        ],
        page_num=1,
        chunk_index=0,
        order=0,
        source="extract",
    )

    assert emitted[0].kind == "figure"
    assert emitted[0].caption is not None
    assert "Mechanoreceptors" in emitted[0].caption
    assert all("Hairy skinGlabrous skin" not in (signal.text or "") for signal in emitted[1:] if hasattr(signal, "text"))
    assert emitted[1].kind == "heading"
    assert emitted[1].text == "Cutaneous receptors"


def test_build_docx_signals_uses_heading_as_title(minimal_docx_bytes: bytes) -> None:
    signals = _build_docx_signals(
        minimal_docx_bytes,
        source_hint="memo.docx",
        image_filenames=[],
    )
    assert signals.format == "docx"
    assert signals.title == "Sample Memo"
    assert signals.blocks
    assert signals.blocks[0].kind == "paragraph"


def test_build_xlsx_signals_builds_sheet_heading_and_table() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "Year"
    ws["B1"] = "Rate"
    ws["A2"] = "2024"
    ws["B2"] = "7.1"
    buf = io.BytesIO()
    wb.save(buf)

    signals, has_formulas = _build_xlsx_signals(
        buf.getvalue(),
        source_hint="rates.xlsx",
        image_filenames=[],
    )
    assert has_formulas is False
    assert signals.format == "xlsx"
    assert signals.title == "Rates"
    assert len(signals.blocks) == 2
    assert signals.blocks[0].kind == "heading"
    assert signals.blocks[1].kind == "table"


def test_build_xlsx_signals_marks_formula_workbooks() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GPA"
    ws["A1"] = "Grade"
    ws["B1"] = 3.7
    ws["A2"] = "Weighted GPA"
    ws["B2"] = "=B1"
    buf = io.BytesIO()
    wb.save(buf)

    signals, has_formulas = _build_xlsx_signals(
        buf.getvalue(),
        source_hint="gpa.xlsx",
        image_filenames=[],
    )
    assert signals.format == "xlsx"
    assert has_formulas is True


def test_build_pdf_signals_uses_llamaparse_markdown_and_saves_artifact(
    monkeypatch,
    tmp_data_dir: Path,
) -> None:
    monkeypatch.setenv("LLAMAPARSE_TIER", "agentic")
    monkeypatch.setenv("LLAMAPARSE_VERSION", "2026-04-22")
    get_settings.cache_clear()

    async def fake_extract(doc_bytes, *, source_hint):
        assert doc_bytes == b"%PDF"
        assert source_hint == "annual-report.pdf"
        return LlamaParseDocument(
            pages=(
                LlamaParsePage(
                    page_number=1,
                    markdown=(
                        "# Annual Report\n\n"
                        "Opening prose before the first section.\n\n"
                        "| Year | Rate |\n"
                        "| --- | --- |\n"
                        "| 2024 | 7.1 |\n"
                    ),
                    text=None,
                    items=(),
                    metadata={"confidence": 0.99},
                ),
                LlamaParsePage(
                    page_number=2,
                    markdown="## Details\n\n- First\n- Second\n",
                    text=None,
                    items=(),
                    metadata={},
                ),
            ),
            markdown_full=None,
            text_full=None,
            job_metadata={"credits": 20},
            raw={},
        )

    monkeypatch.setattr(
        "backend.app.structured_render_v2._extract_llamaparse_document",
        fake_extract,
    )

    import asyncio

    try:
        signals = asyncio.run(
            _build_pdf_signals(
                b"%PDF",
                source_hint="annual-report.pdf",
                image_filenames=[],
                pdf_backend="llamaparse",
                sha256="abc123",
            )
        )
        assert signals.title == "Annual Report"
        assert signals.page_count == 2
        assert signals.metadata == {
            "pdf_backend": "llamaparse",
            "llamaparse_tier": "agentic",
            "llamaparse_version": "2026-04-22",
        }
        assert [block.source for block in signals.blocks] == ["llamaparse"] * len(signals.blocks)
        assert signals.blocks[0].kind == "paragraph"
        assert signals.blocks[0].text == "Opening prose before the first section."
        table = next(block for block in signals.blocks if block.kind == "table")
        assert table.rows == [["Year", "Rate"], ["2024", "7.1"]]
        assert table.header_rows == [0]

        artifact = tmp_data_dir / "artifacts" / "abc123" / "llamaparse" / "document.json"
        assert json.loads(artifact.read_text())["job_metadata"] == {"credits": 20}
    finally:
        get_settings.cache_clear()


def test_run_structured_render_v2_maps_llamaparse_items_to_accessible_table(
    minimal_pdf_bytes: bytes,
    monkeypatch,
) -> None:
    async def fake_extract(doc_bytes, *, source_hint):
        assert doc_bytes == minimal_pdf_bytes
        assert source_hint == "metrics.pdf"
        return LlamaParseDocument(
            pages=(
                LlamaParsePage(
                    page_number=1,
                    markdown="# Metrics Report\n\nFallback markdown should not duplicate items.",
                    text=None,
                    items=(
                        {"type": "heading", "level": 1, "value": "Metrics Report", "md": "# Metrics Report"},
                        {"type": "text", "value": "Opening prose before the first section.", "md": "Opening prose before the first section."},
                        {"type": "heading", "level": 3, "value": "Rates", "md": "### Rates"},
                        {
                            "type": "table",
                            "rows": [["Year", "Rate"], ["2024", "7.1"], ["2025", "7.4"]],
                            "md": "| Year | Rate |\n| --- | --- |\n| 2024 | 7.1 |\n| 2025 | 7.4 |",
                        },
                    ),
                    metadata={"confidence": 0.99},
                ),
            ),
            markdown_full=None,
            text_full=None,
            job_metadata={"credits": 3},
            raw={},
        )

    monkeypatch.setattr(
        "backend.app.structured_render_v2._extract_llamaparse_document",
        fake_extract,
    )

    import asyncio

    result = asyncio.run(
        run_structured_render_v2(
            minimal_pdf_bytes,
            source_hint="metrics.pdf",
            fmt=DocFormat.PDF,
            image_filenames=[],
            pdf_backend="llamaparse",
        )
    )

    rendered = result["result"]
    soup = BeautifulSoup(rendered.html, "html.parser")

    assert rendered.title == "Metrics Report"
    assert check_html(rendered.html) == []
    assert "Opening prose before the first section." in soup.get_text(" ", strip=True)
    assert "Fallback markdown should not duplicate items." not in soup.get_text(" ", strip=True)
    assert [th.get_text(" ", strip=True) for th in soup.select("thead th")] == ["Year", "Rate"]
    assert [th.get("scope") for th in soup.select("thead th")] == ["col", "col"]
    body_header = soup.select_one("tbody th")
    assert body_header is not None
    assert body_header.get_text(" ", strip=True) == "2024"
    assert body_header.get("scope") == "row"


def test_run_structured_render_v2_infers_llamaparse_static_form_fields_and_suppresses_images(
    minimal_pdf_bytes: bytes,
    monkeypatch,
) -> None:
    async def fake_extract(doc_bytes, *, source_hint):
        assert doc_bytes == minimal_pdf_bytes
        assert source_hint == "k12-parent-consent.pdf"
        return LlamaParseDocument(
            pages=(
                LlamaParsePage(
                    page_number=1,
                    markdown="Fallback markdown should not duplicate items.",
                    text=None,
                    items=(
                        {
                            "type": "image",
                            "caption": "Los Angeles Community College District logo",
                            "md": "![Los Angeles Community College District logo](page_1_image_1.jpg)",
                            "url": "page_1_image_1.jpg",
                        },
                        {
                            "type": "heading",
                            "level": 1,
                            "value": "Concurrent/College Career Access Pathways Enrollment Parent Consent",
                            "md": "# Concurrent/College Career Access Pathways Enrollment Parent Consent",
                        },
                        {"type": "heading", "level": 2, "value": "Student Information:", "md": "## Student Information:"},
                        {
                            "type": "text",
                            "value": "LACCD Student ID Number\n(if available)",
                            "md": "LACCD Student ID Number\n(if available)",
                        },
                        {
                            "type": "text",
                            "value": "Last Name [] First Name [] MI [] Date of Birth []",
                            "md": "Last Name \\[***] First Name \\[***] MI \\[***] Date of Birth \\[***]",
                        },
                        {"type": "heading", "level": 2, "value": "Parent Information:", "md": "## Parent Information:"},
                        {
                            "type": "text",
                            "value": "Parent's Last Name [] Parent's First Name []",
                            "md": "Parent's Last Name \\[***] Parent's First Name \\[***]",
                        },
                    ),
                    metadata={},
                ),
            ),
            markdown_full=None,
            text_full=None,
            job_metadata={},
            raw={},
        )

    monkeypatch.setattr(
        "backend.app.structured_render_v2._extract_llamaparse_document",
        fake_extract,
    )

    import asyncio

    result = asyncio.run(
        run_structured_render_v2(
            minimal_pdf_bytes,
            source_hint="k12-parent-consent.pdf",
            fmt=DocFormat.PDF,
            image_filenames=[ExtractedImageAsset(filename="img-1.png", page_number=1)],
            pdf_backend="llamaparse",
        )
    )

    rendered = result["result"]
    parsed = result["parsed_document"]
    soup = BeautifulSoup(rendered.html, "html.parser")
    labels = [label.get_text(" ", strip=True) for label in soup.select("label")]

    assert parsed.doc_kind == "form_doc"
    assert check_html(rendered.html) == []
    assert len(soup.select("form")) == 1
    assert not soup.select("img")
    assert "Figure for" not in soup.get_text(" ", strip=True)
    assert "LACCD Student ID Number" in labels
    assert "Last Name" in labels
    assert "First Name" in labels
    assert "Date of Birth" in labels
    assert "Parent's Last Name" in labels


def test_run_structured_render_v2_uses_hybrid_path_for_formula_xlsx(monkeypatch) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GPA"
    ws["A1"] = "Grade"
    ws["B1"] = 3.7
    ws["A2"] = "Weighted GPA"
    ws["B2"] = "=B1"
    buf = io.BytesIO()
    wb.save(buf)

    async def fake_plan_layout_v2(parsed_document, *, source_hint, fmt):
        assert fmt is DocFormat.XLSX
        return {"version": "v2", "hero_treatment": "compact", "density": "balanced", "toc_policy": "hidden", "sections": []}

    async def fake_remediate_document(doc_bytes, *, source_hint, fmt, image_filenames, plan_json):
        assert fmt is DocFormat.XLSX
        return RemediationResult(
            html='<main id="content" lang="en"><h1 id="document-title">GPA Calculator</h1><form></form></main>',
            outline=[OutlineNode(id="document-title", level=1, text="GPA Calculator").model_dump()],
            title="GPA Calculator",
            description="",
            language="en",
            page_count=1,
            render_mode="interactive",
        )

    monkeypatch.setattr("backend.app.structured_render_v2.plan_layout_v2", fake_plan_layout_v2)
    monkeypatch.setattr("backend.app.structured_render_v2.remediate_mod.remediate_document", fake_remediate_document)

    import asyncio

    result = asyncio.run(
        run_structured_render_v2(
            buf.getvalue(),
            source_hint="gpa.xlsx",
            fmt=DocFormat.XLSX,
            image_filenames=[],
        )
    )
    assert result["result"].render_mode == "interactive"
    assert result["parsed_document"].source_format == "xlsx"


def test_run_structured_render_v2_renders_gpa_xlsx_calculator_without_llm(monkeypatch) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GPA"
    ws["A1"] = "Course"
    ws["B1"] = "Grade (0.0-4.0)"
    ws["C1"] = "Credits"
    for i, (course, grade, credits) in enumerate(
        [
            ("English Composition", 3.7, 3),
            ("Calculus I", 4.0, 4),
            ("World History", 3.3, 3),
            ("Intro to Biology", 3.5, 4),
        ],
        start=2,
    ):
        ws[f"A{i}"] = course
        ws[f"B{i}"] = grade
        ws[f"C{i}"] = credits
    ws["A7"] = "Total credits"
    ws["B7"] = "=SUM(C2:C5)"
    ws["A8"] = "Weighted GPA"
    ws["B8"] = "=SUMPRODUCT(B2:B5,C2:C5)/SUM(C2:C5)"
    buf = io.BytesIO()
    wb.save(buf)

    async def fail_remediate_document(*args, **kwargs):
        raise AssertionError("formula xlsx calculator should not call the LLM renderer")

    monkeypatch.setattr(
        "backend.app.structured_render_v2.remediate_mod.remediate_document",
        fail_remediate_document,
    )

    import asyncio

    result = asyncio.run(
        run_structured_render_v2(
            buf.getvalue(),
            source_hint="GPA_Calculator.xlsx",
            fmt=DocFormat.XLSX,
            image_filenames=[],
        )
    )

    rendered = result["result"]
    soup = BeautifulSoup(rendered.html, "html.parser")
    script = soup.find("script")

    assert rendered.render_mode == "interactive"
    assert rendered.title == "GPA Calculator"
    assert soup.find("output", id="xlsx-cell-b7") is not None
    assert soup.find("output", id="xlsx-cell-b8") is not None
    assert script is not None
    script_text = script.get_text()
    assert 'setOutput("xlsx-cell-b7", sum([number("xlsx-cell-c2")' in script_text
    assert 'setOutput("xlsx-cell-b8", sumProduct([number("xlsx-cell-b2")' in script_text
    assert ")/sum([number(\"xlsx-cell-c2\")" in script_text
    assert check_html(rendered.html, render_mode="interactive") == []


def _nonresident_waiver_pdf_bytes() -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    form = c.acroForm
    c.setFont("Helvetica", 10)
    lines = [
        "Los Angeles Community College District",
        "Nonresident Tuition Fee Waiver Application",
        "Note: a separate form must be submitted for each semester",
        "Term Requested: ______________________ Year: ______________________",
        "Last Name: ________________________ First Name: ________________________ Middle Initial: ________",
        "Student ID #: ______________________ Date of Birth: ________________",
        "Home Address: ____________________________________ City: ________________________",
        "ZIP Code: __________ Email Address: ____________________________________",
        "Phone Number: _____________________",
        "Eligibility: Please read carefully and answer the following questions:",
        "1. My immigration status prevents me from establishing residency in the United States: YES NO",
        '2. I am in the United States under a current "F", "J", or "M" Visa: YES NO',
        'STOP now if you answered "Yes" to #2, above',
        "You are not eligible for this waiver except in circumstances of documented severe economic hardship.",
        "Please submit income information to establish economic hardship.",
        "3. My family income* is at or below the income levels in the chart below. YES NO",
        "Family Size 2024 Income",
        "1 $21,870",
        "2 $29,580",
        "3 $37,290",
        "4 $45,000",
        "5 $52,710",
        "6 $60,420",
        "7 $68,130",
        "8 $75,840",
        "Each Additional Family Member $7,710",
        "*These standards are based upon the federal poverty guidelines, as published each year by the US Department of Health and Human Services.",
        "APPLICANT CERTIFICATION: PLEASE READ AND SIGN BELOW",
        "I hereby swear or affirm, under penalty of perjury, that all the information on this form is true and complete to the best of my knowledge.",
        "Applicant's Signature: _________________________ Date: ________________",
        "Parent Signature (required for dependent students under age 19): _________________________ Date: ______________",
        "FOR OFFICE USE ONLY",
        "Action: Approved Noted Emailed Denied",
        "By: ____________________________________ Date: _________________",
    ]
    y = 760
    for line in lines:
        c.drawString(50, y, line)
        y -= 18

    for name, x, y, width in [
        ("Term Requested", 160, 705, 100),
        ("Year", 320, 705, 80),
        ("Last Name", 120, 670, 100),
        ("First Name", 300, 670, 100),
        ("Middle Initial", 500, 670, 50),
        ("Student ID", 130, 650, 100),
        ("Date3_es_:signer:date", 300, 650, 90),
        ("Home Address", 130, 610, 200),
        ("City", 360, 610, 100),
        ("ZIP Code", 120, 590, 60),
        ("Email Address_es_:email", 220, 590, 180),
        ("Phone Number", 140, 570, 100),
        ("Signature1_es_:signer:signature", 180, 180, 140),
        ("Date_es_:date", 360, 180, 90),
        ("Signature2_es_:signer:signature", 260, 160, 140),
        ("Date_2_es_:date", 440, 160, 90),
        ("Noted_", 230, 100, 100),
        ("By", 100, 80, 200),
        ("Date_3_es_:date", 360, 80, 90),
    ]:
        form.textfield(name=name, x=x, y=y, width=width, height=14, borderWidth=1)

    for name, x, y in [
        ("My immigration status prevents me from establishing residency in the United States", 450, 520),
        ("I am in the United States under a current F J or M Visa", 450, 500),
        ("My family income is at or below the income levels in the chart below", 450, 430),
        ("Approved", 120, 100),
        ("Noted", 170, 100),
        ("Emailed", 330, 100),
        ("Denied", 400, 100),
    ]:
        form.checkbox(name=name, x=x, y=y, buttonStyle="check", borderWidth=1)

    c.showPage()
    c.drawString(50, 760, "Instructions")
    c.drawString(
        50,
        730,
        "LACCD Board Policy 5020 states that students who are citizens and residents of a foreign country may be eligible.",
    )
    c.save()
    return buf.getvalue()


def test_laccd_acroform_uses_facsimile_renderer_without_invented_images_or_tables() -> None:
    import asyncio

    pdf_bytes = _nonresident_waiver_pdf_bytes()
    result = asyncio.run(
        run_structured_render_v2(
            pdf_bytes,
            source_hint="Nonresident Tuition Fee Waiver Application.pdf",
            fmt=DocFormat.PDF,
            image_filenames=[],
            pdf_backend="native",
        )
    )

    rendered = result["result"]
    soup = BeautifulSoup(rendered.html, "html.parser")

    assert rendered.title == "Nonresident Tuition Fee Waiver Application"
    assert result["layout_plan"]["renderer"] == "acroform-facsimile"
    assert check_html(rendered.html) == []
    assert [h.get_text(" ", strip=True) for h in soup.find_all("h2")] == [
        "Term information",
        "Student information",
        "Contact information",
        "Eligibility",
        "Applicant certification",
        "For office use only",
        "Instructions",
    ]
    assert len(soup.find_all("img")) == 0
    assert len(soup.find_all("figure")) == 0
    assert len(soup.find_all("form")) == 1
    assert len(soup.find_all("table")) == 1
    assert "2024 income levels by family size" in soup.get_text(" ", strip=True)
    assert "PDF Form Fields" not in soup.get_text(" ", strip=True)
    assert "Instructions table" not in soup.get_text(" ", strip=True)


def test_laccd_docx_form_uses_pdf_style_facsimile_renderer() -> None:
    import asyncio

    docx_path = Path(__file__).resolve().parents[1] / "demo" / "pdfs" / "Nonresident_Tuition_Fee_Waiver_Application.docx"
    if not docx_path.exists():
        pytest.skip("demo DOCX fixture not present; run scripts/download_samples.sh")

    result = asyncio.run(
        run_structured_render_v2(
            docx_path.read_bytes(),
            source_hint=docx_path.name,
            fmt=DocFormat.DOCX,
            image_filenames=[],
            pdf_backend="native",
        )
    )

    rendered = result["result"]
    soup = BeautifulSoup(rendered.html, "html.parser")

    assert rendered.title == "Nonresident Tuition Fee Waiver Application"
    assert result["layout_plan"]["renderer"] == "docx-form-facsimile"
    assert "doc-page--facsimile" in soup.main.get("class", [])
    assert [h.get_text(" ", strip=True) for h in soup.find_all("h2")] == [
        "Student information",
        "Contact information",
        "Eligibility",
        "Applicant certification",
        "Instructions",
    ]
    assert len(soup.find_all("form")) == 1
    assert len(soup.find_all("table")) == 1
    assert "2024 income levels by family size" in soup.get_text(" ", strip=True)
    assert check_html(rendered.html) == []


def test_generic_acroform_uses_single_facsimile_form_without_filename_bridge() -> None:
    import asyncio

    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    form = c.acroForm
    c.setFont("Helvetica", 10)
    c.drawString(50, 760, "Scholarship Appeal Form")
    c.drawString(50, 730, "Term: ____________________ Year: ____________________")
    c.drawString(50, 700, "Student Name: ____________________ Student ID: ____________________")
    c.drawString(50, 670, "Email: ____________________ Phone: ____________________")
    c.drawString(50, 640, "Course Name: ____________________ Class Number: ____________________")
    c.drawString(50, 610, "I certify this information is accurate.")
    c.drawString(50, 580, "Student Signature: ____________________ Date: ____________________")
    for name, x, y in [
        ("Term", 100, 728),
        ("Year", 280, 728),
        ("Student Name", 150, 698),
        ("Student ID", 390, 698),
        ("Email", 100, 668),
        ("Phone", 290, 668),
        ("Course Name", 150, 638),
        ("Class Number", 390, 638),
        ("Student Signature", 170, 578),
        ("Date", 390, 578),
    ]:
        form.textfield(name=name, x=x, y=y, width=120, height=14, borderWidth=1)
    c.save()

    result = asyncio.run(
        run_structured_render_v2(
            buf.getvalue(),
            source_hint="scholarship-appeal.pdf",
            fmt=DocFormat.PDF,
            image_filenames=[],
            pdf_backend="native",
        )
    )
    rendered = result["result"]
    soup = BeautifulSoup(rendered.html, "html.parser")

    assert rendered.title == "Scholarship Appeal Form"
    assert result["layout_plan"]["renderer"] == "acroform-facsimile"
    assert check_html(rendered.html) == []
    assert len(soup.find_all("form")) == 1
    assert [h.get_text(" ", strip=True) for h in soup.find_all("h2")] == [
        "Term information",
        "Student information",
        "Contact information",
        "Course information",
        "Applicant certification",
    ]
    assert "PDF Form Fields" not in soup.get_text(" ", strip=True)


def test_laccd_acroform_ingest_e2e_caches_facsimile_transcript(
    client: TestClient,
    tmp_data_dir: Path,
    monkeypatch,
) -> None:
    import asyncio

    pdf_bytes = _nonresident_waiver_pdf_bytes()
    src = "https://example.test/nonresident-waiver.pdf"
    fetched = FetchedDoc(
        bytes_=pdf_bytes,
        sha256=sha256_of(pdf_bytes),
        source_url=src,
        content_length=len(pdf_bytes),
        fmt=DocFormat.PDF,
    )

    async def fake_fetch_document(url, *, max_bytes, allowed_hosts, headers):
        assert url == src
        return fetched

    async def fake_ensure_store(sha256, bytes_, *, fmt):
        return "stores/nonresident"

    monkeypatch.setenv("GOOGLE_API_KEY", "test-google-key")
    monkeypatch.setenv("STRUCTURED_RENDER_V2_ENABLED", "true")
    monkeypatch.setenv("LAYOUT_AGENT_ENABLED", "false")
    get_settings.cache_clear()
    monkeypatch.setattr("backend.app.routes.transcript.fetch_document", fake_fetch_document)
    monkeypatch.setattr("backend.app.routes.transcript.file_search.ensure_store", fake_ensure_store)

    response = client.get("/api/transcript/ingest", params={"src": src})
    assert response.status_code == 200
    assert 'event: done' in response.text

    cached = asyncio.run(cache.get_by_hash(tmp_data_dir / "cache.db", fetched.sha256))
    assert cached is not None
    assert cached.pipeline_version == STRUCTURED_RENDER_V2_TRANSCRIPT_PIPELINE_VERSION
    assert cached.title == "Nonresident Tuition Fee Waiver Application"

    soup = BeautifulSoup(cached.html, "html.parser")
    assert len(soup.find_all("img")) == 0
    assert len(soup.find_all("figure")) == 0
    assert len(soup.find_all("table")) == 1
    assert "PDF Form Fields" not in soup.get_text(" ", strip=True)
    assert check_html(cached.html) == []


def _render_demo_laccd_pdf(filename: str):
    import asyncio

    pdf_path = Path(__file__).resolve().parents[1] / "demo" / "pdfs" / filename
    if not pdf_path.exists():
        pytest.skip(f"demo PDF fixture not present: {filename}; run scripts/download_samples.sh")

    return asyncio.run(
        run_structured_render_v2(
            pdf_path.read_bytes(),
            source_hint=filename,
            fmt=DocFormat.PDF,
            image_filenames=[],
            pdf_backend="native",
        )
    )


def test_laccd_ab540_uses_tooltips_and_school_attendance_table() -> None:
    result = _render_demo_laccd_pdf("Nonresident_Tuition_Exemption_Request.pdf")
    rendered = result["result"]
    soup = BeautifulSoup(rendered.html, "html.parser")
    text = soup.get_text(" ", strip=True)

    assert result["layout_plan"]["renderer"] == "acroform-facsimile"
    assert len(soup.find_all("form")) == 1
    assert len(soup.find_all(["input", "textarea", "select"])) == 47
    assert soup.find("table") is not None
    assert "Name of CA School" in text
    assert "School attendance and credit/hour information" in text
    assert "pg1-1" not in text
    assert "Signature9" not in text
    assert check_html(rendered.html) == []


def test_laccd_fee_waiver_instruction_text_is_clean() -> None:
    result = _render_demo_laccd_pdf("Nonresident_Tuition_Fee_Waiver_Application.pdf")
    rendered = result["result"]
    soup = BeautifulSoup(rendered.html, "html.parser")
    text = soup.get_text(" ", strip=True)

    assert result["layout_plan"]["renderer"] == "acroform-facsimile"
    assert "g uidelines" not in text
    assert "f ill" not in text
    assert "Recor ds" not in text
    assert "U .S." not in text
    assert "guidelines apply to you, please fill out this form" in text
    assert check_html(rendered.html) == []


def test_laccd_repeated_course_fields_become_tables() -> None:
    pass_no_pass = _render_demo_laccd_pdf("pass_no_pass_petition.pdf")["result"].html
    ew_petition = _render_demo_laccd_pdf("laccd_ew_petition_240209_0.pdf")["result"].html

    pass_soup = BeautifulSoup(pass_no_pass, "html.parser")
    ew_soup = BeautifulSoup(ew_petition, "html.parser")

    assert "Course Name&Number" in pass_soup.get_text(" ", strip=True)
    assert "Class Number" in pass_soup.get_text(" ", strip=True)
    assert len(pass_soup.find_all("table")) == 1
    assert "Subject and Number" in ew_soup.get_text(" ", strip=True)
    assert "Term and Year" in ew_soup.get_text(" ", strip=True)
    assert len(ew_soup.find_all("table")) == 1
    assert check_html(pass_no_pass) == []
    assert check_html(ew_petition) == []


def test_laccd_checkbox_forms_do_not_expose_raw_check_box_labels() -> None:
    for filename in (
        "Supplemental_Residency_Questionnaire.pdf",
        "Certification_of_Homeless_Status_REV_012018.pdf",
    ):
        rendered = _render_demo_laccd_pdf(filename)["result"]
        soup = BeautifulSoup(rendered.html, "html.parser")
        labels = [label.get_text(" ", strip=True) for label in soup.find_all("label")]

        assert not [
            label
            for label in labels
            if re.match(r"^Check Box", label, flags=re.IGNORECASE)
        ]
        assert check_html(rendered.html) == []


def _phase2_document() -> ParsedDocument:
    return ParsedDocument.model_validate(
        {
            "schema_version": "structured-render-v2",
            "source_format": "pdf",
            "title": "Doc",
            "language": "en",
            "page_count": 1,
            "doc_kind": "report",
            "body": [
                {
                    "id": "intro",
                    "type": "section",
                    "heading": "Intro",
                    "level": 2,
                    "page_start": 1,
                    "page_end": 1,
                    "placement": "inline",
                    "blocks": [
                        {
                            "id": "p1",
                            "type": "paragraph",
                            "text": "Lead paragraph",
                            "page_start": 1,
                            "page_end": 1,
                            "placement": "lead",
                        },
                        {
                            "id": "fig1",
                            "type": "figure",
                            "asset_id": "img-1",
                            "caption": "Figure 1",
                            "alt_text_hint": "Figure 1 alt",
                            "page_start": 1,
                            "page_end": 1,
                            "placement": "rail",
                        },
                    ],
                }
            ],
        }
    )


def _phase2_layout() -> LayoutPlan:
    return LayoutPlan(
        hero_treatment="compact",
        density="balanced",
        toc_policy="hidden",
        sections=[
            SectionLayoutPlan(
                section_id="intro",
                container="flow-with-rail",
                rail_side="right",
                blocks=[
                    BlockPlacementPlan(
                        block_index=0,
                        block_type="paragraph",
                        placement="main",
                        container="lead-prose",
                    ),
                    BlockPlacementPlan(
                        block_index=1,
                        block_type="figure",
                        placement="rail",
                        container="figure",
                    ),
                ],
            )
        ],
    )


def test_apply_layout_critique_updates_allowed_fields() -> None:
    revised = _apply_layout_critique(
        _phase2_layout(),
        LayoutCritiqueResponse(
            hero_treatment="summary",
            density="dense",
            sections=[
                LayoutCritiqueSectionPatch(
                    section_id="intro",
                    container="flow",
                    rail_side="none",
                    blocks=[
                        LayoutCritiqueBlockPatch(
                            block_index=1,
                            placement="main",
                            container="figure",
                        )
                    ],
                )
            ],
        ),
        _phase2_document(),
    )
    assert revised.hero_treatment == "summary"
    assert revised.density == "dense"
    assert revised.sections[0].container == "flow"
    assert revised.sections[0].blocks[1].placement == "main"


def test_apply_layout_critique_rejects_unknown_block_index() -> None:
    try:
        _apply_layout_critique(
            _phase2_layout(),
            LayoutCritiqueResponse(
                sections=[
                    LayoutCritiqueSectionPatch(
                        section_id="intro",
                        blocks=[LayoutCritiqueBlockPatch(block_index=99, placement="main")],
                    )
                ]
            ),
            _phase2_document(),
        )
    except ValueError as e:
        assert "unknown block_index" in str(e)
    else:
        raise AssertionError("expected critique to reject an unknown block index")


def test_maybe_apply_layout_agent_revision_persists_artifacts(monkeypatch, tmp_path: Path) -> None:
    document = _phase2_document()
    layout = _phase2_layout()
    render_input = {
        "title": "Doc",
        "language": "en",
        "page_count": 1,
        "sections": [{"id": "intro", "heading": "Intro", "blocks": [{"kind": "paragraph", "text": "Lead paragraph"}]}],
    }
    initial_render = RemediationResult(
        html='<main id="content" lang="en"><h1 id="document-title">Doc</h1></main>',
        outline=[OutlineNode(id="document-title", level=1, text="Doc").model_dump()],
        title="Doc",
        description="",
        language="en",
        page_count=1,
    )

    class FakeRendered:
        html = initial_render.html
        outline = initial_render.outline
        title = "Doc"
        description = ""
        language = "en"
        page_count = 1
        render_mode = "static"

    monkeypatch.setattr(
        "backend.app.structured_render_v2.get_settings",
        lambda: type(
            "S",
            (),
            {
                "layout_agent_enabled": True,
                "layout_agent_max_revisions": 1,
                "public_origin": "http://127.0.0.1:8000",
                "artifacts_dir_for": lambda self, sha: tmp_path / sha,
            },
        )(),
    )
    monkeypatch.setattr(
        "backend.app.structured_render_v2._capture_layout_screenshot",
        lambda html: __import__("asyncio").sleep(0, result=b"png"),
    )
    monkeypatch.setattr(
        "backend.app.structured_render_v2._request_layout_critique",
        lambda **kwargs: __import__("asyncio").sleep(
            0,
            result=LayoutCritiqueAttempt(
                critique=LayoutCritiqueResponse(hero_treatment="summary"),
                attempts=[{"raw_content": "{}", "raw_response": {}, "error": None}],
            ),
        ),
    )
    monkeypatch.setattr(
        "backend.app.structured_render_v2.render_document",
        lambda parsed, layout: FakeRendered(),
    )

    import asyncio

    rendered, revised_layout = asyncio.run(
        _maybe_apply_layout_agent_revision(
            parsed_document=document,
            render_input=render_input,
            layout_plan=layout,
            rendered=FakeRendered(),
            source_hint="doc.pdf",
            fmt=DocFormat.PDF,
            sha256="abc123",
        )
    )
    assert revised_layout.hero_treatment == "summary"
    artifact_dir = tmp_path / "abc123" / "layout-agent"
    assert json.loads((artifact_dir / "initial_layout_plan.json").read_text())["hero_treatment"] == "compact"
    assert json.loads((artifact_dir / "revised_layout_plan.json").read_text())["hero_treatment"] == "summary"
    assert (artifact_dir / "before.png").read_bytes() == b"png"
    assert (artifact_dir / "after.png").read_bytes() == b"png"
    attempts = json.loads((artifact_dir / "critique_attempts.json").read_text())
    assert len(attempts) == 1
    assert attempts[0]["error"] is None
    assert json.loads((artifact_dir / "skip_reason.json").read_text())["skip_reason"] is None


def test_maybe_apply_layout_agent_revision_skips_when_format_not_enabled(monkeypatch, tmp_path: Path) -> None:
    document = _phase2_document()
    layout = _phase2_layout()

    class FakeRendered:
        html = '<main id="content" lang="en"><h1 id="document-title">Doc</h1></main>'
        outline = [OutlineNode(id="document-title", level=1, text="Doc").model_dump()]
        title = "Doc"
        description = ""
        language = "en"
        page_count = 1
        render_mode = "static"

    monkeypatch.setattr(
        "backend.app.structured_render_v2.get_settings",
        lambda: type(
            "S",
            (),
            {
                "layout_agent_enabled": True,
                "layout_agent_apply_to_formats": frozenset({"pdf"}),
                "layout_agent_max_revisions": 1,
                "public_origin": "http://127.0.0.1:8000",
                "artifacts_dir_for": lambda self, sha: tmp_path / sha,
            },
        )(),
    )

    import asyncio

    rendered, revised_layout = asyncio.run(
        _maybe_apply_layout_agent_revision(
            parsed_document=document,
            render_input={"title": "Doc", "language": "en", "page_count": 1, "sections": []},
            layout_plan=layout,
            rendered=FakeRendered(),
            source_hint="memo.docx",
            fmt=DocFormat.DOCX,
            sha256="docx123",
        )
    )

    assert rendered.title == "Doc"
    assert revised_layout.model_dump(mode="json") == layout.model_dump(mode="json")
    artifact_dir = tmp_path / "docx123" / "layout-agent"
    assert json.loads((artifact_dir / "skip_reason.json").read_text())["skip_reason"] == "layout agent disabled for format docx"


def test_layout_critique_response_ignores_extra_keys() -> None:
    payload = LayoutCritiqueResponse.model_validate_json(
        _strip_json_fences(
            """
            {
              "summary": "Tighten spacing",
              "layout_plan": {"ignored": true},
              "sections": [
                {
                  "section_id": "intro",
                  "blocks": [
                    {
                      "block_index": 0,
                      "placement": "main",
                      "container": "prose",
                      "ignored": "x"
                    }
                  ]
                }
              ]
            }
            """
        )
    )
    assert payload.summary == "Tighten spacing"
    assert payload.sections[0].blocks[0].container == "prose"


def test_request_layout_critique_captures_invalid_raw_replies(monkeypatch) -> None:
    async def fake_ollama_chat_json(**_kwargs):
        return (
            '{"summary":"bad","sections":[{"section_id":"intro","blocks":[{"block_index":"oops"}]}]}',
            {"message": {"content": "bad"}},
        )

    monkeypatch.setattr("backend.app.structured_render_v2.ollama_chat_json", fake_ollama_chat_json)
    monkeypatch.setattr(
        "backend.app.structured_render_v2.get_settings",
        lambda: type(
            "S",
            (),
            {
                "layout_agent_model": "kimi-k2.6:cloud",
                "ollama_reasoning_level": "low",
                "layout_agent_timeout_s": 5.0,
            },
        )(),
    )

    import asyncio

    attempt = asyncio.run(
        _request_layout_critique(
            parsed_document=_phase2_document(),
            layout_plan=_phase2_layout(),
            screenshot_bytes=b"png",
            source_hint="doc.pdf",
            fmt=DocFormat.PDF,
        )
    )
    assert isinstance(attempt, LayoutCritiqueAttempt)
    assert attempt.critique is None
    assert len(attempt.attempts) == 2
    assert attempt.attempts[0]["raw_content"]
    assert "block_index" in attempt.attempts[0]["raw_content"]
    assert attempt.attempts[0]["error"]
