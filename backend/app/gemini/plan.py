"""Planner pass: ask Gemini to analyse the PDF's structure and produce a
typed `DocumentPlan` that drives the synthesis pass in `remediate.py`.

The planner is cheap and structured — it never emits HTML, just intent. The
synthesiser then renders HTML that conforms to the plan, pulling text
content from the same PDF. Splitting the work this way keeps heading ids
consistent with the outline, gives us a stable place to enforce layout
quality (forms vs prose vs tables), and makes bad output much easier to
diagnose because we can inspect the plan.
"""
from __future__ import annotations

import asyncio
import io
import json
import logging
import re
from typing import Any, Literal

from pydantic import BaseModel, Field, ValidationError

from backend.app.config import get_settings
from backend.app.documents import DocFormat
from backend.app.gemini import prompts
from backend.app.gemini.client import get_genai_client, thinking_config
from backend.app.images import (
    ImageAssetLike,
    describe_image_assets_for_prompt,
    filter_image_assets_for_page_range,
)
from backend.app.ollama_client import chat_json as ollama_chat_json
from backend.app.ollama_client import extract_pdf_text, render_pdf_page_images

log = logging.getLogger(__name__)


def _is_likely_truncation(error: BaseException | None) -> bool:
    """Detect failures that look like Gemini output got cut off, corrupted, or stalled at the
    token-budget edge. Same patterns as `remediate._is_likely_truncation` (kept inline to avoid a
    cross-module dependency)."""
    if error is None:
        return False
    if isinstance(error, TimeoutError):
        return True
    msg = str(error)
    return (
        "EOF while parsing" in msg
        or "control character" in msg
        or "expected `,` or `}`" in msg
        or "expected `,` or `]`" in msg
        or "expected value" in msg
        or "timed out" in msg
    )


class FieldSpec(BaseModel):
    id: str
    label: str
    type: Literal[
        "text",
        "email",
        "tel",
        "url",
        "number",
        "date",
        "textarea",
        "checkbox",
        "radio",
        "select",
        "signature",
    ]
    group_name: str | None = None
    required: bool = False


class BlockPlan(BaseModel):
    type: Literal[
        "paragraph",
        "list",
        "table",
        "figure",
        "callout",
        "quote",
        "reference_list",
        "meta_note",
        "form_group",
    ]
    text: str | None = None
    items: list[str] = Field(default_factory=list)
    figure_id: str | None = None
    note: str | None = None


class SectionPlan(BaseModel):
    id: str
    heading: str
    caption: str | None = None
    layout: Literal["form", "prose", "table", "list", "mixed"]
    fields: list[FieldSpec] = Field(default_factory=list)
    images: list[str] = Field(default_factory=list)
    page_start: int | None = Field(default=None, ge=1)
    page_end: int | None = Field(default=None, ge=1)
    blocks: list[BlockPlan] = Field(default_factory=list)
    notes: str | None = None


class FigurePlan(BaseModel):
    id: str
    section_id: str | None = None
    asset_stems: list[str] = Field(default_factory=list)
    caption: str | None = None
    alt_text_hint: str | None = None
    page_start: int | None = Field(default=None, ge=1)
    page_end: int | None = Field(default=None, ge=1)
    placement_context: str | None = None


class DocumentPlan(BaseModel):
    title: str
    language: str = "en"
    page_count: int = Field(ge=1)
    description: str | None = None
    subtitle: str | None = None
    kicker: str | None = None
    summary: str | None = None
    doc_kind: Literal[
        "scientific_article",
        "memo",
        "brochure",
        "reference_doc",
        "form_doc",
        "report",
    ] = "report"
    sections: list[SectionPlan]
    figures: list[FigurePlan] = Field(default_factory=list)
    callouts: list[str] = Field(default_factory=list)
    footnotes: list[str] = Field(default_factory=list)
    references: list[str] = Field(default_factory=list)


def _strip_json_fences(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _fallback_title(source_hint: str) -> str:
    stem = source_hint.rstrip("/").rsplit("/", 1)[-1]
    stem = stem.rsplit(".", 1)[0]
    return stem.replace("_", " ").replace("-", " ").title() or "Document"


def _count_input_units(doc_bytes: bytes, fmt: DocFormat) -> int:
    try:
        if fmt is DocFormat.PDF:
            import pypdf

            return max(1, len(pypdf.PdfReader(io.BytesIO(doc_bytes)).pages))
        if fmt is DocFormat.DOCX:
            from docx import Document  # type: ignore[import-not-found]

            doc = Document(io.BytesIO(doc_bytes))
            return max(1, len(doc.paragraphs) // 30 or 1)
        if fmt is DocFormat.XLSX:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(doc_bytes), read_only=True, data_only=False)
            try:
                return max(1, len(wb.sheetnames))
            finally:
                wb.close()
    except Exception:  # noqa: BLE001
        return 1
    return 1


def _coerce_string_list(values: Any) -> list[str]:
    if not isinstance(values, list):
        return []
    out: list[str] = []
    for value in values:
        if isinstance(value, str):
            if value.strip():
                out.append(value.strip())
            continue
        if isinstance(value, dict):
            text = value.get("text") or value.get("caption") or value.get("title") or value.get("id")
            if isinstance(text, str) and text.strip():
                out.append(text.strip())
                continue
        text = str(value).strip()
        if text:
            out.append(text)
    return out


async def plan_pdf(
    pdf_bytes: bytes,
    *,
    source_hint: str,
    image_filenames: list[ImageAssetLike] | None = None,
    fmt: DocFormat = DocFormat.PDF,
    _allow_chunk_fallback: bool = True,
) -> DocumentPlan:
    """Run the planner against `pdf_bytes` and return a validated plan.

    Raises RuntimeError on parse failure after one retry. Historically
    PDF-only — the `fmt` parameter now routes DOCX/XLSX through the same
    code path with the appropriate MIME type + planner preamble.
    """
    settings = get_settings()
    use_ollama = settings.llm_provider == "ollama"
    if use_ollama and fmt is DocFormat.PDF and _allow_chunk_fallback:
        try:
            return await plan_pdf_chunked(
                pdf_bytes,
                source_hint=source_hint,
                image_filenames=image_filenames,
                pages_per_chunk=6,
                max_chunks=100,
            )
        except Exception as e:  # noqa: BLE001
            if not settings.google_api_key:
                raise
            log.warning("ollama chunked planner failed for %s; falling back to Gemini: %s", source_hint, e)
            use_ollama = False
    client = get_genai_client() if not use_ollama else None

    # DOCX/XLSX bypass the file upload (Gemini File API rejects OOXML MIME);
    # same extraction approach as remediate_document.
    if fmt is DocFormat.PDF and not use_ollama:
        pdf_io = io.BytesIO(pdf_bytes)
        uploaded = await asyncio.to_thread(
            client.files.upload,  # type: ignore[union-attr]
            file=pdf_io,
            config={"mime_type": fmt.mime_type, "display_name": source_hint[:200]},
        )
        extracted_text: str | None = None
    else:
        uploaded = None
        if fmt is DocFormat.PDF:
            extracted_text = extract_pdf_text(pdf_bytes)
        else:
            from backend.app.extract import extract_docx_as_markdown, extract_xlsx_as_json

            extracted_text = (
                extract_docx_as_markdown(pdf_bytes)
                if fmt is DocFormat.DOCX
                else extract_xlsx_as_json(pdf_bytes)
            )

    from google.genai import types as gtypes

    image_assets_text = describe_image_assets_for_prompt(image_filenames)
    image_note = (
        "Extracted visual assets are available in document order. Reference the relevant "
        "asset stems (for example `img-1`) in each section's `images` list when that "
        "section contains them. Assets marked as full-page rasters were generated to "
        "preserve vector-drawn figures that were not extractable as standalone images.\n"
        f"{image_assets_text}"
        if image_assets_text
        else "No visual assets were extracted from this document."
    )

    fmt_noun = {
        DocFormat.PDF: "PDF",
        DocFormat.DOCX: "Word document",
        DocFormat.XLSX: "Excel workbook",
    }[fmt]
    user_prompt = (
        f"Analyse this {fmt_noun} and produce a DocumentPlan. "
        f"Source hint: {source_hint}. "
        + image_note
        + " Follow the planner system instructions exactly. Return JSON only."
    )
    if extracted_text is not None:
        user_prompt += (
            "\n\nExtracted source content (authoritative; the source file is NOT "
            "attached — this text IS the document):\n"
            + ("```markdown\n" if fmt is DocFormat.DOCX else "```json\n")
            + extracted_text
            + "\n```"
        )

    system_instruction = prompts.planner_system_for(fmt)

    if use_ollama:
        if fmt is DocFormat.PDF:
            extracted_text = extract_pdf_text(pdf_bytes)
            image_payload = render_pdf_page_images(pdf_bytes)
        else:
            image_payload = []
        user_prompt += "\n\nReturn JSON matching the provided schema exactly."
        if extracted_text is not None:
            user_prompt += (
                "\n\nExtracted source content (authoritative text for the source):\n"
                + ("```markdown\n" if fmt is DocFormat.DOCX else "```text\n")
                + extracted_text
                + "\n```"
            )
        last_error: Exception | None = None
        for attempt, temp in enumerate((0.1, 0.0)):
            try:
                content, _raw = await asyncio.wait_for(
                    ollama_chat_json(
                        model=settings.ollama_plan_model,
                        system_instruction=system_instruction,
                        user_prompt=user_prompt,
                        schema=DocumentPlan.model_json_schema(),
                        images=image_payload,
                        temperature=temp,
                        think=settings.ollama_reasoning_level,
                    ),
                    timeout=settings.gemini_call_timeout,
                )
                raw_plan = json.loads(_strip_json_fences(content))
                if isinstance(raw_plan, dict):
                    raw_plan.setdefault("title", _fallback_title(source_hint))
                    raw_plan.setdefault("language", "en")
                    raw_plan.setdefault("page_count", _count_input_units(pdf_bytes, fmt))
                    raw_plan.setdefault("sections", [])
                    raw_plan["callouts"] = _coerce_string_list(raw_plan.get("callouts"))
                    raw_plan["footnotes"] = _coerce_string_list(raw_plan.get("footnotes"))
                    raw_plan["references"] = _coerce_string_list(raw_plan.get("references"))
                return DocumentPlan.model_validate(raw_plan)
            except TimeoutError:
                last_error = RuntimeError(
                    f"ollama planner call timed out after {settings.gemini_call_timeout}s"
                )
                log.warning("ollama planner attempt %d timed out", attempt + 1)
                continue
            except (ValidationError, ValueError) as e:
                last_error = e
                log.warning("ollama planner parse attempt %d failed: %s", attempt + 1, e)
            except Exception as e:  # noqa: BLE001
                last_error = e
                log.warning("ollama planner attempt %d failed: %s", attempt + 1, e)
        if _allow_chunk_fallback and _is_likely_truncation(last_error):
            log.info(
                "ollama planner failed, falling back to page-chunked planning for %s",
                source_hint,
            )
            return await plan_pdf_chunked(
                pdf_bytes, source_hint=source_hint, image_filenames=image_filenames
            )
        raise RuntimeError(f"planner returned invalid JSON: {last_error}") from last_error

    def _call(*, temperature: float) -> Any:
        contents = [uploaded, user_prompt] if uploaded is not None else [user_prompt]
        return client.models.generate_content(  # type: ignore[union-attr]
            model=settings.gemini_ask_model,
            contents=contents,
            config=gtypes.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=DocumentPlan,
                system_instruction=system_instruction,
                temperature=temperature,
                max_output_tokens=65_536,
                thinking_config=thinking_config(settings.gemini_ask_model),
            ),
        )

    last_error: Exception | None = None
    for attempt, temp in enumerate((0.1, 0.0)):
        try:
            response = await asyncio.wait_for(
                asyncio.to_thread(_call, temperature=temp),
                timeout=settings.gemini_call_timeout,
            )
        except TimeoutError:
            last_error = RuntimeError(
                f"planner Gemini call timed out after {settings.gemini_call_timeout}s"
            )
            log.warning("planner attempt %d timed out", attempt + 1)
            continue
        try:
            plan = getattr(response, "parsed", None)
            if plan is None:
                plan = DocumentPlan.model_validate_json(response.text)
            return plan
        except (ValidationError, ValueError) as e:
            last_error = e
            log.warning("planner parse attempt %d failed: %s", attempt + 1, e)

    if _allow_chunk_fallback and _is_likely_truncation(last_error):
        log.info("plan parse failed, falling back to page-chunked planning for %s", source_hint)
        return await plan_pdf_chunked(
            pdf_bytes, source_hint=source_hint, image_filenames=image_filenames
        )
    raise RuntimeError(f"planner returned invalid JSON: {last_error}") from last_error


async def plan_pdf_chunked(
    pdf_bytes: bytes,
    *,
    source_hint: str,
    image_filenames: list[ImageAssetLike] | None = None,
    pages_per_chunk: int = 10,
    max_chunks: int = 100,
) -> DocumentPlan:
    """Plan a large PDF by splitting it into page-range chunks and merging the results.

    Each chunk is planned independently (without the chunked fallback to avoid
    recursion). Sections are concatenated in document order; duplicate IDs get a
    chunk-index prefix.
    """
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    total_pages = len(reader.pages)

    all_sections: list[SectionPlan] = []
    title: str | None = None
    language = "en"
    description: str | None = None
    seen_ids: set[str] = set()

    chunk_starts = list(range(0, max(total_pages, 1), pages_per_chunk))[:max_chunks]
    for chunk_idx, start in enumerate(chunk_starts):
        end = min(start + pages_per_chunk, total_pages)
        writer = pypdf.PdfWriter()
        for page_num in range(start, end):
            writer.add_page(reader.pages[page_num])
        buf = io.BytesIO()
        writer.write(buf)
        chunk_bytes = buf.getvalue()

        try:
            chunk_plan = await plan_pdf(
                chunk_bytes,
                source_hint=f"{source_hint} (pages {start + 1}–{end})",
                image_filenames=filter_image_assets_for_page_range(
                    image_filenames, start + 1, end
                ),
                _allow_chunk_fallback=False,
            )
        except Exception:
            log.warning("chunk planner failed for pages %d-%d of %s", start + 1, end, source_hint)
            continue

        if title is None:
            title = chunk_plan.title
            language = chunk_plan.language
            description = chunk_plan.description

        for section in chunk_plan.sections:
            sid = section.id if section.id not in seen_ids else f"c{chunk_idx}-{section.id}"
            seen_ids.add(sid)
            update = {
                "page_start": start + 1,
                "page_end": end,
            }
            if sid != section.id:
                update["id"] = sid
            all_sections.append(section.model_copy(update=update))

    return DocumentPlan(
        title=title or "Document",
        language=language,
        page_count=total_pages,
        description=description,
        sections=all_sections,
    )
