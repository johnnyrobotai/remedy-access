from __future__ import annotations

import asyncio
import io
import json
import logging
import re
from collections.abc import AsyncIterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field, ValidationError

from backend.app.config import get_settings
from backend.app.documents import DocFormat
from backend.app.gemini import prompts
from backend.app.gemini.client import get_genai_client, thinking_config
from backend.app.images import ImageAssetLike, image_asset_filenames

log = logging.getLogger(__name__)


class OutlineNode(BaseModel):
    id: str
    level: int = Field(ge=1, le=6)
    text: str
    parent: str | None = None


class RemediationOutput(BaseModel):
    # Extra keys are tolerated (Gemini occasionally emits a `render_mode`
    # field on non-XLSX jobs; we just ignore it when the format didn't
    # ask for it).
    model_config = ConfigDict(extra="ignore")

    html: str
    outline: list[OutlineNode]
    title: str
    description: str = ""
    language: str = "en"
    page_count: int = Field(ge=1)
    # Only meaningful for XLSX. "static" (default) = plain accessible table.
    # "interactive" = calculator-shaped workbook rendered as a working HTML
    # form with an inline <script> block that recomputes outputs.
    render_mode: Literal["static", "interactive"] = "static"


class RemediationChunkOutput(BaseModel):
    """One batch of rendered <section> blocks — no <main> wrapper."""

    model_config = ConfigDict(extra="ignore")

    html: str
    outline: list[OutlineNode]


@dataclass
class RemediationResult:
    html: str
    outline: list[dict[str, Any]]
    title: str
    description: str
    language: str
    page_count: int
    render_mode: str = "static"


@dataclass
class IngestEvent:
    step: Literal["fetched", "uploaded", "indexed", "generating", "done", "error"]
    pct: int
    message: str = ""
    payload: dict[str, Any] | None = None


def _count_pages(doc_bytes: bytes, fmt: DocFormat = DocFormat.PDF) -> int:
    """Best-effort page/sheet count for the inbound document.

    We prefer a real count but swallow library errors and return 1 —
    the count is used only to give the model a hint and to populate the
    transcript metadata, never for correctness-critical logic.
    """
    try:
        if fmt is DocFormat.PDF:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(doc_bytes))
            return max(1, len(reader.pages))
        if fmt is DocFormat.DOCX:
            from docx import Document  # type: ignore[import-not-found]

            doc = Document(io.BytesIO(doc_bytes))
            # DOCX has no reliable rendered-page count without a layout
            # engine; use paragraph count / 30 as a coarse proxy.
            paras = sum(1 for _ in doc.paragraphs)
            return max(1, paras // 30 or 1)
        if fmt is DocFormat.XLSX:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(doc_bytes), read_only=True, data_only=False)
            try:
                return max(1, len(wb.sheetnames))
            finally:
                wb.close()
    except Exception as e:  # noqa: BLE001
        log.warning("page count failed for fmt=%s: %s", fmt, e)
    return 1


_IMG_REF_RE = re.compile(r'src\s*=\s*"(img-\d+)(?:\.[A-Za-z0-9]+)?"')


def rewrite_image_srcs(
    html: str,
    sha256: str,
    image_filenames: list[ImageAssetLike] | None,
) -> str:
    """Rewrite `src="img-N"` (with or without extension) to
    `src="/images/{sha256}/img-N.{ext}"` based on the saved filename list.

    Anchored to `img-N` to avoid touching any other src value Gemini might emit.
    """
    filenames = image_asset_filenames(image_filenames)
    if not filenames:
        return html
    by_stem = {Path(f).stem: f for f in filenames}

    def _sub(m: re.Match[str]) -> str:
        stem = m.group(1)
        fname = by_stem.get(stem)
        if not fname:
            return 'src=""'
        return f'src="/images/{sha256}/{fname}"'

    return _IMG_REF_RE.sub(_sub, html)


def _validate_wcag_baseline(html: str) -> list[str]:
    """Static structural WCAG 2.1 AA checks. Runs on every transcript before
    it reaches the cache — axe-core runs the full dynamic audit separately."""
    from backend.app.a11y import check_html

    return check_html(html)


_PLACEHOLDER_PHRASES = (
    "refer to the original",
    "refer to the source",
    "see the original document",
    "see the source document",
    "see the source pdf",
    "for full details",
    "content continues",
    "continues in the full",
    "omitted for brevity",
    "summarized for brevity",
    "summary of ",
    "abridged",
    "for brevity",
    "original document for",
    "too lengthy",
    "too large to reproduce",
    "consult the original",
)

_CORRUPTION_PATTERNS = (
    re.compile(r"\.{40,}"),  # wall of dots
    re.compile(r"`{3,}"),  # stray markdown fences
    re.compile(r"-{40,}"),  # wall of dashes
    re.compile(r"={40,}"),  # wall of equals
)


def _looks_like_placeholder_language(html: str) -> bool:
    """Detect model-summary/placeholder or corruption patterns in a chunk's HTML.

    Gemini sometimes replies with disclaimer prose like "refer to the original document"
    or with runs of literal dots when it runs out of budget mid-section. Neither is a
    valid transcript body; we reject the chunk and retry with a stronger prompt.
    """
    lowered = html.lower()
    for phrase in _PLACEHOLDER_PHRASES:
        if phrase in lowered:
            return True
    for rx in _CORRUPTION_PATTERNS:
        if rx.search(html):
            return True
    return False


def _is_likely_truncation(error: BaseException) -> bool:
    """Detect failures that look like Gemini output got cut off, corrupted, stalled, or
    repeatedly summarised at the token-budget edge. Any of these patterns means the
    chunked/placeholder fallback is worth trying."""
    if isinstance(error, TimeoutError):
        return True
    msg = str(error)
    return (
        "EOF while parsing" in msg
        or "control character" in msg
        or "expected `,` or `}`" in msg
        or "expected `,` or `]`" in msg
        or "expected value" in msg
        or "timed out" in msg
        or "placeholder/summary language" in msg
        or "503 UNAVAILABLE" in msg
        or "RESOURCE_EXHAUSTED" in msg
    )


def _fix_and_check(html: str, render_mode: str = "static") -> tuple[str, list[str]]:
    """Apply automatic structural fixups, then run the WCAG checker.

    Returns the (possibly modified) HTML and the remaining issues. The fixups
    repair the most common Gemini compliance gaps (missing <th scope>, label
    for/input id mismatches) so output that's structurally close gets through.
    PDF-extraction artifacts (hyphen line-breaks, multi-column newlines) are
    scrubbed first so they never reach downstream consumers.
    """
    from backend.app.a11y import check_html, fix_html
    from backend.app.text_cleanup import clean_extraction_artifacts

    cleaned = clean_extraction_artifacts(html)
    fixed = fix_html(cleaned)
    return fixed, check_html(fixed, render_mode=render_mode)


async def remediate_document(
    doc_bytes: bytes,
    source_hint: str,
    *,
    fmt: DocFormat = DocFormat.PDF,
    image_filenames: list[ImageAssetLike] | None = None,
    plan_json: str | None = None,
) -> RemediationResult:
    """Run the synthesis pass and return validated HTML + outline.

    `fmt` picks the system-prompt preamble and the MIME type used when
    uploading the file to Gemini. For XLSX the response may come back
    with `render_mode == "interactive"`, which loosens the a11y gate to
    permit the calculator's inline `<script>` block.

    `image_filenames` is the list of server-extracted images in document
    order (PDF-only today — DOCX/XLSX paths pass an empty list). They're
    mentioned in the prompt so Gemini can reference them as
    `<img src="img-N" alt="...">`; the caller rewrites those relative src
    values to served URLs afterward.

    `plan_json` is the structured DocumentPlan emitted by the planner
    pass. When present, the synthesiser is told to conform to it strictly.
    When absent we fall back to prompting the model to infer structure on
    its own.
    """
    settings = get_settings()
    client = get_genai_client()

    # When a PDF plan with sections is available, skip straight to chunked
    # per-section synthesis. One Gemini call per section prevents silent
    # summarisation — each call has its own token budget and an explicit
    # "reproduce ALL content" prompt. Full-document synthesis is retained
    # below only as a fallback when no plan is available or the format
    # isn't PDF.
    if fmt is DocFormat.PDF and plan_json:
        try:
            plan_preview = json.loads(plan_json)
            if plan_preview.get("sections"):
                log.info(
                    "plan has %d sections; routing %s through chunked synthesis",
                    len(plan_preview["sections"]),
                    source_hint,
                )
                return await remediate_pdf_chunked(
                    doc_bytes,
                    source_hint,
                    plan_json,
                    image_filenames=image_filenames,
                    chunk_size=1,
                )
        except (ValueError, json.JSONDecodeError):
            log.warning("plan_json was unparseable; continuing to full synthesis")

    # Gemini's File API accepts PDF but rejects OOXML MIME types, so for
    # DOCX/XLSX we extract the content server-side (python-docx / openpyxl)
    # and pass it in as a text part. Same prompts, just a different input
    # channel. XLSX preserves formulas verbatim so the calculator path
    # can reason about the formula graph.
    if fmt is DocFormat.PDF:
        doc_io = io.BytesIO(doc_bytes)
        uploaded = await asyncio.to_thread(
            client.files.upload,
            file=doc_io,
            config={"mime_type": fmt.mime_type, "display_name": source_hint[:200]},
        )
        extracted_text: str | None = None
    else:
        uploaded = None
        from backend.app.extract import extract_docx_as_markdown, extract_xlsx_as_json

        if fmt is DocFormat.DOCX:
            extracted_text = extract_docx_as_markdown(doc_bytes)
        else:
            extracted_text = extract_xlsx_as_json(doc_bytes)

    from google.genai import types as gtypes

    image_filenames = image_asset_filenames(image_filenames)
    if image_filenames:
        filenames_for_model = ", ".join(Path(f).stem for f in image_filenames)
        images_clause = (
            f" The document contains {len(image_filenames)} extracted image(s): "
            f"{filenames_for_model}. Reference them as <img src=\"img-N\"> using "
            "the filename without extension."
        )
    else:
        images_clause = " The document contains no extractable images; do not emit any <img> tags."

    plan_clause = (
        "\n\nDocumentPlan (JSON — the source of truth for structure, ids, and "
        "form field shapes):\n" + plan_json
        if plan_json
        else "\n\nNo DocumentPlan was provided; infer structure from the source yourself."
    )

    fmt_noun = {
        DocFormat.PDF: "PDF",
        DocFormat.DOCX: "Word document",
        DocFormat.XLSX: "Excel workbook",
    }[fmt]
    unit_noun = "sheet" if fmt is DocFormat.XLSX else "page"
    prompt = (
        f"Remediate this {fmt_noun} into accessible HTML per your system instructions. "
        f"Source hint (URL or filename): {source_hint}. "
        f"The document has {_count_pages(doc_bytes, fmt)} {unit_noun}(s) total."
        + images_clause
        + plan_clause
    )
    if extracted_text is not None:
        content_block = (
            "\n\nExtracted source content (authoritative; the source file is NOT "
            "attached — this text IS the document):\n"
            + ("```markdown\n" if fmt is DocFormat.DOCX else "```json\n")
            + extracted_text
            + "\n```"
        )
        prompt = prompt + content_block

    system_instruction = prompts.remediation_system_for(fmt)

    def _call(*, temperature: float) -> Any:
        contents = [uploaded, prompt] if uploaded is not None else [prompt]
        return client.models.generate_content(
            model=settings.gemini_remediation_model,
            contents=contents,
            config=gtypes.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=RemediationOutput,
                system_instruction=system_instruction,
                temperature=temperature,
                max_output_tokens=32_768,
                thinking_config=thinking_config(settings.gemini_remediation_model),
            ),
        )

    parsed: RemediationOutput | None = None
    last_error: Exception | None = None
    for attempt, temp in enumerate((0.2, 0.0)):
        try:
            response = await asyncio.wait_for(
                asyncio.to_thread(_call, temperature=temp),
                timeout=settings.gemini_call_timeout,
            )
        except TimeoutError:
            last_error = RuntimeError(
                f"remediation Gemini call timed out after {settings.gemini_call_timeout}s"
            )
            log.warning("remediation attempt %d timed out", attempt + 1)
            continue
        try:
            p: RemediationOutput | None = response.parsed  # type: ignore[assignment]
            if p is None:
                p = RemediationOutput.model_validate_json(response.text)
        except (ValidationError, ValueError) as e:
            last_error = RuntimeError(f"remediation model returned invalid JSON: {e}")
            log.warning("remediation parse attempt %d failed: %s", attempt + 1, e)
            continue
        # Only XLSX may opt into interactive; for other formats ignore any
        # render_mode the model drifted into emitting.
        render_mode = p.render_mode if fmt is DocFormat.XLSX else "static"
        p.html, issues = _fix_and_check(p.html, render_mode=render_mode)
        if issues:
            last_error = RuntimeError(
                f"remediation output failed structural checks: {', '.join(issues)}"
            )
            log.warning("remediation WCAG check attempt %d failed: %s", attempt + 1, issues)
            continue
        # XLSX interactive only: the a11y gate passed, but the inline
        # <script> calculator body may still have literal JS parse errors
        # (Gemini has been observed emitting `""—"` and similar). Parse
        # with acorn; if broken, fall back to static rendering (which
        # strips/rejects the script) rather than shipping a silently
        # dead calculator.
        if fmt is DocFormat.XLSX and render_mode == "interactive":
            from backend.app.js_validate import validate_inline_scripts

            js_errors = validate_inline_scripts(p.html)
            if js_errors:
                log.warning(
                    "remediation attempt %d produced invalid inline JS; "
                    "coercing render_mode to static: %s",
                    attempt + 1,
                    js_errors,
                )
                p.html, static_issues = _fix_and_check(p.html, render_mode="static")
                if static_issues:
                    last_error = RuntimeError(
                        "interactive transcript had invalid JS "
                        f"({'; '.join(js_errors)}) and failed static fallback: "
                        f"{', '.join(static_issues)}"
                    )
                    continue
                render_mode = "static"
        p.render_mode = render_mode
        parsed = p
        break

    if parsed is None:
        assert last_error is not None
        # Chunked fallback is PDF-only: DOCX/XLSX are usually one-shot-sized,
        # and chunked synthesis assumes the <main>-wrapper stitching pattern
        # that was designed around PDF plans.
        if fmt is DocFormat.PDF and _is_likely_truncation(last_error) and plan_json:
            for attempt_size in (3, 1):
                log.info(
                    "synthesis JSON parse failed, falling back to chunked synthesis "
                    "(chunk_size=%d) for %s",
                    attempt_size,
                    source_hint,
                )
                try:
                    return await remediate_pdf_chunked(
                        doc_bytes,
                        source_hint,
                        plan_json,
                        image_filenames=image_filenames,
                        chunk_size=attempt_size,
                    )
                except RuntimeError as e:
                    if not _is_likely_truncation(e):
                        raise
                    last_error = e
                    log.warning(
                        "chunked synthesis at chunk_size=%d still truncated for %s",
                        attempt_size,
                        source_hint,
                    )
        raise last_error

    return RemediationResult(
        html=parsed.html,
        outline=[n.model_dump() for n in parsed.outline],
        title=parsed.title,
        description=parsed.description,
        language=parsed.language,
        page_count=parsed.page_count,
        render_mode=parsed.render_mode,
    )


async def remediate_pdf(
    pdf_bytes: bytes,
    source_hint: str,
    *,
    image_filenames: list[ImageAssetLike] | None = None,
    plan_json: str | None = None,
) -> RemediationResult:
    """Back-compat wrapper — tests and older callers still use this name."""
    return await remediate_document(
        pdf_bytes,
        source_hint,
        fmt=DocFormat.PDF,
        image_filenames=image_filenames,
        plan_json=plan_json,
    )


async def _remediate_chunk(
    pdf_bytes: bytes,
    source_hint: str,
    chunk_plan_json: str,
    *,
    image_filenames: list[ImageAssetLike] | None = None,
) -> RemediationChunkOutput:
    """Render a subset of plan sections as bare <section> blocks (no <main> wrapper)."""
    settings = get_settings()
    client = get_genai_client()

    pdf_io = io.BytesIO(pdf_bytes)
    uploaded = await asyncio.to_thread(
        client.files.upload,
        file=pdf_io,
        config={"mime_type": "application/pdf", "display_name": source_hint[:200]},
    )

    from google.genai import types as gtypes

    image_filenames = image_asset_filenames(image_filenames)
    images_clause = (
        f" Extracted image(s): {', '.join(Path(f).stem for f in image_filenames)}."
        if image_filenames
        else " No extractable images."
    )

    base_prompt = (
        "Render ONLY the sections in this DocumentPlan chunk as <section> blocks. "
        "Reproduce EVERY paragraph, list item, and table cell from the source verbatim. "
        "Do not summarise, abridge, or omit content. Do not say 'see the original' or "
        "'refer to the source document' — reproduce the actual text.\n"
        f"Source: {source_hint}.{images_clause}\n\n"
        f"DocumentPlan chunk (JSON):\n{chunk_plan_json}"
    )
    escalated_prompt = (
        "Your previous response was rejected because it contained placeholder or summary "
        "language instead of reproducing the source content. This is a hard error. "
        "Render every paragraph, sentence, and table cell of the source verbatim. "
        "Do NOT output phrases like 'refer to the original', 'for brevity', "
        "'continues in the full document', or any other disclaimer. "
        "If a section is long, output ALL of it — the token budget is 65,536 output tokens.\n"
        + base_prompt
    )

    def _call(*, temperature: float, prompt: str) -> Any:
        return client.models.generate_content(
            model=settings.gemini_remediation_model,
            contents=[uploaded, prompt],
            config=gtypes.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=RemediationChunkOutput,
                system_instruction=prompts.REMEDIATION_CHUNK_SYSTEM,
                temperature=temperature,
                max_output_tokens=65_536,
                thinking_config=thinking_config(settings.gemini_remediation_model),
            ),
        )

    last_error: Exception | None = None
    for attempt, (temp, prompt) in enumerate(
        [(0.2, base_prompt), (0.0, escalated_prompt), (0.0, escalated_prompt)]
    ):
        try:
            response = await asyncio.wait_for(
                asyncio.to_thread(_call, temperature=temp, prompt=prompt),
                timeout=settings.gemini_call_timeout,
            )
        except TimeoutError:
            last_error = RuntimeError(
                f"chunk Gemini call timed out after {settings.gemini_call_timeout}s"
            )
            log.warning("chunk attempt %d timed out", attempt + 1)
            continue
        try:
            chunk: RemediationChunkOutput | None = response.parsed  # type: ignore[assignment]
            if chunk is None:
                chunk = RemediationChunkOutput.model_validate_json(response.text)
        except (ValidationError, ValueError) as e:
            last_error = e
            log.warning("chunk parse attempt %d failed: %s", attempt + 1, e)
            continue
        if _looks_like_placeholder_language(chunk.html):
            last_error = RuntimeError(
                "chunk output contained placeholder/summary language — retrying with "
                "escalated prompt"
            )
            log.warning(
                "chunk attempt %d produced placeholder/summary language; retrying",
                attempt + 1,
            )
            continue
        return chunk

    raise RuntimeError(f"chunk remediation returned invalid JSON: {last_error}") from last_error


async def remediate_pdf_chunked(
    pdf_bytes: bytes,
    source_hint: str,
    plan_json: str,
    *,
    image_filenames: list[ImageAssetLike] | None = None,
    chunk_size: int = 5,
) -> RemediationResult:
    """Split-and-stitch synthesis for PDFs that exceed the per-call output token budget.

    Splits plan sections into batches of `chunk_size`, renders each batch as bare
    <section> blocks, then assembles the full <main> with skip link and <h1>.
    """
    plan = json.loads(plan_json)
    sections = plan.get("sections", [])
    title = plan.get("title", "Document")
    language = plan.get("language", "en")
    page_count = plan.get("page_count", 1)
    description = plan.get("description", "")

    if not sections:
        log.warning("remediate_pdf_chunked received a plan with zero sections for %s", source_hint)
        skip = '<a href="#content" class="sr-skip">Skip to main content</a>'
        h1 = f'<h1 id="document-title">{title}</h1>'
        body = (
            '<section id="too-large">'
            '<h2 id="too-large">Document too large to render inline</h2>'
            '<p>This PDF could not be rendered as accessible HTML in the demo because '
            'it exceeds the model\'s output budget. The original PDF remains available '
            'for download from the page that linked here.</p>'
            '</section>'
        )
        full_html = f'<main id="content" lang="{language}">\n{skip}\n{h1}\n{body}\n</main>'
        return RemediationResult(
            html=full_html,
            outline=[
                {"id": "document-title", "level": 1, "text": title},
                {"id": "too-large", "level": 2, "text": "Document too large to render inline"},
            ],
            title=title,
            description=description,
            language=language,
            page_count=page_count,
        )

    all_html_parts: list[str] = []
    all_outline: list[dict[str, Any]] = []

    consecutive_failures = 0
    aborted = False
    for start in range(0, max(len(sections), 1), chunk_size):
        if consecutive_failures >= 3:
            log.warning(
                "aborting chunked synthesis for %s after %d consecutive failures",
                source_hint,
                consecutive_failures,
            )
            aborted = True
            break
        chunk_sections = sections[start : start + chunk_size]
        chunk_plan = {**plan, "sections": chunk_sections}
        try:
            chunk = await _remediate_chunk(
                pdf_bytes,
                source_hint,
                json.dumps(chunk_plan),
                image_filenames=image_filenames,
            )
            all_html_parts.append(chunk.html)
            all_outline.extend(n.model_dump() for n in chunk.outline)
            consecutive_failures = 0
        except RuntimeError as e:
            if not _is_likely_truncation(e):
                raise
            consecutive_failures += 1
            log.warning(
                "chunk failed for sections %s in %s — emitting placeholders",
                [s.get("id") for s in chunk_sections],
                source_hint,
            )
            for sec in chunk_sections:
                sid = sec.get("id", "section")
                heading = sec.get("heading", "Section")
                placeholder = (
                    f'<section id="{sid}">'
                    f'<h2 id="{sid}">{heading}</h2>'
                    f'<p><em>This section\'s content was too large to render inline. '
                    f'See the original PDF for the full content.</em></p>'
                    f'</section>'
                )
                all_html_parts.append(placeholder)
                all_outline.append({"id": sid, "level": 2, "text": heading})

    if aborted:
        # Replace any rendered content with a single "too large" page so the user
        # sees a coherent message rather than a half-rendered document.
        skip = '<a href="#content" class="sr-skip">Skip to main content</a>'
        h1 = f'<h1 id="document-title">{title}</h1>'
        body = (
            '<section id="too-large">'
            '<h2 id="too-large">Document too large to render inline</h2>'
            '<p>This PDF could not be rendered as accessible HTML in the demo because '
            'it exceeds the model\'s output budget. The original PDF remains available '
            'for download from the page that linked here.</p>'
            '</section>'
        )
        full_html = f'<main id="content" lang="{language}">\n{skip}\n{h1}\n{body}\n</main>'
        return RemediationResult(
            html=full_html,
            outline=[
                {"id": "document-title", "level": 1, "text": title},
                {"id": "too-large", "level": 2, "text": "Document too large to render inline"},
            ],
            title=title,
            description=description,
            language=language,
            page_count=page_count,
        )

    skip = '<a href="#content" class="sr-skip">Skip to main content</a>'
    h1 = f'<h1 id="document-title">{title}</h1>'
    full_html = (
        f'<main id="content" lang="{language}">\n'
        f"{skip}\n{h1}\n"
        + "".join(all_html_parts)
        + "\n</main>"
    )

    full_html, issues = _fix_and_check(full_html)
    if issues:
        raise RuntimeError(f"chunked remediation failed structural checks: {', '.join(issues)}")

    doc_outline = [{"id": "document-title", "level": 1, "text": title}] + all_outline

    return RemediationResult(
        html=full_html,
        outline=doc_outline,
        title=title,
        description=description,
        language=language,
        page_count=page_count,
    )


async def remediate_pdf_stream(
    pdf_bytes: bytes, source_hint: str
) -> AsyncIterator[IngestEvent]:
    """Yield progress events so the viewer can render a progress bar.

    Stages and their target percentages:
      fetched   10   (caller emits this before handing bytes to us)
      uploaded  35   file uploaded to Gemini
      indexed   55   FileSearchStore populated (done by caller / file_search.ensure_store)
      generating 70  generate_content kicked off
      done      100  RemediationResult ready
    """
    yield IngestEvent("uploaded", 35, "Uploading to Gemini")
    yield IngestEvent("generating", 70, "Remediating document")

    result = await remediate_pdf(pdf_bytes, source_hint)
    yield IngestEvent(
        "done",
        100,
        "Ready",
        payload={
            "html": result.html,
            "outline": result.outline,
            "title": result.title,
            "description": result.description,
            "language": result.language,
            "page_count": result.page_count,
        },
    )
