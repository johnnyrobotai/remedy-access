from __future__ import annotations

import asyncio
import base64
import io
import json
import re
import tempfile
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field

from backend.app.config import get_settings
from backend.app.documents import DocFormat
from backend.app.extract import extract_docx_as_markdown
from backend.app.gemini import remediate as remediate_mod
from backend.app.gemini.design_v2 import plan_layout_v2
from backend.app.gemini.parse_v2 import make_document_signals, parse_document_v2
from backend.app.images import ImageAssetLike
from backend.app.layout_plan_v2 import (
    BlockContainer,
    BlockPlacement,
    HeroTreatment,
    LayoutDensity,
    LayoutPlan,
    RailSide,
    SectionContainer,
    TocPolicy,
    validate_layout_plan_for_document,
)
from backend.app.ollama_client import chat_json as ollama_chat_json
from backend.app.parser_backends.liteparse_adapter import (
    parse_pdf_with_liteparse,
)
from backend.app.parser_backends.llamaparse_adapter import (
    LlamaParseDocument,
    LlamaParseOutputError,
    parse_pdf_with_llamaparse,
)
from backend.app.render_v2 import RenderedDocument, render_document
from backend.app.structured_parser import normalize_document_assets
from backend.app.structured_parser.builder import FormField, RawBlockSignal
from backend.app.structured_parser.types import ParsedDocument

_LIST_ITEM_RE = re.compile(r"^\s*(?:[-*•]\s+|\d+[\.\)]\s+|[A-Za-z][\.\)]\s+)(.+?)\s*$")
_CALL_OUT_RE = re.compile(r"^(note|warning|caution|important|tip)[:\s-]*$", re.IGNORECASE)
_REFERENCE_HEADINGS = {"references", "reference", "bibliography", "works cited", "citations"}
_REVISION_LINE_RE = re.compile(r"^(?:\d+\s+)?revised\b", re.IGNORECASE)
_NUMBERED_FIELD_RE = re.compile(r"^\s*(?:\d+[\.\)]\s*)?(.+?):\s*(.*)$")
_FILL_LINE_RE = re.compile(r"(?P<label>[A-Za-z][A-Za-z0-9 /&().,'?-]*?)\s*_{3,}")
_BRACKET_FILL_RE = re.compile(r"\\?\[\s*(?:_{2,}|\*{2,})?\s*\\?\]")
_CHECK_OPTION_RE = re.compile(r"(?:_{2,}|□|☐|☑|☒|\[\s*[xX ]?\s*\])\s*(?P<label>[^_□☐☑☒]+)")
_FIGURE_START_RE = re.compile(r"^(figure|fig\.)\s*\d+[A-Za-z]?\s*[:.\-]\s*(.*)$", re.IGNORECASE)
_XLSX_CELL_REF_RE = re.compile(r"\$?[A-Z]{1,3}\$?\d+", re.IGNORECASE)
_XLSX_RANGE_REF_RE = re.compile(
    r"(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?\d+)",
    re.IGNORECASE,
)
_MAX_PARAGRAPH_SIGNAL_TEXT = 18_000
_TITLE_SKIP_PATTERNS = (
    re.compile(r"^provided proper attribution", re.IGNORECASE),
    re.compile(r"^reproduce the tables", re.IGNORECASE),
    re.compile(r"^scholarly works", re.IGNORECASE),
    re.compile(r"^\d+\s+revised\b", re.IGNORECASE),
    re.compile(r"^abstract$", re.IGNORECASE),
)
SignalSource = Literal["extract", "liteparse", "llamaparse"]

LAYOUT_CRITIC_SYSTEM = """\
You are the layout critic for an accessible document-to-web pipeline.
You receive:
- a rendered screenshot of the current transcript view
- a constrained ParsedDocument summary
- the current LayoutPlan

You may suggest only layout changes. You may not change text, section order,
block identity, or content semantics.

Allowed changes:
- hero_treatment
- density
- toc_policy
- section container
- rail_side
- block placement
- block container

Return a PATCH only, not a full layout plan.

Rules:
- Only include fields you want to change.
- Only include sections that need changes.
- Inside a changed section, include only blocks that need changes.
- If you cannot point to a specific `section_id` or `block_index`, do not propose a change.
- Prefer a no-op unless the screenshot shows a clear, visible layout problem.
- A change must be justified by a concrete visual issue, not a stylistic preference.
- Never return the full current layout back to the caller.
- Do not return `layout_plan`, `parsed_document`, `title`, `schema_version`, or any unrelated keys.
- Do not invent sections or blocks. Refer only to provided section ids and block indexes.

Valid minimal no-op example:
{
  "summary": "No beneficial layout changes because the current layout is already readable.",
  "issues": [],
  "sections": []
}

Valid patch example:
{
  "summary": "Move the figure into the rail and tighten the overall density.",
  "issues": [
    {
      "area": "figure_placement",
      "severity": "medium",
      "note": "The figure interrupts the main text flow."
    }
  ],
  "density": "dense",
  "sections": [
    {
      "section_id": "results",
      "container": "flow-with-rail",
      "rail_side": "right",
      "blocks": [
        {
          "block_index": 1,
          "placement": "rail",
          "container": "figure"
        }
      ]
    }
  ]
}
"""


def _strip_json_fences(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        text = text.removeprefix("```json").removeprefix("```").removesuffix("```")
    return text.strip()


class LayoutCritiqueIssue(BaseModel):
    model_config = ConfigDict(extra="ignore")

    area: Literal["hierarchy", "whitespace", "figure_placement", "table_density", "form_clarity"]
    severity: Literal["low", "medium", "high"]
    note: str


class LayoutCritiqueBlockPatch(BaseModel):
    model_config = ConfigDict(extra="ignore")

    block_index: int = Field(ge=0)
    placement: BlockPlacement | None = None
    container: BlockContainer | None = None


class LayoutCritiqueSectionPatch(BaseModel):
    model_config = ConfigDict(extra="ignore")

    section_id: str
    container: SectionContainer | None = None
    rail_side: RailSide | None = None
    blocks: list[LayoutCritiqueBlockPatch] = Field(default_factory=list)


class LayoutCritiqueResponse(BaseModel):
    model_config = ConfigDict(extra="ignore")

    summary: str = ""
    issues: list[LayoutCritiqueIssue] = Field(default_factory=list)
    hero_treatment: HeroTreatment | None = None
    density: LayoutDensity | None = None
    toc_policy: TocPolicy | None = None
    sections: list[LayoutCritiqueSectionPatch] = Field(default_factory=list)


@dataclass
class LayoutCritiqueAttempt:
    critique: LayoutCritiqueResponse | None
    attempts: list[dict[str, Any]]


@dataclass
class FormTableModel:
    caption: str
    headers: list[str]
    rows: list[list[str]] | None = None
    field_rows: list[list[FormField]] | None = None
    row_headers: list[str] | None = None


@dataclass
class FormSectionModel:
    legend: str
    fields: list[FormField]
    intro: str | None = None
    tables: list[FormTableModel] | None = None


@dataclass
class FormDocumentModel:
    title: str
    language: str
    page_count: int
    kicker: str | None
    subtitle: str | None
    note: str | None
    description: str
    sections: list[FormSectionModel]
    trailing_sections: list[tuple[str, str]]
    instructions_first: bool = False


async def run_structured_render_v2(
    doc_bytes: bytes,
    *,
    source_hint: str,
    fmt: DocFormat,
    image_filenames: Sequence[ImageAssetLike] | None = None,
    sha256: str | None = None,
    pdf_backend: str = "native",
) -> dict[str, Any]:
    parsed_document_override = None
    if fmt is DocFormat.PDF:
        direct_form = await asyncio.to_thread(
            _try_render_acroform_facsimile,
            doc_bytes,
            source_hint=source_hint,
        )
        if direct_form is not None:
            return direct_form
        signals = await _build_pdf_signals(
            doc_bytes,
            source_hint=source_hint,
            image_filenames=image_filenames,
            pdf_backend=pdf_backend,
            sha256=sha256,
        )
    elif fmt is DocFormat.DOCX:
        direct_form = await asyncio.to_thread(
            _try_render_docx_form_facsimile,
            doc_bytes,
            source_hint=source_hint,
        )
        if direct_form is not None:
            return direct_form
        signals = await asyncio.to_thread(
            _build_docx_signals,
            doc_bytes,
            source_hint=source_hint,
            image_filenames=image_filenames,
        )
    elif fmt is DocFormat.XLSX:
        signals, has_formulas = await asyncio.to_thread(
            _build_xlsx_signals,
            doc_bytes,
            source_hint=source_hint,
            image_filenames=image_filenames,
        )
        if has_formulas:
            parsed_document_override = parse_document_v2(signals)
            direct_calculator = await asyncio.to_thread(
                _try_render_xlsx_calculator,
                doc_bytes,
                source_hint=source_hint,
            )
            if direct_calculator is not None:
                direct_calculator["parsed_document"] = parsed_document_override
                return direct_calculator
    else:
        raise ValueError(f"structured_render_v2 does not support {fmt.value} inputs")
    parsed_document = (
        parsed_document_override
        if parsed_document_override is not None
        else parse_document_v2(signals)
    )
    layout_plan = await plan_layout_v2(
        parsed_document,
        source_hint=source_hint,
        fmt=fmt,
    )
    if fmt is DocFormat.XLSX and 'has_formulas' in locals() and has_formulas:
        result = await remediate_mod.remediate_document(
            doc_bytes,
            source_hint=source_hint,
            fmt=fmt,
            image_filenames=[],
            plan_json=None,
        )
        return {
            "result": result,
            "parsed_document": parsed_document,
            "layout_plan": layout_plan,
        }
    render_input = _render_input_from_parsed(
        parsed_document,
        image_assets=image_filenames,
        sha256=sha256,
    )
    render_layout = _render_layout_from_v2_plan(layout_plan, parsed_document)
    result = render_document(render_input, render_layout)
    result, layout_plan = await _maybe_apply_layout_agent_revision(
        parsed_document=parsed_document,
        render_input=render_input,
        layout_plan=layout_plan,
        rendered=result,
        source_hint=source_hint,
        fmt=fmt,
        sha256=sha256,
    )
    return {
        "result": result,
        "parsed_document": parsed_document,
        "layout_plan": layout_plan,
    }


async def _maybe_apply_layout_agent_revision(
    *,
    parsed_document: ParsedDocument,
    render_input: dict[str, Any],
    layout_plan: LayoutPlan,
    rendered: RenderedDocument,
    source_hint: str,
    fmt: DocFormat,
    sha256: str | None,
) -> tuple[RenderedDocument, LayoutPlan]:
    settings = get_settings()
    if not settings.layout_agent_enabled or settings.layout_agent_max_revisions < 1:
        return rendered, layout_plan
    apply_to_formats = getattr(settings, "layout_agent_apply_to_formats", frozenset({"pdf"}))
    if fmt.value not in apply_to_formats:
        await _persist_layout_agent_artifacts(
            sha256=sha256,
            initial_layout=layout_plan,
            critique_attempt=LayoutCritiqueAttempt(critique=None, attempts=[]),
            revised_layout=layout_plan,
            before_screenshot=None,
            after_screenshot=None,
            skip_reason=f"layout agent disabled for format {fmt.value}",
        )
        return rendered, layout_plan
    if rendered.render_mode != "static":
        await _persist_layout_agent_artifacts(
            sha256=sha256,
            initial_layout=layout_plan,
            critique_attempt=LayoutCritiqueAttempt(critique=None, attempts=[]),
            revised_layout=layout_plan,
            before_screenshot=None,
            after_screenshot=None,
            skip_reason="render_mode is not static",
        )
        return rendered, layout_plan
    if not layout_plan.sections:
        await _persist_layout_agent_artifacts(
            sha256=sha256,
            initial_layout=layout_plan,
            critique_attempt=LayoutCritiqueAttempt(critique=None, attempts=[]),
            revised_layout=layout_plan,
            before_screenshot=None,
            after_screenshot=None,
            skip_reason="layout plan has no sections",
        )
        return rendered, layout_plan
    if sum(len(section.blocks) for section in layout_plan.sections) < 2:
        await _persist_layout_agent_artifacts(
            sha256=sha256,
            initial_layout=layout_plan,
            critique_attempt=LayoutCritiqueAttempt(critique=None, attempts=[]),
            revised_layout=layout_plan,
            before_screenshot=None,
            after_screenshot=None,
            skip_reason="layout plan has fewer than two blocks across sections",
        )
        return rendered, layout_plan

    before_bytes = await _capture_layout_screenshot(rendered.html)
    if before_bytes is None:
        await _persist_layout_agent_artifacts(
            sha256=sha256,
            initial_layout=layout_plan,
            critique_attempt=LayoutCritiqueAttempt(critique=None, attempts=[]),
            revised_layout=layout_plan,
            before_screenshot=None,
            after_screenshot=None,
            skip_reason="failed to capture screenshot",
        )
        return rendered, layout_plan

    critique_attempt = await _request_layout_critique(
        parsed_document=parsed_document,
        layout_plan=layout_plan,
        screenshot_bytes=before_bytes,
        source_hint=source_hint,
        fmt=fmt,
    )
    critique = critique_attempt.critique
    if critique is None:
        await _persist_layout_agent_artifacts(
            sha256=sha256,
            initial_layout=layout_plan,
            critique_attempt=critique_attempt,
            revised_layout=layout_plan,
            before_screenshot=before_bytes,
            after_screenshot=before_bytes,
            skip_reason="critique returned no usable patch",
        )
        return rendered, layout_plan

    revised_layout = _apply_layout_critique(layout_plan, critique, parsed_document)
    revised_render = rendered
    after_bytes = before_bytes
    if revised_layout.model_dump(mode="json") != layout_plan.model_dump(mode="json"):
        revised_render = render_document(
            render_input,
            _render_layout_from_v2_plan(revised_layout, parsed_document),
        )
        captured_after = await _capture_layout_screenshot(revised_render.html)
        if captured_after is not None:
            after_bytes = captured_after

    await _persist_layout_agent_artifacts(
        sha256=sha256,
        initial_layout=layout_plan,
        critique_attempt=critique_attempt,
        revised_layout=revised_layout,
        before_screenshot=before_bytes,
        after_screenshot=after_bytes,
        skip_reason=None,
    )
    return revised_render, revised_layout


async def _request_layout_critique(
    *,
    parsed_document: ParsedDocument,
    layout_plan: LayoutPlan,
    screenshot_bytes: bytes,
    source_hint: str,
    fmt: DocFormat,
) -> LayoutCritiqueAttempt:
    settings = get_settings()
    prompt = _build_layout_critique_prompt(
        parsed_document=parsed_document,
        layout_plan=layout_plan,
        source_hint=source_hint,
        fmt=fmt,
    )
    image_b64 = base64.b64encode(screenshot_bytes).decode("ascii")
    attempts: list[dict[str, Any]] = []
    for _attempt in range(2):
        try:
            content, _raw = await asyncio.wait_for(
                ollama_chat_json(
                    model=settings.layout_agent_model,
                    system_instruction=LAYOUT_CRITIC_SYSTEM,
                    user_prompt=prompt,
                    schema=LayoutCritiqueResponse.model_json_schema(),
                    images=[image_b64],
                    temperature=0.0,
                    think=settings.ollama_reasoning_level,
                ),
                timeout=settings.layout_agent_timeout_s,
            )
            attempt_info = {
                "raw_content": content,
                "raw_response": _raw,
                "error": None,
            }
            attempts.append(attempt_info)
            try:
                critique = LayoutCritiqueResponse.model_validate_json(_strip_json_fences(content))
                return LayoutCritiqueAttempt(critique=critique, attempts=attempts)
            except Exception as e:  # noqa: BLE001
                attempt_info["error"] = str(e)
                continue
        except Exception as e:  # noqa: BLE001
            attempts.append({"raw_content": None, "raw_response": None, "error": str(e)})
    if attempts:
        # Best-effort feature; do not fail ingest on critique errors.
        import logging

        logging.getLogger(__name__).warning("layout critique failed: %s", attempts[-1]["error"])
    return LayoutCritiqueAttempt(critique=None, attempts=attempts)


def _apply_layout_critique(
    layout_plan: LayoutPlan,
    critique: LayoutCritiqueResponse,
    parsed_document: ParsedDocument,
) -> LayoutPlan:
    updated = layout_plan.model_copy(deep=True)
    if critique.hero_treatment is not None:
        updated.hero_treatment = critique.hero_treatment
    if critique.density is not None:
        updated.density = critique.density
    if critique.toc_policy is not None:
        updated.toc_policy = critique.toc_policy

    section_by_id = {section.section_id: section for section in updated.sections}
    for section_patch in critique.sections:
        section = section_by_id.get(section_patch.section_id)
        if section is None:
            raise ValueError(f"unknown section_id in critique: {section_patch.section_id}")
        if section_patch.container is not None:
            section.container = section_patch.container
        if section_patch.rail_side is not None:
            section.rail_side = section_patch.rail_side
        block_by_index = {block.block_index: block for block in section.blocks}
        for block_patch in section_patch.blocks:
            block = block_by_index.get(block_patch.block_index)
            if block is None:
                raise ValueError(
                    f"unknown block_index {block_patch.block_index} for section {section_patch.section_id}"
                )
            if block_patch.placement is not None:
                block.placement = block_patch.placement
            if block_patch.container is not None:
                block.container = block_patch.container
    return validate_layout_plan_for_document(updated, parsed_document)


def _build_layout_critique_prompt(
    *,
    parsed_document: ParsedDocument,
    layout_plan: LayoutPlan,
    source_hint: str,
    fmt: DocFormat,
) -> str:
    schema_json = json.dumps(LayoutCritiqueResponse.model_json_schema(), indent=2)
    summary = {
        "title": parsed_document.title,
        "doc_kind": parsed_document.doc_kind,
        "source_format": parsed_document.source_format,
        "sections": [
            {
                "id": section.id,
                "heading": section.heading,
                "blocks": [
                    {"index": idx, "type": block.type}
                    for idx, block in enumerate(section.blocks)
                ],
            }
            for section in _iter_sections_for_summary(parsed_document.body)
        ],
        "layout_plan": layout_plan.model_dump(mode="json"),
    }
    return (
        f"Source hint: {source_hint}\n"
        f"Format: {fmt.value}\n"
        "Review the rendered transcript screenshot and suggest only layout improvements.\n"
        "Return a constrained patch. Use only the provided section ids and block indexes.\n"
        "Critique response JSON schema:\n"
        + schema_json
        + "\nParsed document summary and current layout:\n"
        + json.dumps(summary, indent=2)
    )


def _iter_sections_for_summary(body: Sequence[Any]) -> list[Any]:
    sections: list[Any] = []
    for block in body:
        if getattr(block, "type", "") == "section":
            sections.append(block)
            sections.extend(_iter_sections_for_summary(getattr(block, "blocks", [])))
    return sections


async def _capture_layout_screenshot(html: str) -> bytes | None:
    return await asyncio.to_thread(_capture_layout_screenshot_sync, html)


def _capture_layout_screenshot_sync(html: str) -> bytes | None:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return None

    settings = get_settings()
    css_path = Path(__file__).resolve().parents[2] / "viewer" / "src" / "styles.css"
    css = css_path.read_text(encoding="utf-8") if css_path.exists() else ""
    doc = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        f"<base href='{settings.public_origin.rstrip('/')}/'>"
        f"<style>{css}</style></head><body><div class='transcript-pane'>{html}</div></body></html>"
    )
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": 1440, "height": 1100})
            page.set_content(doc, wait_until="load")
            page.wait_for_load_state("networkidle", timeout=30000)
            image = page.screenshot(full_page=False)
            browser.close()
            return image
    except Exception:
        return None


async def _persist_layout_agent_artifacts(
    *,
    sha256: str | None,
    initial_layout: LayoutPlan,
    critique_attempt: LayoutCritiqueAttempt,
    revised_layout: LayoutPlan,
    before_screenshot: bytes | None,
    after_screenshot: bytes | None,
    skip_reason: str | None,
) -> None:
    if not sha256:
        return
    settings = get_settings()
    artifact_dir = settings.artifacts_dir_for(sha256) / "layout-agent"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    await asyncio.to_thread(
        (artifact_dir / "initial_layout_plan.json").write_text,
        json.dumps(initial_layout.model_dump(mode="json"), indent=2),
        "utf-8",
    )
    await asyncio.to_thread(
        (artifact_dir / "revised_layout_plan.json").write_text,
        json.dumps(revised_layout.model_dump(mode="json"), indent=2),
        "utf-8",
    )
    await asyncio.to_thread(
        (artifact_dir / "critique.json").write_text,
        json.dumps(
            critique_attempt.critique.model_dump(mode="json") if critique_attempt.critique is not None else {},
            indent=2,
        ),
        "utf-8",
    )
    await asyncio.to_thread(
        (artifact_dir / "critique_attempts.json").write_text,
        json.dumps(critique_attempt.attempts, indent=2),
        "utf-8",
    )
    await asyncio.to_thread(
        (artifact_dir / "skip_reason.json").write_text,
        json.dumps({"skip_reason": skip_reason}, indent=2),
        "utf-8",
    )
    if before_screenshot is not None:
        await asyncio.to_thread((artifact_dir / "before.png").write_bytes, before_screenshot)
    if after_screenshot is not None:
        await asyncio.to_thread((artifact_dir / "after.png").write_bytes, after_screenshot)


async def _build_pdf_signals(
    doc_bytes: bytes,
    *,
    source_hint: str,
    image_filenames: Sequence[ImageAssetLike] | None,
    pdf_backend: str,
    sha256: str | None = None,
):
    if pdf_backend == "llamaparse":
        settings = get_settings()
        llama_doc = await _extract_llamaparse_document(doc_bytes, source_hint=source_hint)
        await _persist_llamaparse_artifact(sha256=sha256, document=llama_doc)
        title_hint, blocks = _raw_blocks_from_llamaparse_document(
            llama_doc,
            source_hint=source_hint,
        )
        page_texts = [
            (page.page_number, page.markdown or page.text or "")
            for page in llama_doc.pages
        ]
        title_hint = title_hint or _choose_document_title(
            page_texts,
            source_hint=source_hint,
            metadata_title=None,
        )
        if not blocks:
            raise LlamaParseOutputError("LlamaParse returned no usable markdown or text blocks")
        has_form_signals = any(block.kind == "form_field" for block in blocks)
        return make_document_signals(
            format="pdf",
            title=title_hint,
            language="en",
            page_count=max(1, len(llama_doc.pages)),
            blocks=blocks,
            image_assets=[] if has_form_signals else list(image_filenames or []),
            metadata={
                "pdf_backend": "llamaparse",
                "llamaparse_tier": settings.llamaparse_tier,
                "llamaparse_version": settings.llamaparse_version,
            },
        )

    if pdf_backend == "liteparse":
        page_texts = await _extract_liteparse_page_texts(doc_bytes)
        metadata_title = None
        acroform_fields: list[FormField] = []
    else:
        page_texts, metadata_title, acroform_fields = await asyncio.to_thread(
            _extract_native_pdf_signals, doc_bytes
        )

    signal_source: SignalSource = (
        "liteparse" if pdf_backend == "liteparse" else "extract"
    )
    title_hint = _choose_document_title(
        page_texts,
        source_hint=source_hint,
        metadata_title=metadata_title,
    )
    has_acroform = bool(acroform_fields)
    blocks = _raw_blocks_from_page_texts(
        page_texts,
        source=signal_source,
        title_hint=title_hint,
        disable_form_heuristics=has_acroform,
    )
    if acroform_fields:
        blocks.extend(_acroform_fields_to_signals(acroform_fields, start_order=len(blocks)))
    has_form_signals = has_acroform or any(block.kind == "form_field" for block in blocks)
    signal_kwargs: dict[str, Any] = {
        "format": "pdf",
        "title": title_hint,
        "language": "en",
        "page_count": max(1, len(page_texts)),
        "image_assets": [] if has_form_signals else list(image_filenames or []),
        "metadata": {"pdf_backend": pdf_backend},
    }
    if signal_source == "liteparse":
        signal_kwargs["liteparse_blocks"] = blocks
    else:
        signal_kwargs["blocks"] = blocks
    return make_document_signals(**signal_kwargs)


def _build_docx_signals(
    doc_bytes: bytes,
    *,
    source_hint: str,
    image_filenames: Sequence[ImageAssetLike] | None,
):
    markdown = extract_docx_as_markdown(doc_bytes)
    title_hint, blocks = _raw_blocks_from_markdown(markdown)
    return make_document_signals(
        format="docx",
        title=title_hint or _form_title_from_source_hint(source_hint) or _fallback_title(source_hint),
        language="en",
        page_count=1,
        blocks=blocks,
        image_assets=list(image_filenames or []),
        metadata={},
    )


def _build_xlsx_signals(
    doc_bytes: bytes,
    *,
    source_hint: str,
    image_filenames: Sequence[ImageAssetLike] | None,
):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(doc_bytes), data_only=False)
    try:
        has_formulas = _workbook_has_formula(wb)

        blocks: list[RawBlockSignal] = []
        order = 0
        for page_num, sheet_name in enumerate(wb.sheetnames, start=1):
            ws = wb[sheet_name]
            blocks.append(
                RawBlockSignal(
                    id=f"heading-{order}",
                    kind="heading",
                    order=order,
                    source="extract",
                    source_key=f"sheet-{page_num}-heading",
                    page_start=page_num,
                    page_end=page_num,
                    level=2,
                    text=sheet_name,
                )
            )
            order += 1

            rows = _xlsx_sheet_rows(ws)
            if rows:
                header_rows = [0] if len(rows) > 1 else []
                blocks.append(
                    RawBlockSignal(
                        id=f"table-{order}",
                        kind="table",
                        order=order,
                        source="extract",
                        source_key=f"sheet-{page_num}-table",
                        page_start=page_num,
                        page_end=page_num,
                        rows=rows,
                        header_rows=header_rows,
                    )
                )
                order += 1
        return make_document_signals(
            format="xlsx",
            title=_fallback_title(source_hint),
            language="en",
            page_count=max(1, len(wb.sheetnames)),
            blocks=blocks,
            image_assets=list(image_filenames or []),
            metadata={},
        ), has_formulas
    finally:
        wb.close()


def _try_render_xlsx_calculator(
    doc_bytes: bytes,
    *,
    source_hint: str,
) -> dict[str, Any] | None:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(doc_bytes), data_only=False)
    try:
        formula_cells = [
            (ws, cell)
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        if not formula_cells:
            return None

        compiled_outputs: list[dict[str, Any]] = []
        input_refs: set[str] = set()
        formula_refs: set[str] = set()
        for ws, cell in formula_cells:
            coord = _normalize_xlsx_ref(cell.coordinate)
            formula_refs.add(coord)
            expression = _compile_xlsx_formula_to_js(str(cell.value))
            if expression is None:
                return None
            refs = _xlsx_formula_references(str(cell.value))
            if not refs:
                return None
            input_refs.update(refs)
            compiled_outputs.append(
                {
                    "sheet": ws.title,
                    "coord": coord,
                    "label": _xlsx_output_label(ws, cell),
                    "expression": expression,
                    "refs": refs,
                }
            )

        input_refs.difference_update(formula_refs)
        if not input_refs:
            return None

        sheet = formula_cells[0][0]
        title = _xlsx_calculator_title(source_hint)
        html, outline = _render_xlsx_calculator_html(
            sheet=sheet,
            title=title,
            input_refs=input_refs,
            outputs=compiled_outputs,
            page_count=max(1, len(wb.sheetnames)),
        )
        result = remediate_mod.RemediationResult(
            html=html,
            outline=outline,
            title=title,
            description="Interactive HTML spreadsheet calculator.",
            language="en",
            page_count=max(1, len(wb.sheetnames)),
            render_mode="interactive",
        )
        return {
            "result": result,
            "parsed_document": {
                "source_format": "xlsx",
                "title": title,
                "language": "en",
                "page_count": result.page_count,
                "doc_kind": "calculator",
                "renderer": "xlsx-calculator",
            },
            "layout_plan": {
                "template": "calculator",
                "visual_density": "dense",
                "toc_policy": "hidden",
                "renderer": "xlsx-calculator",
            },
        }
    finally:
        wb.close()


def _xlsx_calculator_title(source_hint: str) -> str:
    title = _fallback_title(source_hint)
    return re.sub(r"\bGpa\b", "GPA", title)


def _render_xlsx_calculator_html(
    *,
    sheet,
    title: str,
    input_refs: set[str],
    outputs: Sequence[dict[str, Any]],
    page_count: int,
) -> tuple[str, list[dict[str, Any]]]:
    sorted_inputs = sorted(input_refs, key=_xlsx_ref_sort_key)
    input_rows = sorted({row for row, _col in (_xlsx_ref_row_col(ref) for ref in sorted_inputs)})
    input_cols = sorted({col for _row, col in (_xlsx_ref_row_col(ref) for ref in sorted_inputs)})
    min_input_row = input_rows[0]
    header_row = min_input_row - 1 if min_input_row > 1 else None
    label_col = 1 if any(_xlsx_cell_text(sheet.cell(row=row, column=1)) for row in input_rows) else min(input_cols)
    table_cols = sorted(set(input_cols) | {label_col})

    table_header = "".join(
        f'<th scope="col">{escape(_xlsx_column_label(sheet, header_row, col))}</th>'
        for col in table_cols
    )
    body_rows: list[str] = []
    for row in input_rows:
        cells: list[str] = []
        row_label = _xlsx_cell_text(sheet.cell(row=row, column=label_col))
        for col in table_cols:
            coord = _normalize_xlsx_ref(f"{_xlsx_column_letter(col)}{row}")
            value = sheet.cell(row=row, column=col).value
            if coord in input_refs:
                label = _xlsx_input_label(sheet, header_row, row, col, row_label)
                input_type = "number" if isinstance(value, int | float) else "text"
                attrs = ' inputmode="decimal" step="any"' if input_type == "number" else ""
                cells.append(
                    "<td>"
                    f'<label class="sr-only" for="{_xlsx_cell_id(coord)}">{escape(label)}</label>'
                    f'<input id="{_xlsx_cell_id(coord)}" type="{input_type}"{attrs} value="{escape(_xlsx_cell_text(sheet.cell(row=row, column=col)))}" />'
                    "</td>"
                )
            elif col == label_col:
                cells.append(f'<th scope="row">{escape(_xlsx_cell_text(sheet.cell(row=row, column=col)) or f"Row {row}")}</th>')
            else:
                cells.append(f"<td>{escape(_xlsx_cell_text(sheet.cell(row=row, column=col)))}</td>")
        body_rows.append(f"<tr>{''.join(cells)}</tr>")

    output_html = []
    update_lines = []
    for output in outputs:
        coord = output["coord"]
        output_id = _xlsx_cell_id(coord)
        label = str(output["label"])
        refs = " ".join(_xlsx_cell_id(ref) for ref in output["refs"])
        output_html.append(
            '<div class="doc-form__field doc-form__field--compact">'
            f'<label for="{output_id}">{escape(label)}</label>'
            f'<output id="{output_id}" for="{escape(refs)}"></output>'
            "</div>"
        )
        update_lines.append(f"setOutput({json.dumps(output_id)}, {output['expression']});")

    form_id = "xlsx-calculator"
    script = _xlsx_calculator_script(form_id=form_id, update_lines=update_lines)
    html = (
        f'<main id="content" lang="en" class="doc-page doc-page--calculator doc-page--density-dense">'
        '<a href="#content" class="sr-skip">Skip to main content</a>'
        f'<header class="doc-hero doc-hero--compact"><h1 id="document-title" class="doc-hero__title">{escape(title)}</h1>'
        f'<p class="doc-hero__subtitle">{page_count} sheet{"s" if page_count != 1 else ""}</p></header>'
        f'<form id="{form_id}" class="doc-form doc-form--calculator">'
        '<section class="doc-section doc-section--single-column" id="calculator-inputs" aria-labelledby="calculator-inputs-heading">'
        '<header class="doc-section__header"><h2 id="calculator-inputs-heading">Calculator inputs</h2></header>'
        '<div class="doc-table"><table>'
        f'<caption>{escape(sheet.title)} input cells</caption><thead><tr>{table_header}</tr></thead>'
        f'<tbody>{"".join(body_rows)}</tbody></table></div></section>'
        '<section class="doc-section doc-section--single-column" id="calculator-results" aria-labelledby="calculator-results-heading">'
        '<header class="doc-section__header"><h2 id="calculator-results-heading">Calculator results</h2></header>'
        f'<div class="doc-form-grid doc-form-grid--2">{"".join(output_html)}</div></section>'
        f"{script}</form></main>"
    )
    outline = [
        {"id": "document-title", "level": 1, "text": title},
        {"id": "calculator-inputs-heading", "level": 2, "text": "Calculator inputs"},
        {"id": "calculator-results-heading", "level": 2, "text": "Calculator results"},
    ]
    return html, outline


def _xlsx_calculator_script(*, form_id: str, update_lines: Sequence[str]) -> str:
    body = "\n".join(f"    {line}" for line in update_lines)
    return (
        "<script>"
        "(function(){\n"
        f"  const form = document.getElementById({json.dumps(form_id)});\n"
        "  if (!form) return;\n"
        "  const number = (id) => {\n"
        "    const el = document.getElementById(id);\n"
        "    const raw = el && 'value' in el ? el.value : el ? el.textContent : '';\n"
        "    const value = Number(raw);\n"
        "    return Number.isFinite(value) ? value : 0;\n"
        "  };\n"
        "  const sum = (values) => values.reduce((total, value) => total + value, 0);\n"
        "  const sumProduct = (left, right) => left.reduce((total, value, index) => total + value * (right[index] === undefined ? 0 : right[index]), 0);\n"
        "  const formatNumber = (value) => Number.isFinite(value) ? Number((Math.round(value * 1000) / 1000).toFixed(3)).toString() : '';\n"
        "  const setOutput = (id, value) => { const output = document.getElementById(id); if (output) output.textContent = formatNumber(value); };\n"
        "  const update = () => {\n"
        f"{body}\n"
        "  };\n"
        "  form.addEventListener('input', update);\n"
        "  window.__docboxCalc = window.__docboxCalc || {};\n"
        f"  window.__docboxCalc[{json.dumps(form_id)}] = {{ update }};\n"
        "  update();\n"
        "})();"
        "</script>"
    )


def _compile_xlsx_formula_to_js(formula: str) -> str | None:
    expression = re.sub(r"\s+", "", formula.removeprefix("=")).replace("$", "").upper()

    def sumproduct_repl(match: re.Match[str]) -> str:
        left = _expand_xlsx_range(f"{match.group(1)}:{match.group(2)}")
        right = _expand_xlsx_range(f"{match.group(3)}:{match.group(4)}")
        if not left or len(left) != len(right):
            raise ValueError("unsupported SUMPRODUCT range")
        return (
            "sumProduct(["
            + ",".join(_xlsx_number_call(ref) for ref in left)
            + "],["
            + ",".join(_xlsx_number_call(ref) for ref in right)
            + "])"
        )

    def sum_repl(match: re.Match[str]) -> str:
        refs = _expand_xlsx_range(match.group(1))
        if not refs:
            raise ValueError("unsupported SUM range")
        return "sum([" + ",".join(_xlsx_number_call(ref) for ref in refs) + "])"

    try:
        expression = re.sub(
            rf"SUMPRODUCT\({_XLSX_RANGE_REF_RE.pattern},{_XLSX_RANGE_REF_RE.pattern}\)",
            sumproduct_repl,
            expression,
            flags=re.IGNORECASE,
        )
        expression = re.sub(
            rf"SUM\(({_XLSX_RANGE_REF_RE.pattern})\)",
            sum_repl,
            expression,
            flags=re.IGNORECASE,
        )
    except ValueError:
        return None

    expression = _replace_xlsx_cell_refs_outside_strings(expression)
    probe = re.sub(r'"[^"]*"', "0", expression)
    probe = probe.replace("sumProduct", "").replace("number", "").replace("sum", "")
    if re.search(r"[A-Za-z_]", probe) or not re.fullmatch(r"[0-9+\-*/().,\[\] ]*", probe):
        return None
    return expression


def _replace_xlsx_cell_refs_outside_strings(expression: str) -> str:
    parts = re.split(r'("[^"]*")', expression)
    for index in range(0, len(parts), 2):
        parts[index] = _XLSX_CELL_REF_RE.sub(
            lambda match: _xlsx_number_call(_normalize_xlsx_ref(match.group(0))),
            parts[index],
        )
    return "".join(parts)


def _xlsx_formula_references(formula: str) -> list[str]:
    expression = formula.removeprefix("=").replace("$", "").upper()
    refs: list[str] = []
    consumed_ranges: list[tuple[int, int]] = []
    for match in _XLSX_RANGE_REF_RE.finditer(expression):
        consumed_ranges.append(match.span())
        refs.extend(_expand_xlsx_range(match.group(0)))
    for match in _XLSX_CELL_REF_RE.finditer(expression):
        if any(start <= match.start() < end for start, end in consumed_ranges):
            continue
        refs.append(_normalize_xlsx_ref(match.group(0)))
    return _dedupe(refs)


def _expand_xlsx_range(range_ref: str) -> list[str]:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    try:
        min_col, min_row, max_col, max_row = range_boundaries(range_ref.replace("$", "").upper())
    except ValueError:
        return []
    return [
        f"{get_column_letter(col)}{row}"
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    ]


def _xlsx_output_label(worksheet, cell) -> str:
    for col in range(cell.column - 1, 0, -1):
        text = _xlsx_cell_text(worksheet.cell(row=cell.row, column=col))
        if text:
            return text
    return f"Result {cell.coordinate}"


def _xlsx_input_label(worksheet, header_row: int | None, row: int, col: int, row_label: str) -> str:
    header = _xlsx_column_label(worksheet, header_row, col)
    if row_label and row_label != header:
        return f"{row_label} {header}"
    return header or f"Cell {_xlsx_column_letter(col)}{row}"


def _xlsx_column_label(worksheet, header_row: int | None, col: int) -> str:
    if header_row:
        text = _xlsx_cell_text(worksheet.cell(row=header_row, column=col))
        if text:
            return text
    return f"Column {_xlsx_column_letter(col)}"


def _xlsx_cell_text(cell) -> str:
    return "" if cell.value is None else str(cell.value)


def _xlsx_ref_row_col(ref: str) -> tuple[int, int]:
    from openpyxl.utils.cell import coordinate_to_tuple

    return coordinate_to_tuple(_normalize_xlsx_ref(ref))


def _xlsx_ref_sort_key(ref: str) -> tuple[int, int]:
    return _xlsx_ref_row_col(ref)


def _xlsx_column_letter(col: int) -> str:
    from openpyxl.utils.cell import get_column_letter

    return get_column_letter(col)


def _normalize_xlsx_ref(ref: str) -> str:
    return ref.replace("$", "").upper()


def _xlsx_cell_id(ref: str) -> str:
    return f"xlsx-cell-{_normalize_xlsx_ref(ref).lower()}"


def _xlsx_number_call(ref: str) -> str:
    return f"number({json.dumps(_xlsx_cell_id(ref))})"


async def _extract_liteparse_page_texts(doc_bytes: bytes) -> list[tuple[int, str]]:
    def _run() -> list[tuple[int, str]]:
        with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp:
            tmp.write(doc_bytes)
            tmp.flush()
            raw = parse_pdf_with_liteparse(tmp.name)
        return [(page.page_num, page.text or "") for page in raw.pages]

    return await asyncio.to_thread(_run)


async def _extract_llamaparse_document(
    doc_bytes: bytes,
    *,
    source_hint: str,
) -> LlamaParseDocument:
    settings = get_settings()
    return await parse_pdf_with_llamaparse(
        doc_bytes,
        api_key=settings.llama_cloud_api_key,
        base_url=settings.llamaparse_base_url,
        tier=settings.llamaparse_tier,
        version=settings.llamaparse_version,
        max_pages=settings.llamaparse_max_pages,
        target_pages=settings.llamaparse_target_pages,
        use_cost_optimizer=settings.llamaparse_use_cost_optimizer,
        disable_cache=settings.llamaparse_disable_cache,
        do_not_cache=settings.llamaparse_do_not_cache,
        timeout_s=settings.llamaparse_timeout_s,
        ocr_languages=settings.llamaparse_ocr_languages,
        filename=_llamaparse_filename(source_hint),
    )


async def _persist_llamaparse_artifact(
    *,
    sha256: str | None,
    document: LlamaParseDocument,
) -> None:
    if not sha256:
        return
    settings = get_settings()
    artifact_dir = settings.artifacts_dir_for(sha256) / "llamaparse"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "pages": [
            {
                "page_number": page.page_number,
                "markdown": page.markdown,
                "text": page.text,
                "items": list(page.items),
                "metadata": page.metadata,
            }
            for page in document.pages
        ],
        "markdown_full": document.markdown_full,
        "text_full": document.text_full,
        "job_metadata": document.job_metadata,
    }
    await asyncio.to_thread(
        (artifact_dir / "document.json").write_text,
        json.dumps(payload, indent=2),
        "utf-8",
    )


def _llamaparse_filename(source_hint: str) -> str:
    candidate = Path(source_hint.rstrip("/")).name or "document.pdf"
    return candidate if candidate.lower().endswith(".pdf") else f"{candidate}.pdf"


def _extract_native_page_texts(doc_bytes: bytes) -> list[tuple[int, str]]:
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(doc_bytes))
    return [
        (page_num, page.extract_text() or "")
        for page_num, page in enumerate(reader.pages, start=1)
    ]


def _extract_native_pdf_signals(doc_bytes: bytes) -> tuple[list[tuple[int, str]], str | None, list[FormField]]:
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(doc_bytes))
    meta = reader.metadata or {}
    metadata_title = getattr(meta, "title", None) if not isinstance(meta, dict) else meta.get("/Title")
    page_texts: list[tuple[int, str]] = []
    for page_num, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
        except TypeError:
            text = page.extract_text() or ""
        page_texts.append((page_num, text))
    return page_texts, _clean_metadata_title(metadata_title), _extract_acroform_fields(reader)


def _extract_acroform_fields(reader: Any) -> list[FormField]:
    try:
        raw_fields = reader.get_fields() or {}
    except Exception:  # noqa: BLE001
        return []

    fields: list[FormField] = []
    seen_ids: set[str] = set()
    positions = _acroform_field_positions(reader)
    ordered_fields = sorted(
        raw_fields.items(),
        key=lambda item: positions.get(str(item[0]), (10_000, 10_000.0, 10_000.0)),
    )
    for fallback_name, field in ordered_fields:
        field_type = str(field.get("/FT") or "").strip("/")
        if not field_type:
            continue
        label = _clean_field_label(_acroform_field_label(field, fallback_name))
        if not label:
            continue
        control = _control_type_for_acrofield(field_type, label, field)
        if control == "button":
            continue
        raw_id = _slugify(label)
        field_id = _unique_slug(raw_id, seen_ids)
        options = _field_options(field) if control in {"radio", "checkbox", "select"} else []
        if control == "radio" and not options:
            control = "checkbox"
        fields.append(
            FormField(
                id=field_id,
                label=label,
                field_type=control,
                required=bool(int(field.get("/Ff") or 0) & 2),
                options=options,
            )
        )
    return fields


def _acroform_field_label(field: Any, fallback_name: str) -> str:
    tooltip = str(field.get("/TU") or "").strip()
    technical_name = str(field.get("/T") or fallback_name or "").strip()
    if tooltip and technical_name and _should_prefer_acroform_technical_name(tooltip, technical_name):
        return technical_name
    if tooltip and tooltip.lower() != "undefined":
        return tooltip
    return technical_name or fallback_name or "Field"


def _should_prefer_acroform_technical_name(tooltip: str, technical_name: str) -> bool:
    tooltip_clean = _clean_field_label(tooltip)
    name_clean = _clean_field_label(technical_name)
    if not tooltip_clean or tooltip_clean.lower() == "undefined":
        return True
    if _looks_like_instructional_tooltip(tooltip_clean) and not _looks_like_technical_field_name(name_clean):
        return True
    if re.fullmatch(r"(?:mm|m){1,2}/(?:dd|d){1,2}/(?:yyyy|yyy|yy)", tooltip_clean, flags=re.IGNORECASE):
        return True

    tooltip_parts = _repeated_row_field_parts(tooltip_clean)
    name_parts = _repeated_row_field_parts(name_clean)
    if tooltip_parts and name_parts:
        tooltip_header, tooltip_row = tooltip_parts
        name_header, name_row = name_parts
        if _slugify(tooltip_header) == _slugify(name_header) and tooltip_row != name_row:
            return True
    if name_parts:
        name_header, _name_row = name_parts
        if _slugify(tooltip_clean) == _slugify(name_header):
            return True
    return False


def _looks_like_instructional_tooltip(label: str) -> bool:
    return bool(
        re.match(
            r"^(?:enter|select|click|choose|type|provide)\b.+\b(?:here|form|data|entered)\b",
            label.strip(),
            flags=re.IGNORECASE,
        )
    )


def _looks_like_technical_field_name(label: str) -> bool:
    return bool(
        re.match(r"^(?:pg\d+[-_]\d+|enter text|check box|text|date|signature)\b", label, flags=re.IGNORECASE)
    )


def _acroform_field_positions(reader: Any) -> dict[str, tuple[int, float, float]]:
    positions: dict[str, tuple[int, float, float]] = {}
    for page_num, page in enumerate(getattr(reader, "pages", []), start=1):
        try:
            annots = page.get("/Annots") or []
        except Exception:  # noqa: BLE001
            continue
        for annot_ref in annots:
            try:
                annot = annot_ref.get_object()
            except Exception:  # noqa: BLE001
                continue
            name = _annotation_field_name(annot)
            if not name:
                continue
            rect = annot.get("/Rect") or []
            try:
                x0 = float(rect[0])
                top = max(float(rect[1]), float(rect[3]))
            except Exception:  # noqa: BLE001
                continue
            key = (page_num, -round(top), x0)
            current = positions.get(name)
            if current is None or key < current:
                positions[name] = key
    return positions


def _annotation_field_name(annot: Any) -> str | None:
    name = annot.get("/T")
    if name:
        return str(name)
    parent = annot.get("/Parent")
    if parent is None:
        return None
    try:
        parent_obj = parent.get_object()
    except Exception:  # noqa: BLE001
        return None
    parent_name = parent_obj.get("/T")
    return str(parent_name) if parent_name else None


def _try_render_acroform_facsimile(
    doc_bytes: bytes,
    *,
    source_hint: str,
) -> dict[str, Any] | None:
    """Render fillable PDF forms without guessing layout from whitespace.

    AcroForm PDFs expose authoritative field names and control types. The
    general text parser is intentionally conservative for these forms because
    PDF layout whitespace often makes ordinary labels look like tables,
    figures, or duplicate checkbox groups.
    """
    try:
        import pypdf

        reader = pypdf.PdfReader(io.BytesIO(doc_bytes))
    except Exception:  # noqa: BLE001
        return None

    fields = _extract_acroform_fields(reader)
    if not fields:
        return None

    meta = reader.metadata or {}
    metadata_title = getattr(meta, "title", None) if not isinstance(meta, dict) else meta.get("/Title")
    page_texts: list[tuple[int, str]] = []
    for page_num, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
        except TypeError:
            text = page.extract_text() or ""
        page_texts.append((page_num, text))

    title = _choose_document_title(
        page_texts,
        source_hint=source_hint,
        metadata_title=_clean_metadata_title(metadata_title),
    )
    title = _repair_form_title(title, source_hint=source_hint, page_texts=page_texts)

    model = _build_acroform_document_model(
        page_texts=page_texts,
        fields=fields,
        title=title,
        source_hint=source_hint,
        page_count=max(1, len(reader.pages)),
    )
    html, outline = _render_form_document_html(model)
    result = remediate_mod.RemediationResult(
        html=html,
        outline=outline,
        title=model.title,
        description=model.description,
        language=model.language,
        page_count=max(1, len(reader.pages)),
        render_mode="static",
    )
    return {
        "result": result,
        "parsed_document": {
            "source_format": "pdf",
            "title": model.title,
            "language": model.language,
            "page_count": result.page_count,
            "doc_kind": "form_doc",
            "renderer": "acroform-facsimile",
            "has_fillable_form": True,
        },
        "layout_plan": {
            "template": "form",
            "visual_density": "dense",
            "toc_policy": "hidden",
            "renderer": "acroform-facsimile",
        },
    }


def _try_render_docx_form_facsimile(
    doc_bytes: bytes,
    *,
    source_hint: str,
) -> dict[str, Any] | None:
    markdown = extract_docx_as_markdown(doc_bytes)
    lines = [line.strip() for line in markdown.splitlines() if line.strip()]
    if not any(_line_looks_like_docx_form_signal(line) for line in lines):
        return None

    signal_lines = _merge_docx_yes_no_lines(lines)
    emitted, _next_order = _signals_from_chunk(
        signal_lines,
        page_num=1,
        chunk_index=0,
        order=0,
        source="extract",
    )
    fields = _normalize_docx_form_fields(
        signal.field
        for signal in emitted
        if signal.kind == "form_field" and signal.field is not None
    )
    if len(fields) < 3:
        return None

    title = _form_title_from_source_hint(source_hint) or _fallback_title(source_hint)
    visual_lines = [_clean_pdf_visual_line(line) for line in lines]
    fieldsets = _group_acroform_fields(fields, page_texts=[(1, "\n".join(visual_lines))])
    tables = _form_tables_from_lines(visual_lines)
    if tables:
        for section in fieldsets:
            if section.legend == "Eligibility":
                section.tables = [*(section.tables or []), *tables]
                break
        else:
            fieldsets.append(FormSectionModel(legend="Tables", fields=[], tables=tables))

    trailing_sections: list[tuple[str, str]] = []
    instructions = _docx_instruction_text(visual_lines)
    if instructions:
        trailing_sections.append(("Instructions", instructions))

    model = FormDocumentModel(
        title=title,
        language=_infer_form_language(source_hint=source_hint, title=title, page_texts=[(1, markdown)]),
        page_count=1,
        kicker=_first_matching_line(visual_lines[:12], ("community college district", "college district", "los angeles")),
        subtitle=_source_form_note(visual_lines[:16]),
        note=None,
        description=_form_description(title, fieldsets),
        sections=fieldsets,
        trailing_sections=trailing_sections,
        instructions_first=False,
    )
    html, outline = _render_form_document_html(model)
    result = remediate_mod.RemediationResult(
        html=html,
        outline=outline,
        title=model.title,
        description=model.description,
        language=model.language,
        page_count=model.page_count,
        render_mode="static",
    )
    return {
        "result": result,
        "parsed_document": {
            "source_format": "docx",
            "title": model.title,
            "language": model.language,
            "page_count": model.page_count,
            "doc_kind": "form_doc",
            "renderer": "docx-form-facsimile",
            "has_fillable_form": True,
        },
        "layout_plan": {
            "template": "form",
            "visual_density": "dense",
            "toc_policy": "hidden",
            "renderer": "docx-form-facsimile",
        },
    }


def _line_looks_like_docx_form_signal(line: str) -> bool:
    stripped = line.strip()
    return (
        "___" in stripped
        or bool(re.search(r"\bYES\s+NO\b", stripped, flags=re.IGNORECASE))
        or stripped.lower().startswith(("section i", "section ii", "section iii"))
    )


def _merge_docx_yes_no_lines(lines: Sequence[str]) -> list[str]:
    merged: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if (
            index + 1 < len(lines)
            and re.fullmatch(r"yes\s+no", lines[index + 1].strip(), flags=re.IGNORECASE)
            and line.strip()
            and not _looks_like_form_section_start(line)
        ):
            merged.append(f"{line} YES NO")
            index += 2
            continue
        merged.append(line)
        index += 1
    return merged


def _normalize_docx_form_fields(fields: Sequence[FormField]) -> list[FormField]:
    normalized: list[FormField] = []
    for field in fields:
        label = _clean_field_label(field.label)
        if not label:
            continue
        if label == "First Name Middle Initial":
            normalized.extend(
                [
                    field.model_copy(update={"id": "first-name", "label": "First Name"}),
                    field.model_copy(update={"id": "middle-initial", "label": "Middle Initial"}),
                ]
            )
            continue
        if label == "Applicant’s Signature: Date":
            normalized.extend(
                [
                    field.model_copy(
                        update={"id": "applicants-signature", "label": "Applicant’s Signature", "field_type": "signature"}
                    ),
                    field.model_copy(update={"id": "date", "label": "Date", "field_type": "date"}),
                ]
            )
            continue
        if label.startswith("Parent Signature") and field.options == ["Date"]:
            normalized.extend(
                [
                    field.model_copy(update={"id": "parent-signature", "label": label, "field_type": "signature", "options": []}),
                    field.model_copy(update={"id": "date", "label": "Date", "field_type": "date", "options": []}),
                ]
            )
            continue
        if len(label) > 180 and field.field_type == "checkbox":
            continue
        label = re.sub(r"\s+l$", "", label)
        normalized.append(field.model_copy(update={"id": _slugify(label), "label": label}))
    return normalized


def _docx_instruction_text(lines: Sequence[str]) -> str | None:
    for line in lines:
        if line.lower().startswith("laccd board policy"):
            return line
    intro_lines = [
        line
        for line in lines[:10]
        if len(line) > 80 and not _line_looks_like_docx_form_signal(line)
    ]
    return " ".join(intro_lines) if intro_lines else None


def _build_acroform_document_model(
    *,
    page_texts: Sequence[tuple[int, str]],
    fields: Sequence[FormField],
    title: str,
    source_hint: str,
    page_count: int,
) -> FormDocumentModel:
    visual_lines = [
        _clean_pdf_visual_line(line)
        for _, text in page_texts
        for line in text.splitlines()
        if line.strip()
    ]
    language = _infer_form_language(source_hint=source_hint, title=title, page_texts=page_texts)
    fieldsets = _group_acroform_fields(fields, page_texts=page_texts)
    if "nonresident tuition exemption" in title.lower():
        _merge_application_identity_section(fieldsets)
    tables = _form_tables_from_lines(visual_lines)
    if tables:
        for section in fieldsets:
            if section.legend == "Eligibility":
                section.tables = [*(section.tables or []), *tables]
                break
        else:
            fieldsets.append(FormSectionModel(legend="Tables", fields=[], tables=tables))

    trailing_sections: list[tuple[str, str]] = []
    instructions = _extract_instructions_text(_instruction_lines(visual_lines))
    if instructions:
        trailing_sections.append(("Instructions", instructions))
    instructions_first = _instructions_precede_first_form_section(visual_lines)

    kicker = _first_matching_line(
        visual_lines[:8],
        ("community college district", "college district", "los angeles"),
    )
    source_note = _source_form_note(visual_lines[:16])
    description = _form_description(title, fieldsets)
    return FormDocumentModel(
        title=title,
        language=language,
        page_count=page_count,
        kicker=kicker,
        subtitle=source_note,
        note=None,
        description=description,
        sections=fieldsets,
        trailing_sections=trailing_sections,
        instructions_first=instructions_first,
    )


def _group_acroform_fields(
    fields: Sequence[FormField],
    *,
    page_texts: Sequence[tuple[int, str]],
) -> list[FormSectionModel]:
    buckets: dict[str, list[FormField]] = {}
    section_order: list[str] = []
    all_text = " ".join(text for _, text in page_texts)
    checkbox_label_candidates = _checkbox_label_candidates_from_text(all_text)
    repaired_fields = [
        _repair_acroform_field(
            _replace_generic_checkbox_label(field, checkbox_label_candidates),
            all_text=all_text,
        )
        for field in fields
    ]
    sequence_tables, consumed_sequence_ids = _sequence_tables_from_repeated_fields(repaired_fields)
    for field in repaired_fields:
        if field.id in consumed_sequence_ids:
            continue
        section = _section_for_field(field)
        if section not in buckets:
            buckets[section] = []
            section_order.append(section)
        buckets[section].append(field)
    for section in sequence_tables:
        if section.legend not in buckets:
            buckets[section.legend] = []
            section_order.append(section.legend)

    preferred_order = [
        "Application",
        "Term information",
        "Student information",
        "Contact information",
        "Nonimmigrant visa status",
        "Eligibility",
        "Attendance and graduation requirements",
        "School attendance and credit/hour information",
        "Course information",
        "Applicant certification",
        "For office use only",
        "Form fields",
        "Form tables",
    ]
    order_rank = {legend: index for index, legend in enumerate(preferred_order)}
    seen_rank = {legend: index for index, legend in enumerate(section_order)}
    ordered_sections = sorted(
        section_order,
        key=lambda legend: (order_rank.get(legend, len(order_rank)), seen_rank[legend]),
    )
    sections = [
        FormSectionModel(
            legend=legend,
            fields=buckets[legend],
            tables=[
                table
                for sequence_section in sequence_tables
                if sequence_section.legend == legend
                for table in sequence_section.tables or []
            ]
            or None,
        )
        for legend in ordered_sections
        if buckets[legend] or any(sequence_section.legend == legend for sequence_section in sequence_tables)
    ]
    _move_repeated_row_fields_to_tables(sections)
    _move_office_workflow_fields_to_table(sections)
    _move_applicant_identity_and_signature_fields_to_tables(sections)
    _move_student_course_summary_fields_to_table(sections)
    _move_evidence_result_fields_to_table(sections)
    sections[:] = [
        section
        for section in sections
        if section.fields or section.tables
    ]
    if len(sections) == 1 and sections[0].legend == "Form fields":
        sections[0].legend = "Document form"
    return sections


def _repair_acroform_field(field: FormField, *, all_text: str) -> FormField:
    label = _humanize_field_label(field.label)
    field_type = field.field_type
    options = list(field.options)
    if label.lower() == "specify the college" and "i, the undersigned" in all_text.lower():
        label = (
            "I, the undersigned, am applying for the California Nonresident Tuition "
            "Exemption at (specify the college)"
        )
    line = _line_for_field_label(label, all_text)
    line_options = _options_from_text(line or "")
    if field_type == "checkbox" and not options and _line_has_yes_no_choice(label, line):
        field_type = "radio"
        options = ["Yes", "No"]
    elif field_type == "checkbox" and line_options and len(line_options) > 1:
        normalized_label = label.strip().lower()
        normalized_options = {option.strip().lower().rstrip("*") for option in line_options}
        if normalized_label not in normalized_options:
            field_type = _choice_type_from_label_options(label, line_options)
            options = line_options
    return field.model_copy(
        update={
            "label": label,
            "field_type": field_type,
            "options": options,
        }
    )


def _replace_generic_checkbox_label(
    field: FormField,
    candidates: list[str],
) -> FormField:
    if field.field_type != "checkbox" or not _is_generic_checkbox_label(field.label):
        return field
    while candidates:
        label = _clean_field_label(candidates.pop(0))
        if label and not _is_generic_checkbox_label(label):
            return field.model_copy(update={"id": _slugify(label), "label": label})
    fallback = _fallback_checkbox_label(field.label)
    return field.model_copy(update={"id": _slugify(fallback), "label": fallback})


def _is_generic_checkbox_label(label: str) -> bool:
    return bool(re.match(r"^check\s+box[\w\\-]*$", _clean_field_label(label), flags=re.IGNORECASE))


def _fallback_checkbox_label(label: str) -> str:
    suffix = re.sub(r"^check\s+box", "", _clean_field_label(label), flags=re.IGNORECASE).strip()
    suffix = suffix or "option"
    return f"Additional checkbox option {suffix}"


def _checkbox_label_candidates_from_text(text: str) -> list[str]:
    lines = [
        _clean_pdf_visual_line(line)
        for line in text.splitlines()
        if _clean_pdf_visual_line(line)
    ]
    candidates: list[str] = []
    in_check_one_block = False
    for raw_line in lines:
        line = re.sub(r"\s+", " ", raw_line).strip()
        if not line:
            continue
        lowered = line.lower()
        if "check one" in lowered or "check any" in lowered:
            in_check_one_block = True

        yes_no_match = re.search(r"\bYES\s+NO\b", line, flags=re.IGNORECASE)
        if yes_no_match:
            label = f"{line[:yes_no_match.start()]} {line[yes_no_match.end():]}".strip(" :-")
            label = _strip_form_item_marker(label)
            label = re.sub(r"\s+", " ", label).strip()
            if label:
                candidates.extend([f"{label} - Yes", f"{label} - No"])
            continue

        option = _checkbox_option_label_from_visual_line(line)
        if option:
            candidates.append(option)
            continue

        if in_check_one_block and _looks_like_standalone_checkbox_candidate(line):
            candidates.append(_strip_form_item_marker(line))
            continue
        if _looks_like_standalone_checkbox_candidate(line) and line.lower().startswith(
            ("a director", "a mckinney", "an unaccompanied")
        ):
            candidates.append(_strip_form_item_marker(line))
            continue

        numbered_option = _numbered_checkbox_option_from_visual_line(line)
        if numbered_option:
            candidates.append(numbered_option)
            continue

        if _looks_like_documentation_option(line):
            candidates.append(line)
            continue

        if _looks_like_non_option_prose(line):
            in_check_one_block = False
    return _dedupe(candidates)


def _checkbox_option_label_from_visual_line(line: str) -> str | None:
    match = re.match(r"^(?:[a-z]\.|[-*•])\s*(?P<label>.+)$", line.strip(), flags=re.IGNORECASE)
    if not match:
        return None
    label = match.group("label").strip()
    if not label or _looks_like_non_option_prose(label):
        return None
    return _strip_form_item_marker(label)


def _numbered_checkbox_option_from_visual_line(line: str) -> str | None:
    match = re.match(r"^(?P<number>\d{1,2})\.\s*(?P<label>.+)$", line.strip())
    if not match:
        return None
    number = int(match.group("number"))
    label = match.group("label").strip()
    if number < 19 or number > 34:
        return None
    if not label.lower().startswith("i "):
        return None
    return _strip_form_item_marker(label)


def _looks_like_documentation_option(line: str) -> bool:
    lowered = line.lower().strip()
    return lowered.startswith(
        (
            "federal or california",
            "documentation of purchase",
            "california automobile",
            "california voter",
            "utility bills",
            "california driver",
            "california vehicle",
            "active california",
            "selective service",
            "receipt of benefits",
            "california occupational",
            "other documentation",
        )
    )


def _strip_form_item_marker(label: str) -> str:
    return re.sub(r"^\s*(?:\d+[\.\)]\s*|[a-z]\.\s*)", "", label, flags=re.IGNORECASE).strip()


def _looks_like_standalone_checkbox_candidate(line: str) -> bool:
    stripped = _strip_form_item_marker(line)
    if not stripped or len(stripped) > 160:
        return False
    lowered = stripped.lower()
    if lowered.startswith(("if ", "provide ", "this ", "as per ", "page ")):
        return False
    return bool(re.search(r"\b(?:director|liaison|unaccompanied|resident|applicant|beneficiary|status|program|shelter)\b", lowered))


def _looks_like_non_option_prose(line: str) -> bool:
    stripped = line.strip()
    lowered = stripped.lower()
    return (
        len(stripped) > 180
        or stripped.endswith(".")
        and not re.match(r"^(?:[a-z]\.|[-*•])\s+", stripped, flags=re.IGNORECASE)
        and not lowered.startswith(("a ", "an "))
    )


def _section_for_field(field: FormField) -> str:
    label = field.label.lower()
    field_id = field.id.lower()
    haystack = f"{label} {field_id}"
    if "check one box only" in haystack or "nonimmigrant visa" in haystack:
        return "Nonimmigrant visa status"
    if any(token in haystack for token in ("attendance at a california", "attendance and graduation", "attended or attained", "high school coursework", "elementary schools", "associate", "high school diploma", "minimum requirements", "community college for transfer")):
        return "Attendance and graduation requirements"
    repeated_parts = _repeated_row_field_parts(field.label)
    if repeated_parts:
        repeated_header = repeated_parts[0].lower()
        if any(token in haystack for token in ("school", "credits", "hours", "month/year")) or repeated_header == "city":
            return "School attendance and credit/hour information"
        if any(token in haystack for token in ("class number", "course", "semester", "year", "college")):
            return "Course information"
        return "Form tables"
    if any(token in haystack for token in ("office", "approved", "denied", "noted", "emailed", "received", "rec'd", "recd", "evaluator", "counselor", "date evidence", "date section", "credit entered", "processed by")) or label in {"by"}:
        return "For office use only"
    if "middle initial" in haystack:
        return "Student information"
    if any(token in haystack for token in ("signature", "student initials", "initials", "certif", "parent signature")):
        return "Applicant certification"
    if any(token in haystack for token in ("course", "class number", "subject", "section", "semester year", "units", "grade", "grading method")):
        return "Course information"
    if (
        any(token in haystack for token in ("address", "city", "zip", "email", "e-mail", "phone", "telephone", "contact"))
        or re.search(r"\bstate\b", haystack)
    ):
        return "Contact information"
    if any(token in haystack for token in ("term", "semester", "fall", "winter", "spring", "summer", "year")):
        return "Term information"
    if any(token in haystack for token in ("immigration", "visa", "income", "race", "ethnic", "hispanic", "latino", "gender", "sex", "language", "goal", "ferpa", "homeless", "military", "veteran", "citizen", "resident")):
        return "Eligibility"
    if "specify the college" in haystack or "college student id" in haystack:
        return "Application"
    if any(token in haystack for token in ("name", "student", "birth", "dob", "age", "social security", "ssn", "campus", "college")):
        return "Student information"
    if re.fullmatch(r"date\s*\d*", label):
        return "Applicant certification"
    return "Form fields"


def _move_repeated_row_fields_to_tables(sections: list[FormSectionModel]) -> None:
    grouped: dict[str, dict[int, list[tuple[int, int, str, FormField]]]] = {}
    remove_by_section: dict[int, set[int]] = {}
    first_section_for_group: dict[str, int] = {}

    for section_index, section in enumerate(sections):
        for field_index, field in enumerate(section.fields):
            parts = _repeated_row_field_parts(field.label)
            if parts is None:
                continue
            header, row_num = parts
            table_key = _table_key_for_repeated_field(header, section.legend)
            grouped.setdefault(table_key, {}).setdefault(row_num, []).append(
                (section_index, field_index, header, field)
            )
            first_section_for_group.setdefault(table_key, section_index)

    tables_by_section: dict[int, list[FormTableModel]] = {}
    for table_key, rows_by_num in grouped.items():
        if len(rows_by_num) < 2:
            continue
        headers = _sort_repeated_table_headers(table_key, _headers_for_repeated_rows(rows_by_num))
        if len(headers) < 2:
            continue
        table_rows: list[list[FormField]] = []
        for row_num in sorted(rows_by_num):
            by_header = {header: (section_index, field_index, field) for section_index, field_index, header, field in rows_by_num[row_num]}
            if len(by_header) < max(2, len(headers) // 2):
                continue
            row_fields: list[FormField] = []
            for header in headers:
                entry = by_header.get(header)
                if entry is None:
                    field = FormField(
                        id=_slugify(f"{header} row {row_num}"),
                        label=f"{header}, row {row_num}",
                        field_type="text",
                    )
                else:
                    section_index, field_index, field = entry
                    remove_by_section.setdefault(section_index, set()).add(field_index)
                row_fields.append(field.model_copy(update={"label": f"{header}, row {row_num}"}))
            if row_fields:
                table_rows.append(row_fields)
        if table_rows:
            section_index = first_section_for_group[table_key]
            tables_by_section.setdefault(section_index, []).append(
                FormTableModel(
                    caption=_caption_for_repeated_table(table_key, sections[section_index].legend),
                    headers=headers,
                    field_rows=table_rows,
                )
            )

    if tables_by_section:
        for section_index, remove_indexes in remove_by_section.items():
            section = sections[section_index]
            section.fields = [
                field for index, field in enumerate(section.fields) if index not in remove_indexes
            ]
        for section_index, tables in tables_by_section.items():
            section = sections[section_index]
            section.tables = [*(section.tables or []), *tables]

    # Drop table-only staging sections after their fields have been moved.
    sections[:] = [
        section
        for section in sections
        if section.legend != "Form tables" or section.fields or section.tables
    ]


def _sequence_tables_from_repeated_fields(
    fields: Sequence[FormField],
) -> tuple[list[FormSectionModel], set[str]]:
    candidates: list[tuple[str, FormField]] = []
    for field in fields:
        header = _sequence_table_header(field.label)
        if header:
            candidates.append((header, field))
    if not candidates:
        return [], set()

    counts = Counter(header for header, _field in candidates)
    repeated_headers = {header for header, count in counts.items() if count >= 2}
    if len(repeated_headers) < 2:
        return [], set()

    header_order: list[str] = []
    seen_headers: set[str] = set()
    for header, _field in candidates:
        if header in repeated_headers and header not in seen_headers:
            header_order.append(header)
            seen_headers.add(header)

    if not _looks_like_repeated_course_sequence(header_order):
        return [], set()
    header_order = sorted(
        header_order,
        key=lambda header: (
            {"subject and number": 0, "title": 1, "term and year": 2}.get(header.lower(), 99),
            header.lower(),
        ),
    )

    occurrences: Counter[str] = Counter()
    rows_by_num: dict[int, dict[str, FormField]] = {}
    for header, field in candidates:
        if header not in repeated_headers:
            continue
        occurrences[header] += 1
        rows_by_num.setdefault(occurrences[header], {})[header] = field

    table_rows: list[list[FormField]] = []
    consumed_ids: set[str] = set()
    for row_num in sorted(rows_by_num):
        by_header = rows_by_num[row_num]
        if len(by_header) < max(2, len(header_order) // 2):
            continue
        row_fields: list[FormField] = []
        for header in header_order:
            field = by_header.get(header)
            if field is None:
                field = FormField(
                    id=_slugify(f"{header} row {row_num}"),
                    label=f"{header}, row {row_num}",
                    field_type="text",
                )
            else:
                consumed_ids.add(field.id)
            row_fields.append(field.model_copy(update={"label": f"{header}, row {row_num}"}))
        table_rows.append(row_fields)

    if len(table_rows) < 2:
        return [], set()
    return [
        FormSectionModel(
            legend="Course information",
            fields=[],
            tables=[
                FormTableModel(
                    caption="Course information",
                    headers=header_order,
                    field_rows=table_rows,
                )
            ],
        )
    ], consumed_ids


def _move_office_workflow_fields_to_table(sections: list[FormSectionModel]) -> None:
    date_fields: dict[str, FormField] = {}
    processed_fields: dict[str, FormField] = {}
    remove_ids: set[str] = set()

    for section in sections:
        for field in section.fields:
            date_step = _office_workflow_date_step(field.label)
            if date_step:
                date_fields.setdefault(date_step, field)
            processed_step = _office_workflow_processed_step(field.label)
            if processed_step:
                processed_fields.setdefault(processed_step, field)

    preferred_steps = [
        "Received",
        "Evidence Requested",
        "Evidence Received",
        "Evidence Evaluated",
        "Result Accepted",
        "Section Created",
        "Credit Entered",
    ]
    all_steps = [
        step
        for step in preferred_steps
        if step in date_fields and step in processed_fields
    ]
    if len(all_steps) < 3:
        return

    field_rows: list[list[FormField]] = []
    for step in all_steps:
        date_field = date_fields[step].model_copy(update={"label": f"{step} date"})
        processed_field = processed_fields[step].model_copy(
            update={"label": f"{step} processed by name and role"}
        )
        field_rows.append([date_field, processed_field])
        remove_ids.add(date_fields[step].id)
        remove_ids.add(processed_fields[step].id)

    for section in sections:
        section.fields = [field for field in section.fields if field.id not in remove_ids]

    office_section = next((section for section in sections if section.legend == "For office use only"), None)
    if office_section is None:
        office_section = FormSectionModel(legend="For office use only", fields=[])
        sections.append(office_section)
    office_section.tables = [
        *(office_section.tables or []),
        FormTableModel(
            caption="Office processing workflow",
            headers=["Date", "Processed by (Name/Role)"],
            field_rows=field_rows,
            row_headers=all_steps,
        ),
    ]


def _move_applicant_identity_and_signature_fields_to_tables(sections: list[FormSectionModel]) -> None:
    identity_specs = [
        ("Full name", lambda field: _normalized_label(field.label) == "full name"),
        ("Campus ID number", lambda field: _normalized_label(field.label) == "campus id number"),
        ("Email address", lambda field: _normalized_label(field.label) == "email address"),
    ]
    identity_locations = [
        (header, _find_first_field_location(sections, predicate))
        for header, predicate in identity_specs
    ]
    if not all(location for _header, location in identity_locations):
        return
    identity_fields = _pop_located_fields(sections, identity_locations)

    certification = _ensure_section(sections, "Applicant certification")
    certification.tables = [
        *(certification.tables or []),
        FormTableModel(
            caption="Declaration applicant information",
            headers=[header for header, _field in identity_fields],
            field_rows=[
                [
                    field.model_copy(update={"label": header})
                    for header, field in identity_fields
                    if field is not None
                ]
            ],
        ),
    ]

    signature = _pop_first_field(
        [certification],
        lambda field: "signature" in _normalized_label(field.label),
    )
    date = _pop_first_field(
        [certification],
        lambda field: _normalized_label(field.label) == "date",
    )
    if signature is not None and date is not None:
        certification.tables.append(
            FormTableModel(
                caption="Signature and date",
                headers=["Signature", "Date"],
                field_rows=[
                    [
                        signature.model_copy(update={"label": "Signature"}),
                        date.model_copy(update={"label": "Date"}),
                    ]
                ],
            )
        )
    else:
        if signature is not None:
            certification.fields.append(signature)
        if date is not None:
            certification.fields.append(date)


def _move_student_course_summary_fields_to_table(sections: list[FormSectionModel]) -> None:
    specs = [
        ("Last name", lambda field: _normalized_label(field.label) == "last name"),
        ("First name", lambda field: _normalized_label(field.label) == "first name"),
        ("LACCD student ID#", lambda field: "laccd student id" in _normalized_label(field.label)),
        ("Course", lambda field: "subject number and title" in _normalized_label(field.label)),
        ("Grading method preference", lambda field: _normalized_label(field.label) == "grading method"),
    ]
    locations = [(header, _find_first_field_location(sections, predicate)) for header, predicate in specs]
    if sum(1 for _header, location in locations if location is not None) < 4:
        return
    picked = _pop_located_fields(sections, locations)
    course_section = _ensure_section(sections, "Course information")
    course_section.tables = [
        FormTableModel(
            caption="Student and course information",
            headers=[header for header, field in picked if field is not None],
            field_rows=[
                [
                    field.model_copy(update={"label": header})
                    for header, field in picked
                    if field is not None
                ]
            ],
        ),
        *(course_section.tables or []),
    ]


def _move_evidence_result_fields_to_table(sections: list[FormSectionModel]) -> None:
    specs = [
        ("Evidence type", lambda field: _normalized_label(field.label) == "select evidence here"),
        ("Evidence requested detail", lambda field: "detail evidence" in _normalized_label(field.label)),
        ("Evidence result", lambda field: _normalized_label(field.label) == "evidence result"),
        ("Denial reason", lambda field: "denial reason" in _normalized_label(field.label)),
        ("Result acceptance", lambda field: _normalized_label(field.label) == "result acceptance"),
        ("Date acknowledged", lambda field: _normalized_label(field.label) == "date acknowledged"),
        ("Notes", lambda field: _normalized_label(field.label) == "notes"),
        (
            "Evidence type specifications",
            lambda field: "evidence type specifications" in _normalized_label(field.label),
        ),
    ]
    locations = [(label, _find_first_field_location(sections, predicate)) for label, predicate in specs]
    if sum(1 for _label, location in locations if location is not None) < 3:
        return
    picked = _pop_located_fields(sections, locations)

    detail_section = _ensure_section(sections, "Form fields")
    detail_section.tables = [
        FormTableModel(
            caption="Evidence and result details",
            headers=["Details"],
            field_rows=[
                [field.model_copy(update={"label": label})]
                for label, field in picked
                if field is not None
            ],
            row_headers=[label for label, field in picked if field is not None],
        ),
        *(detail_section.tables or []),
    ]


def _normalized_label(label: str) -> str:
    return re.sub(r"[^a-z0-9#]+", " ", label.lower()).strip()


def _ensure_section(sections: list[FormSectionModel], legend: str) -> FormSectionModel:
    section = next((candidate for candidate in sections if candidate.legend == legend), None)
    if section is None:
        section = FormSectionModel(legend=legend, fields=[])
        sections.append(section)
    return section


def _append_field_to_section(sections: list[FormSectionModel], legend: str, field: FormField) -> None:
    _ensure_section(sections, legend).fields.append(field)


def _pop_first_field(
    sections: Sequence[FormSectionModel],
    predicate: Any,
) -> FormField | None:
    for section in sections:
        for index, field in enumerate(section.fields):
            if predicate(field):
                return section.fields.pop(index)
    return None


def _find_first_field_location(
    sections: Sequence[FormSectionModel],
    predicate: Any,
) -> tuple[int, int] | None:
    for section_index, section in enumerate(sections):
        for field_index, field in enumerate(section.fields):
            if predicate(field):
                return section_index, field_index
    return None


def _pop_field_at(
    sections: Sequence[FormSectionModel],
    location: tuple[int, int],
) -> FormField:
    section_index, field_index = location
    return sections[section_index].fields.pop(field_index)


def _pop_located_fields(
    sections: Sequence[FormSectionModel],
    located: Sequence[tuple[str, tuple[int, int] | None]],
) -> list[tuple[str, FormField]]:
    popped: dict[str, FormField] = {}
    for label, location in sorted(
        [(label, location) for label, location in located if location is not None],
        key=lambda item: item[1] or (-1, -1),
        reverse=True,
    ):
        if location is not None:
            popped[label] = _pop_field_at(sections, location)
    return [
        (label, popped[label])
        for label, location in located
        if location is not None and label in popped
    ]


def _office_workflow_date_step(label: str) -> str | None:
    match = re.match(r"^date\s+(?P<step>.+)$", _clean_field_label(label), flags=re.IGNORECASE)
    if not match:
        return None
    return _canonical_office_workflow_step(match.group("step"))


def _office_workflow_processed_step(label: str) -> str | None:
    match = re.search(
        r"processed\s+by\s*(?:\([^)]*\))?\s*[:_-]\s*(?P<step>.+)$",
        _clean_field_label(label),
        flags=re.IGNORECASE,
    )
    if not match:
        return None
    return _canonical_office_workflow_step(match.group("step"))


def _canonical_office_workflow_step(step: str) -> str | None:
    cleaned = re.sub(r"\s+", " ", step.replace("_", " ")).strip(" :-")
    lowered = cleaned.lower()
    mapping = {
        "received": "Received",
        "evidence requested": "Evidence Requested",
        "evidence received": "Evidence Received",
        "evidence evaluated": "Evidence Evaluated",
        "result accepted": "Result Accepted",
        "section created": "Section Created",
        "credit entered": "Credit Entered",
    }
    return mapping.get(lowered)


def _sequence_table_header(label: str) -> str | None:
    cleaned = _humanize_field_label(label)
    cleaned = re.sub(r",?\s*example:.*$", "", cleaned, flags=re.IGNORECASE).strip(" :-")
    lowered = cleaned.lower()
    if lowered in {"subject and number", "title", "term and year"}:
        return cleaned
    return None


def _looks_like_repeated_course_sequence(headers: Sequence[str]) -> bool:
    lowered = {header.lower() for header in headers}
    return (
        "subject and number" in lowered
        and "title" in lowered
        and "term and year" in lowered
    )


def _move_repeated_row_fields_to_tables_old(sections: list[FormSectionModel]) -> None:
    for section in sections:
        grouped: dict[str, dict[int, list[tuple[int, str, FormField]]]] = {}
        remove_indexes: set[int] = set()
        for index, field in enumerate(section.fields):
            parts = _repeated_row_field_parts(field.label)
            if parts is None:
                continue
            header, row_num = parts
            table_key = _table_key_for_repeated_field(header, section.legend)
            grouped.setdefault(table_key, {}).setdefault(row_num, []).append((index, header, field))

        tables: list[FormTableModel] = []
        for table_key, rows_by_num in grouped.items():
            if len(rows_by_num) < 2:
                continue
            headers = _headers_for_repeated_rows(rows_by_num)
            if len(headers) < 2:
                continue
            table_rows: list[list[FormField]] = []
            for row_num in sorted(rows_by_num):
                by_header = {header: field for index, header, field in rows_by_num[row_num]}
                if len(by_header) < max(2, len(headers) // 2):
                    continue
                row_fields: list[FormField] = []
                for header in headers:
                    field = by_header.get(header)
                    if field is None:
                        field = FormField(
                            id=_slugify(f"{header} row {row_num}"),
                            label=f"{header}, row {row_num}",
                            field_type="text",
                        )
                    else:
                        source_index = next(
                            index
                            for index, candidate_header, candidate in rows_by_num[row_num]
                            if candidate_header == header and candidate is field
                        )
                        remove_indexes.add(source_index)
                    row_fields.append(field.model_copy(update={"label": f"{header}, row {row_num}"}))
                if row_fields:
                    table_rows.append(row_fields)
            if table_rows:
                tables.append(
                    FormTableModel(
                        caption=_caption_for_repeated_table(table_key, section.legend),
                        headers=headers,
                        field_rows=table_rows,
                    )
                )

        if tables:
            section.fields = [
                field for index, field in enumerate(section.fields) if index not in remove_indexes
            ]
            section.tables = [*(section.tables or []), *tables]


def _repeated_row_field_parts(label: str) -> tuple[str, int] | None:
    normalized = _humanize_field_label(label)
    match = re.search(r"^(?P<header>.+?)(?:\s*[:_-]?\s*)(?:row)\s*(?P<row>\d+)$", normalized, flags=re.IGNORECASE)
    if match:
        header = _clean_field_label(match.group("header"))
        if not header:
            return None
        return header, int(match.group("row"))
    terminal = _terminal_number_field_parts(normalized)
    if terminal is None:
        return None
    header, row_num = terminal
    if header.lower() in {"student initials", "signature", "date", "text entry"}:
        return None
    if not header:
        return None
    return header, row_num


def _terminal_number_field_parts(label: str) -> tuple[str, int] | None:
    match = re.search(r"^(?P<header>.+?)\s+(?P<row>\d+)$", _clean_field_label(label))
    if not match:
        return None
    header = _clean_field_label(match.group("header"))
    if not header or len(header) < 2:
        return None
    return header, int(match.group("row"))


def _table_key_for_repeated_field(header: str, section_legend: str) -> str:
    section = section_legend.lower()
    if "school attendance" in section:
        return "school-attendance"
    if "course" in section:
        return "course-rows"
    lowered = header.lower()
    if lowered in {"from", "to", "state country", "state/country", "statecountry"}:
        return "residence-history"
    if any(token in lowered for token in ("school", "credits", "hours", "month/year")):
        return "school-attendance"
    if any(token in lowered for token in ("class number", "course", "semester", "year", "college")):
        return "course-rows"
    return _slugify(section_legend or "form table")


def _headers_for_repeated_rows(rows_by_num: dict[int, list[tuple[Any, ...]]]) -> list[str]:
    header_order: list[tuple[int, str]] = []
    seen: set[str] = set()
    for row_num in sorted(rows_by_num):
        for entry in sorted(rows_by_num[row_num], key=lambda item: item[-3]):
            index = int(entry[-3])
            header = str(entry[-2])
            key = _slugify(header)
            if key in seen:
                continue
            seen.add(key)
            header_order.append((index, header))
    return [header for _index, header in sorted(header_order, key=lambda item: item[0])]


def _sort_repeated_table_headers(table_key: str, headers: list[str]) -> list[str]:
    if table_key == "school-attendance":
        preferred = [
            ("name", "school"),
            ("type", "school"),
            ("city",),
            ("state", "country"),
            ("from",),
            ("to",),
            ("number", "credits"),
            ("number", "hours"),
        ]
    elif table_key == "course-rows":
        preferred = [
            ("course",),
            ("class", "number"),
            ("semester",),
            ("year",),
            ("college",),
            ("units",),
            ("grade",),
        ]
    else:
        return headers

    def rank(header: str) -> tuple[int, str]:
        lowered = header.lower()
        for index, tokens in enumerate(preferred):
            if all(token in lowered for token in tokens):
                return index, lowered
        return len(preferred), lowered

    return sorted(headers, key=rank)


def _caption_for_repeated_table(table_key: str, fallback: str) -> str:
    if table_key == "school-attendance":
        return "School attendance and credit/hour information"
    if table_key == "residence-history":
        return "Residence history"
    if table_key == "course-rows":
        return "Course information"
    return fallback


def _render_form_document_html(model: FormDocumentModel) -> tuple[str, list[dict[str, Any]]]:
    used_ids: set[str] = set()
    outline = [{"id": "document-title", "level": 1, "text": model.title}]
    form_sections: list[str] = []
    leading_sections: list[str] = []
    for section in model.sections:
        section_id = _unique_slug(section.legend, used_ids)
        heading_id = _unique_slug(f"{section_id}-heading", used_ids)
        legend_id = _unique_slug(f"{section_id}-legend", used_ids)
        outline.append({"id": heading_id, "level": 2, "text": section.legend})
        fields = "".join(_render_facsimile_field(field, used_ids=used_ids) for field in section.fields)
        table_html = "".join(_render_form_table(table) for table in section.tables or [])
        intro_html = f'<p class="doc-form__intro">{escape(section.intro)}</p>' if section.intro else ""
        fieldset_html = (
            '<fieldset class="doc-form__fieldset">'
            f'<legend id="{legend_id}">{escape(section.legend)}</legend>'
            f'{intro_html}<div class="doc-form-grid doc-form-grid--2">{fields}</div>'
            '</fieldset>'
            if section.fields
            else intro_html
        )
        form_sections.append(
            '<section class="doc-section doc-section--single-column '
            f'doc-section--{section_id}" '
            f'id="{section_id}" aria-labelledby="{heading_id}">'
            '<header class="doc-section__header">'
            f'<h2 id="{heading_id}">{escape(section.legend)}</h2></header>'
            '<div class="doc-section__stack">'
            f'{fieldset_html}'
            f'{table_html}'
            '</div></section>'
        )

    trailing_html: list[str] = []
    for heading, body in model.trailing_sections:
        section_id = _unique_slug(heading, used_ids)
        heading_id = _unique_slug(f"{section_id}-heading", used_ids)
        outline.append({"id": heading_id, "level": 2, "text": heading})
        rendered = (
            '<section class="doc-section doc-section--single-column" '
            f'id="{section_id}" aria-labelledby="{heading_id}">'
            '<header class="doc-section__header">'
            f'<h2 id="{heading_id}">{escape(heading)}</h2></header>'
            f'<div class="doc-section__stack"><p>{escape(body)}</p></div></section>'
        )
        if heading.strip().lower() == "instructions" and model.instructions_first:
            leading_sections.append(rendered)
        else:
            trailing_html.append(rendered)

    kicker = model.kicker or ""
    kicker_html = f'<p class="doc-hero__kicker">{escape(kicker)}</p>' if kicker else ""
    subtitle_html = (
        f'<p class="doc-hero__subtitle">{escape(model.subtitle)}</p>'
        if model.subtitle
        else ""
    )
    html = (
        f'<main id="content" lang="{escape(model.language)}" '
        'class="doc-page doc-page--form doc-page--facsimile doc-page--density-dense">'
        '<a href="#content" class="sr-skip">Skip to main content</a>'
        '<header class="doc-hero doc-hero--compact">'
        f'{kicker_html}'
        f'<h1 id="document-title" class="doc-hero__title">{escape(model.title)}</h1>'
        f'{subtitle_html}'
        "</header>"
        '<div class="doc-page__body">'
        f'{"".join(leading_sections)}'
        f'<form class="doc-form doc-form--facsimile">{"".join(form_sections)}</form>'
        f'{"".join(trailing_html)}'
        "</div></main>"
    )
    return html, outline


def _source_form_note(lines: Sequence[str]) -> str | None:
    for line in lines:
        stripped = line.strip()
        if not stripped.lower().startswith("note:"):
            continue
        if "transcribed forms" in stripped.lower():
            continue
        return re.sub(r"\s+", " ", stripped)
    return None


def _merge_application_identity_section(sections: list[FormSectionModel]) -> None:
    application_index = next(
        (
            index
            for index, section in enumerate(sections)
            if section.legend == "Application"
            and any("college student id" in field.label.lower() for field in section.fields)
            and any("specify the college" in field.label.lower() for field in section.fields)
        ),
        None,
    )
    student_index = next(
        (
            index
            for index, section in enumerate(sections)
            if section.legend == "Student information"
            and section.fields
            and not section.tables
            and all("name" in field.label.lower() for field in section.fields)
        ),
        None,
    )
    if application_index is None or student_index is None:
        return
    application = sections[application_index]
    student = sections[student_index]
    application.fields = [*student.fields, *application.fields]
    del sections[student_index]


def _render_facsimile_field(field: FormField, *, used_ids: set[str]) -> str:
    field_id = _unique_slug(field.id or field.label, used_ids)
    label = escape(field.label)
    required = " required" if field.required else ""
    if field.field_type == "select" and field.options:
        options = "".join(
            f'<option value="{escape(option)}">{escape(option)}</option>'
            for option in field.options
        )
        return (
            '<div class="doc-form__field doc-form__field--compact">'
            f'<label for="{field_id}">{label}</label>'
            f'<select id="{field_id}"{required}>{options}</select></div>'
        )
    if field.field_type in {"radio", "checkbox"} and field.options:
        legend_id = _unique_slug(f"{field_id}-legend", used_ids)
        choices = []
        control_type = "radio" if field.field_type == "radio" else "checkbox"
        for option in field.options:
            option_id = _unique_slug(f"{field_id}-{option}", used_ids)
            choices.append(
                '<div class="doc-form__choice">'
                f'<input id="{option_id}" type="{control_type}" name="{field_id}" '
                f'value="{escape(option)}"{required} />'
                f'<label for="{option_id}">{escape(option)}</label></div>'
            )
        return (
            '<fieldset class="doc-form__group doc-form__group--inline" '
            f'aria-labelledby="{legend_id}">'
            f'<legend id="{legend_id}">{label}</legend>'
            f'<div class="doc-form__choices">{"".join(choices)}</div></fieldset>'
        )
    if field.field_type == "checkbox":
        return (
            '<div class="doc-form__field doc-form__field--compact">'
            '<div class="doc-form__choice">'
            f'<input id="{field_id}" type="checkbox"{required} />'
            f'<label for="{field_id}">{label}</label></div></div>'
        )
    if field.field_type in {"textarea", "signature"}:
        return (
            '<div class="doc-form__field doc-form__field--compact">'
            f'<label for="{field_id}">{label}</label>'
            f'<textarea id="{field_id}" rows="2"{required}></textarea></div>'
        )
    input_type = field.field_type if field.field_type in {"email", "tel", "url", "number", "date"} else "text"
    return (
        '<div class="doc-form__field doc-form__field--compact">'
        f'<label for="{field_id}">{label}</label>'
        f'<input id="{field_id}" type="{input_type}"{required} /></div>'
    )


def _render_form_table(table: FormTableModel) -> str:
    head_cells = []
    if table.field_rows and table.row_headers:
        head_cells.append('<th scope="col">Process step</th>')
    head_cells.extend(f'<th scope="col">{escape(header)}</th>' for header in table.headers)
    head = "".join(head_cells)
    rows = []
    if table.field_rows:
        used_ids: set[str] = set()
        for row_index, row in enumerate(table.field_rows, start=1):
            cells = []
            if table.row_headers and row_index <= len(table.row_headers):
                cells.append(f'<th scope="row">{escape(table.row_headers[row_index - 1])}</th>')
            for col_index, field in enumerate(row):
                header = table.headers[col_index] if col_index < len(table.headers) else field.label
                cells.append(
                    '<td>'
                    f'{_render_facsimile_table_control(field, header=header, row_index=row_index, used_ids=used_ids)}'
                    '</td>'
                )
            rows.append(f'<tr>{"".join(cells)}</tr>')
    else:
        for row in table.rows or []:
            cells = []
            for index, cell in enumerate(row):
                if index == 0:
                    cells.append(f'<th scope="row">{escape(cell)}</th>')
                else:
                    cells.append(f'<td>{escape(cell)}</td>')
            rows.append(f'<tr>{"".join(cells)}</tr>')
    return (
        '<div class="doc-table"><table>'
        f'<caption>{escape(table.caption)}</caption>'
        f'<thead><tr>{head}</tr></thead>'
        f'<tbody>{"".join(rows)}</tbody>'
        '</table></div>'
    )


def _render_facsimile_table_control(
    field: FormField,
    *,
    header: str,
    row_index: int,
    used_ids: set[str],
) -> str:
    field_id = _unique_slug(field.id or f"{header} row {row_index}", used_ids)
    label = field.label if field.label else f"{header}, row {row_index}"
    required = " required" if field.required else ""
    control = field.field_type
    if control in {"textarea", "signature"}:
        return (
            f'<label class="sr-only" for="{field_id}">{escape(label)}</label>'
            f'<textarea id="{field_id}" rows="2"{required}></textarea>'
        )
    input_type = control if control in {"email", "tel", "url", "number", "date"} else "text"
    return (
        f'<label class="sr-only" for="{field_id}">{escape(label)}</label>'
        f'<input id="{field_id}" type="{input_type}"{required} />'
    )


def _form_tables_from_lines(lines: Sequence[str]) -> list[FormTableModel]:
    income_rows = _income_table_rows(lines)
    if not income_rows:
        return []
    return [
        FormTableModel(
            caption="2024 income levels by family size",
            headers=["Family size", "Income"],
            rows=[[size, income] for size, income in income_rows],
        )
    ]


def _instruction_lines(lines: Sequence[str]) -> list[str]:
    for index, line in enumerate(lines):
        if line.strip().lower() == "instructions":
            collected = [line]
            for next_line in lines[index + 1:]:
                if _looks_like_form_section_start(next_line):
                    break
                collected.append(next_line)
            return collected
    return []


def _looks_like_form_section_start(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    lowered = stripped.lower().rstrip(":")
    explicit_headings = {
        "application",
        "term information",
        "student information",
        "contact information",
        "eligibility",
        "course information",
        "school information",
        "college information",
        "applicant certification",
        "certification",
        "for office use only",
        "office use only",
        "parent/guardian information",
        "student certification",
        "petition",
        "request",
    }
    if lowered in explicit_headings:
        return True
    return bool(re.match(r"^\d+\.\)\s+", stripped))


def _instructions_precede_first_form_section(lines: Sequence[str]) -> bool:
    instruction_index: int | None = None
    first_section_index: int | None = None
    for index, line in enumerate(lines):
        if line.strip().lower() == "instructions" and instruction_index is None:
            instruction_index = index
        if _looks_like_form_section_start(line) and first_section_index is None:
            first_section_index = index
    return instruction_index is not None and (
        first_section_index is None or instruction_index < first_section_index
    )


def _first_matching_line(lines: Sequence[str], needles: Sequence[str]) -> str | None:
    for line in lines:
        lowered = line.lower()
        if any(needle in lowered for needle in needles):
            return line
    return None


def _form_description(title: str, sections: Sequence[FormSectionModel]) -> str:
    field_count = sum(len(section.fields) for section in sections)
    return f"Accessible HTML form facsimile for {title} with {field_count} transcribed fields."


def _infer_form_language(
    *,
    source_hint: str,
    title: str,
    page_texts: Sequence[tuple[int, str]],
) -> str:
    haystack = " ".join([source_hint, title, *(text[:3000] for _, text in page_texts)]).lower()
    spanish_markers = (
        "spanish",
        "español",
        "apellido",
        "dirección",
        "fecha de nacimiento",
        "seguro social",
        "solicitud",
    )
    return "es" if any(marker in haystack for marker in spanish_markers) else "en"


def _repair_form_title(
    title: str,
    *,
    source_hint: str,
    page_texts: Sequence[tuple[int, str]],
) -> str:
    source_title = _form_title_from_source_hint(source_hint)
    visible_title = _form_title_candidate(
        [_normalize_title_text(line.strip()) for line in page_texts[0][1].splitlines() if line.strip()]
    ) if page_texts else None
    if source_title and "(AB 540)" in source_title and "nonresident tuition exemption" in title.lower():
        return source_title
    if visible_title and source_title and len(visible_title) > len(source_title) + 8:
        return visible_title
    if source_title and _is_bad_form_title(title):
        return source_title
    if source_title and len(title.split()) <= 3 and "laccd" in source_title.lower():
        return source_title
    return title


def _is_bad_form_title(title: str) -> bool:
    lowered = title.lower().strip()
    return (
        not lowered
        or "_" in title
        or "____" in title
        or ".indd" in lowered
        or " fall winter" in lowered
        or lowered.startswith(("please complete", "student information", "reclassification requested"))
        or lowered.startswith(("a proveer", "to provide"))
        or lowered in {"los angeles community college district", "form", "application"}
        or len(title) > 100
    )


def _form_title_from_source_hint(source_hint: str) -> str | None:
    from urllib.parse import unquote

    name = unquote(Path(source_hint.rstrip("/")).name)
    if not name:
        return None
    stem = re.sub(r"\.(pdf|docx|xlsx)$", "", name, flags=re.IGNORECASE)
    stem = stem.replace("_", " ")
    stem = re.sub(r"\s+", " ", stem).strip()
    stem = re.sub(r"\b(?:v|REV)?\s*\d{5,8}\b$", "", stem, flags=re.IGNORECASE).strip()
    stem = re.sub(r"\b2\.\d+\s+Fillable\b", "", stem, flags=re.IGNORECASE).strip()
    stem = re.sub(r"\s+v\d+$", "", stem, flags=re.IGNORECASE).strip()
    if not stem:
        return None
    if re.search(r"\bnonresident tuition exemption request\b", stem, flags=re.IGNORECASE):
        return "California Nonresident Tuition Exemption Request (AB 540)"
    words = []
    acronyms = {"LACCD", "AB", "EW", "SSN", "ID", "PDF", "K-12"}
    lowercase_words = {"a", "an", "and", "for", "of", "the", "to", "in"}
    for word in stem.split():
        upper = word.upper()
        if upper in acronyms or re.fullmatch(r"\d+", word):
            words.append(upper if upper != "K-12" else "K-12")
        elif word.lower() in lowercase_words:
            words.append(word.lower())
        else:
            words.append(word.capitalize())
    title = " ".join(words)
    title = re.sub(r"\bApplication for Noncredit Admission Spanish\b", "Application for Noncredit Admission (Spanish)", title)
    return title.strip()


def _should_prefer_source_form_title(source_title: str, visible_candidate: str) -> bool:
    form_keywords = {
        "application",
        "certification",
        "consent",
        "form",
        "petition",
        "questionnaire",
        "request",
        "waiver",
    }
    source_lower = source_title.lower()
    candidate_lower = visible_candidate.lower()
    return (
        not _is_strong_visible_title(visible_candidate)
        and any(keyword in source_lower for keyword in form_keywords)
        and not any(keyword in candidate_lower for keyword in form_keywords)
    )


def _humanize_field_label(label: str) -> str:
    cleaned = _clean_field_label(label)
    cleaned = re.sub(r"\bes\s*:\s*[^ ]+", "", cleaned)
    cleaned = re.sub(r"\bundefined\s*\d*\b", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^enter text\s*\d*$", "Text entry", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"([a-z])([A-Z])", r"\1 \2", cleaned)
    cleaned = re.sub(r"\bRow\s*(\d+)\b", r"Row \1", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" :-_")
    cleaned = re.sub(r"^(Signature|Date|Text)\s*\d+$", r"\1", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\bEMail\b", "Email", cleaned)
    cleaned = re.sub(r"\bLACCD\b", "LACCD", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" :-_")
    return cleaned or "Form field"


def _line_for_field_label(label: str, all_text: str) -> str | None:
    if not label or label == "Form field":
        return None
    needle = re.sub(r"[^a-z0-9]+", " ", label.lower()).strip()
    if not needle:
        return None
    for raw_line in all_text.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        normalized = re.sub(r"[^a-z0-9]+", " ", line.lower()).strip()
        if needle and needle in normalized:
            return line
    return None


def _line_has_yes_no_choice(label: str, line: str | None) -> bool:
    if not line:
        return False
    lowered = line.lower()
    return bool(re.search(r"\byes\b.*\bno\b|\bno\b.*\byes\b", lowered)) and len(label.split()) >= 4


def _label_before_yes_no(line: str) -> str | None:
    text = re.sub(r"\s+", " ", line).strip()
    match = re.search(r"\b(?:yes|no)\b\s+\b(?:yes|no)\b\s*$", text, flags=re.IGNORECASE)
    if not match:
        return None
    label = text[: match.start()].strip(" :-")
    if not label or len(label.split()) > 30:
        return None
    return re.sub(r"^\s*\d+[\.\)]\s*", "", label).strip()


def _clean_pdf_visual_line(line: str) -> str:
    text = re.sub(r"\s+", " ", line).strip()
    replacements = {
        "Nonresident Tuition Fee Waiver Application": "Nonresident Tuition Fee Waiver Application",
        "Eligib ility": "Eligibility",
        "care fully": "carefully",
        "immigra tion": "immigration",
        "reside ncy": "residency",
        "est ablishing": "establishing",
        "fam ily i ncome": "family income",
        "The se": "These",
        "pove rty": "poverty",
        "pub lished": "published",
        "Re gulation s": "Regulations",
        "s tuden t": "student",
        "P overty Guideline s": "Poverty Guidelines",
        "U .S.": "U.S.",
        "CERTIFIC ATION": "CERTIFICATION",
        "P LEASE": "PLEASE",
        "th e": "the",
        "th is": "this",
        "g uidelines": "guidelines",
        "responsib le": "responsible",
        "reimbursi ng": "reimbursing",
        "colleg e": "college",
        "f ill": "fill",
        "f ees": "fees",
        "suspen sion": "suspension",
        "expulsi on": "expulsion",
        "Bo ard": "Board",
        "for m": "form",
        "Re cor ds": "Records",
        "Recor ds": "Records",
        "guidelines apply , to you": "guidelines apply to you,",
        "guidelines apply ,": "guidelines apply,",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    return text.strip()


def _extract_instructions_text(lines: Sequence[str]) -> str:
    parts = [line for line in lines if line.lower() != "instructions"]
    return re.sub(r"\s+", " ", " ".join(parts)).strip()


def _income_table_rows(lines: Sequence[str]) -> list[tuple[str, str]]:
    rows: list[tuple[str, str]] = []
    for line in lines:
        if "|" in line:
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if len(cells) >= 2:
                size, income = cells[0], cells[1]
                if (
                    not re.fullmatch(r"-+", size)
                    and not size.lower().startswith("family")
                    and re.fullmatch(r"(?:Each Additional Family Member|\d+)", size, flags=re.IGNORECASE)
                    and re.fullmatch(r"\$[\d,]+", income)
                ):
                    rows.append((size, income))
                    continue
        match = re.search(
            r"^(?P<size>Each Additional Family Member|\d+)\s+(?P<income>\$[\d,]+)$",
            line,
            flags=re.IGNORECASE,
        )
        if match:
            rows.append((match.group("size"), match.group("income")))
    return rows


def _control_type_for_acrofield(field_type: str, label: str, field: Any) -> str:
    lowered = label.lower()
    flags = int(field.get("/Ff") or 0)
    if field_type == "Btn" and flags & 65536:
        return "button"
    if field_type == "Btn":
        return "radio" if _field_options(field) else "checkbox"
    if field_type == "Sig" or "signature" in lowered:
        return "signature"
    if "email" in lowered or "e-mail" in lowered:
        return "email"
    if "telephone" in lowered or "phone" in lowered:
        return "tel"
    if "date" in lowered or "mm/dd/yyyy" in lowered or "mmdd" in lowered:
        return "date"
    if flags & 4096:
        return "textarea"
    if field_type == "Ch":
        return "select"
    return "text"


def _field_options(field: Any) -> list[str]:
    raw = field.get("/Opt") or []
    options: list[str] = []
    if isinstance(raw, (list, tuple)):
        for value in raw:
            if isinstance(value, (list, tuple)) and value:
                value = value[-1]
            text = str(value).strip()
            if text:
                options.append(text)
    if not options:
        states = field.get("/_States_") or []
        if isinstance(states, (list, tuple)):
            state_options: list[str] = []
            for value in states:
                text = str(value).strip("/() ")
                text = _normalize_acroform_option_label(text)
                if text and text.lower() not in {"off", "on"}:
                    state_options.append(text)
            if {option.lower() for option in state_options} != {"yes"}:
                options.extend(state_options)
    return _dedupe(options)


def _normalize_acroform_option_label(label: str) -> str:
    cleaned = label.strip()
    cleaned = re.sub(r"_(?:\d+)$", "", cleaned)
    if cleaned.upper() in {"YES", "NO"}:
        return cleaned.capitalize()
    return cleaned.replace("_", " ")


def _clean_field_label(label: str) -> str:
    label = label.replace("_", " ")
    label = re.sub(r"\s+", " ", label)
    label = re.sub(r"\s+None$", "", label, flags=re.IGNORECASE)
    return label.strip(" :-")


def _acroform_fields_to_signals(
    fields: Sequence[FormField],
    *,
    start_order: int,
) -> list[RawBlockSignal]:
    if not fields:
        return []
    signals = [
        RawBlockSignal(
            id=f"heading-{start_order}",
            kind="heading",
            order=start_order,
            source="extract",
            source_key="acroform-heading",
            page_start=1,
            page_end=1,
            level=2,
            text="PDF Form Fields",
        )
    ]
    order = start_order + 1
    for field in fields:
        signals.append(
            RawBlockSignal(
                id=f"form-field-{order}",
                kind="form_field",
                order=order,
                source="extract",
                source_key=f"acroform-{field.id}",
                page_start=1,
                page_end=1,
                field=field,
                group_name="PDF Form Fields",
            )
        )
        order += 1
    return signals


def _raw_blocks_from_markdown(
    markdown: str,
    *,
    source: SignalSource = "extract",
    page_num: int = 1,
    order_start: int = 0,
    source_prefix: str = "markdown",
) -> tuple[str | None, list[RawBlockSignal]]:
    blocks: list[RawBlockSignal] = []
    order = order_start
    title: str | None = None
    lines = markdown.splitlines()
    i = 0

    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()
        if not line:
            i += 1
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            if level == 1 and not title:
                title = text
            else:
                blocks.append(
                    RawBlockSignal(
                        id=f"heading-{order}",
                        kind="heading",
                        order=order,
                        source=source,
                        source_key=f"{source_prefix}-heading-{order}",
                        page_start=page_num,
                        page_end=page_num,
                        level=max(2, level),
                        text=text,
                    )
                )
                order += 1
            i += 1
            continue

        if line.startswith("|"):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [_split_markdown_table_row(table_line) for table_line in table_lines]
            if len(rows) >= 2 and all(set(cell) <= {"-"} for cell in rows[1] if cell):
                header_rows = [0]
                body_rows = rows[2:] if len(rows) > 2 else rows[:1]
                rows = [rows[0], *body_rows]
            else:
                header_rows = [0] if len(rows) > 1 else []
            blocks.append(
                RawBlockSignal(
                    id=f"table-{order}",
                    kind="table",
                    order=order,
                    source=source,
                    source_key=f"{source_prefix}-table-{order}",
                    page_start=page_num,
                    page_end=page_num,
                    rows=rows,
                    header_rows=header_rows,
                )
            )
            order += 1
            continue

        if line.startswith("- "):
            items = [line[2:].strip()]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("- "):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            form_signals, form_next_order = _signals_from_chunk(
                items,
                page_num=page_num,
                chunk_index=order,
                order=order,
                source=source,
            )
            if any(signal.kind == "form_field" for signal in form_signals):
                blocks.extend(form_signals)
                order = form_next_order
                continue
            blocks.append(
                RawBlockSignal(
                    id=f"list-{order}",
                    kind="list",
                    order=order,
                    source=source,
                    source_key=f"{source_prefix}-list-{order}",
                    page_start=page_num,
                    page_end=page_num,
                    items=items,
                )
            )
            order += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or re.match(r"^(#{1,6})\s+", nxt) or nxt.startswith("|") or nxt.startswith("- "):
                break
            paragraph_lines.append(nxt)
            i += 1
        form_signals, form_next_order = _signals_from_chunk(
            paragraph_lines,
            page_num=page_num,
            chunk_index=order,
            order=order,
            source=source,
        )
        if any(signal.kind == "form_field" for signal in form_signals):
            blocks.extend(form_signals)
            order = form_next_order
            continue
        paragraph_signals = _paragraph_signals(
            text=" ".join(paragraph_lines),
            order=order,
            source=source,
            source_key=f"{source_prefix}-paragraph-{order}",
            page_start=page_num,
            page_end=page_num,
        )
        blocks.extend(paragraph_signals)
        order += len(paragraph_signals)

    return title, blocks


def _raw_blocks_from_llamaparse_document(
    document: LlamaParseDocument,
    *,
    source_hint: str,
) -> tuple[str | None, list[RawBlockSignal]]:
    title: str | None = None
    blocks: list[RawBlockSignal] = []
    next_order = 0

    for page in document.pages:
        page_title, item_blocks = _raw_blocks_from_llamaparse_items(
            page.items,
            page_num=page.page_number,
            order_start=next_order,
        )
        if page_title and title is None:
            title = page_title
        if item_blocks:
            blocks.extend(item_blocks)
            next_order = max((block.order for block in blocks), default=next_order) + 1
            continue

        content = (page.markdown or page.text or "").strip()
        if not content:
            continue
        page_title, page_blocks = _raw_blocks_from_markdown(
            content,
            source="llamaparse",
            page_num=page.page_number,
            order_start=next_order,
            source_prefix=f"llamaparse-page-{page.page_number}",
        )
        if page_title and title is None:
            title = page_title
        blocks.extend(page_blocks)
        next_order = max((block.order for block in blocks), default=next_order) + 1

    if title is None:
        title_texts = [
            (page.page_number, page.text or _markdown_text_for_title(page.markdown))
            for page in document.pages
        ]
        title = _choose_document_title(
            title_texts,
            source_hint=source_hint,
            metadata_title=None,
        )
    return title, blocks


def _raw_blocks_from_llamaparse_items(
    items: Sequence[dict[str, Any]],
    *,
    page_num: int,
    order_start: int,
) -> tuple[str | None, list[RawBlockSignal]]:
    title: str | None = None
    blocks: list[RawBlockSignal] = []
    order = order_start

    for item_index, raw_item in enumerate(items):
        if not isinstance(raw_item, Mapping):
            continue
        item = dict(raw_item)
        kind = _llamaparse_item_kind(item)
        source_key = f"llamaparse-page-{page_num}-item-{item_index}"

        if kind == "heading":
            text = _llamaparse_item_text(item)
            if not text:
                continue
            level = _llamaparse_heading_level(item)
            if level == 1 and title is None:
                title = text
                continue
            blocks.append(
                RawBlockSignal(
                    id=f"llamaparse-heading-{order}",
                    kind="heading",
                    order=order,
                    source="llamaparse",
                    source_key=f"{source_key}-heading",
                    page_start=page_num,
                    page_end=page_num,
                    level=2,
                    text=text,
                    metadata={"llamaparse_level": level},
                )
            )
            order += 1
            continue

        if kind == "table":
            rows = _llamaparse_table_rows(item)
            if not rows:
                continue
            blocks.append(
                RawBlockSignal(
                    id=f"llamaparse-table-{order}",
                    kind="table",
                    order=order,
                    source="llamaparse",
                    source_key=f"{source_key}-table",
                    page_start=page_num,
                    page_end=page_num,
                    rows=rows,
                    header_rows=[0] if len(rows) > 1 else [],
                    metadata={"parse_concerns": item.get("parse_concerns") or []},
                )
            )
            order += 1
            continue

        if kind == "list":
            list_items = _llamaparse_list_items(item)
            if not list_items:
                continue
            blocks.append(
                RawBlockSignal(
                    id=f"llamaparse-list-{order}",
                    kind="list",
                    order=order,
                    source="llamaparse",
                    source_key=f"{source_key}-list",
                    page_start=page_num,
                    page_end=page_num,
                    items=list_items,
                    ordered=bool(item.get("ordered")),
                )
            )
            order += 1
            continue

        if kind == "image":
            caption = _clean_llamaparse_text(item.get("caption")) or _llamaparse_item_text(item)
            if not caption:
                continue
            blocks.append(
                RawBlockSignal(
                    id=f"llamaparse-caption-{order}",
                    kind="caption",
                    order=order,
                    source="llamaparse",
                    source_key=f"{source_key}-caption",
                    page_start=page_num,
                    page_end=page_num,
                    text=caption,
                    caption=caption,
                )
            )
            order += 1
            continue

        item_lines = _llamaparse_item_lines(item)
        if not item_lines:
            continue
        item_signals, next_order = _signals_from_chunk(
            item_lines,
            page_num=page_num,
            chunk_index=item_index,
            order=order,
            source="llamaparse",
        )
        blocks.extend(item_signals)
        order = next_order

    return title, blocks


def _llamaparse_item_kind(item: Mapping[str, Any]) -> str:
    raw = item.get("type")
    if isinstance(raw, str) and raw.strip():
        return raw.strip().lower()
    if "rows" in item or "csv" in item:
        return "table"
    if "level" in item:
        return "heading"
    if "items" in item and isinstance(item.get("items"), Sequence):
        return "list"
    if "url" in item and "caption" in item:
        return "image"
    return "text"


def _llamaparse_item_text(item: Mapping[str, Any]) -> str:
    for key in ("value", "text", "md"):
        text = _clean_llamaparse_text(item.get(key))
        if text:
            if key == "md":
                text = re.sub(r"^#{1,6}\s+", "", text)
            return text
    return ""


def _llamaparse_item_lines(item: Mapping[str, Any]) -> list[str]:
    for key in ("value", "text", "md"):
        value = item.get(key)
        if not isinstance(value, str) or not value.strip():
            continue
        lines = [_clean_llamaparse_text(line) for line in value.splitlines()]
        lines = [re.sub(r"^#{1,6}\s+", "", line) for line in lines if line]
        if lines:
            return lines
    return []


def _llamaparse_heading_level(item: Mapping[str, Any]) -> int:
    raw = item.get("level")
    if isinstance(raw, int) and 1 <= raw <= 6:
        return raw
    try:
        value = int(str(raw))
    except (TypeError, ValueError):
        return 2
    return min(6, max(1, value))


def _llamaparse_table_rows(item: Mapping[str, Any]) -> list[list[str]]:
    raw_rows = item.get("rows")
    rows: list[list[str]] = []
    if isinstance(raw_rows, Sequence) and not isinstance(raw_rows, (str, bytes)):
        for raw_row in raw_rows:
            if not isinstance(raw_row, Sequence) or isinstance(raw_row, (str, bytes)):
                continue
            row = [_clean_llamaparse_text(cell) for cell in raw_row]
            if any(row):
                rows.append(row)
    if not rows:
        md = _clean_llamaparse_text(item.get("md"))
        if md:
            table_lines = [line.strip() for line in md.splitlines() if line.strip().startswith("|")]
            if table_lines:
                parsed_rows = [_split_markdown_table_row(line) for line in table_lines]
                if len(parsed_rows) >= 2 and all(set(cell) <= {"-"} for cell in parsed_rows[1] if cell):
                    parsed_rows = [parsed_rows[0], *parsed_rows[2:]]
                rows = parsed_rows
    return _normalize_table_rows(rows) if rows else []


def _llamaparse_list_items(item: Mapping[str, Any]) -> list[str]:
    raw_items = item.get("items")
    values: list[str] = []
    if isinstance(raw_items, Sequence) and not isinstance(raw_items, (str, bytes)):
        for raw_item in raw_items:
            values.extend(_llamaparse_list_item_texts(raw_item))
    if not values:
        md = _clean_llamaparse_text(item.get("md"))
        if md:
            for line in md.splitlines():
                stripped = line.strip()
                if _LIST_ITEM_RE.match(stripped):
                    values.append(_strip_list_marker(stripped))
    return _dedupe([value for value in values if value])


def _llamaparse_list_item_texts(raw_item: Any) -> list[str]:
    if isinstance(raw_item, Mapping):
        item = dict(raw_item)
        if _llamaparse_item_kind(item) == "list":
            values: list[str] = []
            nested = item.get("items")
            if isinstance(nested, Sequence) and not isinstance(nested, (str, bytes)):
                for child in nested:
                    values.extend(_llamaparse_list_item_texts(child))
            return values
        text = _llamaparse_item_text(item)
        return [text] if text else []
    text = _clean_llamaparse_text(raw_item)
    return [text] if text else []


def _clean_llamaparse_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _markdown_text_for_title(markdown: str) -> str:
    lines = []
    for line in markdown.splitlines():
        stripped = line.strip()
        stripped = re.sub(r"^#{1,6}\s+", "", stripped)
        if stripped and not stripped.startswith("|"):
            lines.append(stripped)
    return "\n".join(lines)


def _split_markdown_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _workbook_has_formula(workbook) -> bool:
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    return True
    return False


def _xlsx_sheet_rows(worksheet) -> list[list[str]]:
    rows: list[list[str]] = []
    max_col = worksheet.max_column or 0
    for row in worksheet.iter_rows(max_col=max_col):
        values = ["" if cell.value is None else str(cell.value) for cell in row]
        if any(value.strip() for value in values):
            rows.append(values)
    return rows


def _raw_blocks_from_page_texts(
    page_texts: Sequence[tuple[int, str]],
    *,
    source: SignalSource,
    title_hint: str,
    disable_form_heuristics: bool = False,
) -> list[RawBlockSignal]:
    blocks: list[RawBlockSignal] = []
    order = 0
    title_emitted = False

    for page_num, page_text in page_texts:
        for chunk_index, chunk_lines in enumerate(_chunk_lines(page_text)):
            if not chunk_lines:
                continue

            lines = list(chunk_lines)
            if not title_emitted:
                lines = _prune_title_from_lines(lines, title_hint)
                blocks.append(
                    RawBlockSignal(
                        id=f"title-{order}",
                        kind="title",
                        order=order,
                        source=source,
                        source_key=f"page-{page_num}-chunk-{chunk_index}-title",
                        page_start=page_num,
                        page_end=page_num,
                        text=title_hint,
                    )
                )
                order += 1
                title_emitted = True
                if not lines:
                    continue

            emitted, order = _signals_from_chunk(
                lines,
                page_num=page_num,
                chunk_index=chunk_index,
                order=order,
                source=source,
                disable_form_heuristics=disable_form_heuristics,
            )
            blocks.extend(emitted)

    if title_emitted:
        return blocks

    return [
        RawBlockSignal(
            id="title-0",
            kind="title",
            order=0,
            source=source,
            source_key="fallback-title",
            page_start=1,
            page_end=1,
            text=title_hint,
        )
    ]


def _signals_from_chunk(
    lines: Sequence[str],
    *,
    page_num: int,
    chunk_index: int,
    order: int,
    source: SignalSource,
    disable_form_heuristics: bool = False,
) -> tuple[list[RawBlockSignal], int]:
    emitted: list[RawBlockSignal] = []
    remaining = [
        line.strip()
        for line in lines
        if line.strip() and not _REVISION_LINE_RE.match(line.strip())
    ]

    figure_signals = _figure_signals_from_lines(
        remaining,
        page_num=page_num,
        chunk_index=chunk_index,
        order=order,
        source=source,
    )
    if figure_signals is not None:
        return figure_signals

    if not disable_form_heuristics:
        form_signals = _form_signals_from_lines(
            remaining,
            page_num=page_num,
            chunk_index=chunk_index,
            order=order,
            source=source,
        )
        if form_signals:
            return form_signals, order + len(form_signals)

    if remaining and _looks_like_callout_label(remaining[0]) and len(remaining) > 1:
        label = remaining.pop(0).rstrip(":").strip()
        emitted.append(
            RawBlockSignal(
                id=f"callout-{order}",
                kind="callout",
                order=order,
                source=source,
                source_key=f"page-{page_num}-chunk-{chunk_index}-callout",
                page_start=page_num,
                page_end=page_num,
                text=" ".join(remaining),
                group_name=label,
            )
        )
        return emitted, order + 1

    if remaining and _looks_like_heading(remaining[0]):
        heading = remaining.pop(0)
        emitted.append(
            RawBlockSignal(
                id=f"heading-{order}",
                kind="heading",
                order=order,
                source=source,
                source_key=f"page-{page_num}-chunk-{chunk_index}-heading",
                page_start=page_num,
                page_end=page_num,
                level=2,
                text=heading,
            )
        )
        order += 1
        if not remaining:
            return emitted, order

    if _looks_like_list_block(remaining):
        emitted.append(
            RawBlockSignal(
                id=f"list-{order}",
                kind="list",
                order=order,
                source=source,
                source_key=f"page-{page_num}-chunk-{chunk_index}-list",
                page_start=page_num,
                page_end=page_num,
                items=[_strip_list_marker(line) for line in remaining],
            )
        )
        return emitted, order + 1

    if _looks_like_table_block(remaining):
        rows = _normalize_table_rows([_split_table_row(line) for line in remaining])
        header_rows = [0] if len(rows) > 1 else []
        emitted.append(
            RawBlockSignal(
                id=f"table-{order}",
                kind="table",
                order=order,
                source=source,
                source_key=f"page-{page_num}-chunk-{chunk_index}-table",
                page_start=page_num,
                page_end=page_num,
                rows=rows,
                header_rows=header_rows,
            )
        )
        return emitted, order + 1

    paragraph_signals = _paragraph_signals(
        text=" ".join(remaining),
        order=order,
        source=source,
        source_key=f"page-{page_num}-chunk-{chunk_index}-paragraph",
        page_start=page_num,
        page_end=page_num,
    )
    emitted.extend(paragraph_signals)
    return emitted, order + len(paragraph_signals)


def _figure_signals_from_lines(
    lines: Sequence[str],
    *,
    page_num: int,
    chunk_index: int,
    order: int,
    source: SignalSource,
) -> tuple[list[RawBlockSignal], int] | None:
    if not lines:
        return None
    match = _FIGURE_START_RE.match(lines[0].strip())
    if not match:
        return None

    caption_lines = [match.group(2).strip() or lines[0].strip()]
    index = 1
    while index < len(lines):
        line = lines[index].strip()
        if _looks_like_label_cloud(line):
            index += 1
            continue
        if _looks_like_heading(line) and caption_lines and caption_lines[-1].endswith("."):
            break
        caption_lines.append(line)
        index += 1
        if line.endswith(".") and index < len(lines) and _looks_like_label_cloud(lines[index]):
            break

    caption = _join_caption_lines(caption_lines)
    emitted = [
        RawBlockSignal(
            id=f"figure-{order}",
            kind="figure",
            order=order,
            source=source,
            source_key=f"page-{page_num}-chunk-{chunk_index}-figure",
            page_start=page_num,
            page_end=page_num,
            caption=caption,
            text=caption,
        )
    ]
    next_order = order + 1

    rest = list(lines[index:])
    while len(rest) > 1 and (
        _looks_like_label_cloud(rest[0])
        or (_looks_like_heading(rest[0]) and _looks_like_heading(rest[1]))
    ):
        rest.pop(0)
    if rest:
        rest_signals, next_order = _signals_from_chunk(
            rest,
            page_num=page_num,
            chunk_index=chunk_index,
            order=next_order,
            source=source,
        )
        emitted.extend(rest_signals)
    return emitted, next_order


def _form_signals_from_lines(
    lines: Sequence[str],
    *,
    page_num: int,
    chunk_index: int,
    order: int,
    source: SignalSource,
) -> list[RawBlockSignal]:
    if not lines or not any(_looks_like_form_text(line) for line in lines):
        return []

    emitted: list[RawBlockSignal] = []
    active_group: str | None = None
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue

        option_lines: list[str] = []
        if _is_checkbox_option_line(line):
            label = active_group or "Options"
            while index < len(lines) and _is_checkbox_option_line(lines[index]):
                option_lines.append(lines[index])
                index += 1
            field = _checkbox_group_field(label, option_lines)
            emitted.append(
                _form_field_signal(
                    field,
                    order=order + len(emitted),
                    page_num=page_num,
                    chunk_index=chunk_index,
                    source=source,
                    group_name=label,
                )
            )
            continue

        if _is_instruction_with_option_lines(lines, index):
            label = line.rstrip(".:")
            index += 1
            while index < len(lines) and _is_checkbox_option_line(lines[index]):
                option_lines.append(lines[index])
                index += 1
            field = _checkbox_group_field(label, option_lines)
            emitted.append(
                _form_field_signal(
                    field,
                    order=order + len(emitted),
                    page_num=page_num,
                    chunk_index=chunk_index,
                    source=source,
                    group_name=label,
                )
            )
            continue

        bracket_fields = [] if _has_number_prefix(line) else _bracket_placeholder_fields_from_line(line)
        if bracket_fields:
            for field in bracket_fields:
                emitted.append(
                    _form_field_signal(
                        field,
                        order=order + len(emitted),
                        page_num=page_num,
                        chunk_index=chunk_index,
                        source=source,
                        group_name=active_group or "Form Fields",
                    )
                )
            index += 1
            while index < len(lines) and lines[index].strip().startswith("("):
                index += 1
            continue

        heading_candidate = line.rstrip(":").strip()
        if (
            line.endswith(":")
            and not _has_number_prefix(line)
            and "___" not in line
            and not _options_from_text(line)
        ):
            active_group = heading_candidate
            emitted.append(
                RawBlockSignal(
                    id=f"heading-{order + len(emitted)}",
                    kind="heading",
                    order=order + len(emitted),
                    source=source,
                    source_key=f"page-{page_num}-chunk-{chunk_index}-heading-{index}",
                    page_start=page_num,
                    page_end=page_num,
                    level=2,
                    text=heading_candidate,
                )
            )
            index += 1
            continue

        label_row_fields = [] if _has_number_prefix(line) else _label_row_fields_from_line(line)
        if label_row_fields:
            for field in label_row_fields:
                emitted.append(
                    _form_field_signal(
                        field,
                        order=order + len(emitted),
                        page_num=page_num,
                        chunk_index=chunk_index,
                        source=source,
                        group_name=active_group or "Form Fields",
                    )
                )
            index += 1
            while index < len(lines) and lines[index].strip().startswith("("):
                index += 1
            continue

        yes_no_label = _label_before_yes_no(line)
        if yes_no_label:
            label = _clean_field_label(yes_no_label)
            field = FormField(
                id=_slugify(label),
                label=label,
                field_type="radio",
                options=["Yes", "No"],
            )
            emitted.append(
                _form_field_signal(
                    field,
                    order=order + len(emitted),
                    page_num=page_num,
                    chunk_index=chunk_index,
                    source=source,
                    group_name=active_group,
                )
            )
            index += 1
            continue

        fill_fields = [] if _has_number_prefix(line) else _fill_fields_from_line(line)
        if fill_fields:
            for field in fill_fields:
                emitted.append(
                    _form_field_signal(
                        field,
                        order=order + len(emitted),
                        page_num=page_num,
                        chunk_index=chunk_index,
                        source=source,
                        group_name=active_group or "Form Fields",
                    )
                )
            index += 1
            continue

        inline_options = _options_from_text(line)
        inline_label = _label_before_options(line)
        if inline_options and inline_label:
            label = _clean_field_label(inline_label)
            index += 1
            while index < len(lines) and _is_checkbox_option_line(lines[index]):
                inline_options.extend(_options_from_text(lines[index]))
                index += 1
            inline_options = _dedupe(inline_options)
            field = FormField(
                id=_slugify(label),
                label=label,
                field_type=_choice_type_from_label_options(label, inline_options),
                options=inline_options,
            )
            emitted.append(
                _form_field_signal(
                    field,
                    order=order + len(emitted),
                    page_num=page_num,
                    chunk_index=chunk_index,
                    source=source,
                    group_name=active_group,
                )
            )
            continue

        numbered = _NUMBERED_FIELD_RE.match(line)
        if numbered and (_has_number_prefix(line) or "___" in numbered.group(2)):
            label = _clean_field_label(numbered.group(1))
            remainder_parts = [numbered.group(2).strip()]
            index += 1
            while index < len(lines) and _is_checkbox_option_line(lines[index]):
                remainder_parts.append(lines[index].strip())
                index += 1
            options = _options_from_text(" ".join(remainder_parts))
            field = FormField(
                id=_slugify(label),
                label=label,
                field_type=(
                    _choice_type_from_label_options(label, options)
                    if options
                    else _field_type_from_label(label)
                ),
                options=options,
            )
            emitted.append(
                _form_field_signal(
                    field,
                    order=order + len(emitted),
                    page_num=page_num,
                    chunk_index=chunk_index,
                    source=source,
                    group_name=active_group,
                )
            )
            continue

        emitted.extend(
            _paragraph_signals(
                text=line,
                order=order + len(emitted),
                source=source,
                source_key=f"page-{page_num}-chunk-{chunk_index}-paragraph-{index}",
                page_start=page_num,
                page_end=page_num,
            )
        )
        index += 1

    return emitted


def _form_field_signal(
    field: FormField,
    *,
    order: int,
    page_num: int,
    chunk_index: int,
    source: SignalSource,
    group_name: str | None,
) -> RawBlockSignal:
    field_id = _slugify(field.id or field.label)
    field = field.model_copy(update={"id": f"{field_id}-{order}"})
    return RawBlockSignal(
        id=f"form-field-{order}",
        kind="form_field",
        order=order,
        source=source,
        source_key=f"page-{page_num}-chunk-{chunk_index}-field-{order}",
        page_start=page_num,
        page_end=page_num,
        field=field,
        group_name=group_name,
    )


def _looks_like_label_cloud(line: str) -> bool:
    text = line.strip()
    if not text:
        return False
    if re.search(r"[a-z][A-Z]", text):
        return True
    words = text.split()
    return 1 <= len(words) <= 3 and not text.endswith((".", ":")) and not _looks_like_heading(text)


def _join_caption_lines(lines: Sequence[str]) -> str:
    text = " ".join(line.strip() for line in lines if line.strip())
    text = re.sub(r"-\s+([a-z])", r"\1", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip() or "Figure"


def _looks_like_form_text(line: str) -> bool:
    text = line.strip()
    return (
        "___" in text
        or bool(_CHECK_OPTION_RE.search(text))
        or bool(_label_row_fields_from_line(text))
        or bool(re.match(r"^\s*\d+[\.\)]\s*[^:]{2,80}:\s*$", text))
        or bool(re.search(r"\b(signature|date|email|phone|name)\b.*_{3,}", text, re.IGNORECASE))
    )


def _is_checkbox_option_line(line: str) -> bool:
    text = re.sub(r"^\s*\d+[\.\)]\s*", "", line.strip())
    return bool(_CHECK_OPTION_RE.match(text))


def _is_instruction_with_option_lines(lines: Sequence[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    line = lines[index].strip()
    return (
        not _is_checkbox_option_line(line)
        and "___" not in line
        and _is_checkbox_option_line(lines[index + 1])
    )


def _has_number_prefix(line: str) -> bool:
    return bool(re.match(r"^\s*\d+[\.\)]\s+", line))


def _checkbox_group_field(label: str, option_lines: Sequence[str]) -> FormField:
    options = _options_from_text(" ".join(option_lines))
    if label == "Options" and len(options) == 1:
        option_label = options[0]
        return FormField(
            id=_slugify(option_label),
            label=option_label,
            field_type="checkbox",
        )
    return FormField(
        id=_slugify(label),
        label=_clean_field_label(label),
        field_type=_choice_type_from_label_options(label, options),
        options=options,
    )


def _options_from_text(text: str) -> list[str]:
    options: list[str] = []
    normalized = re.sub(r"^\s*\d+[\.\)]\s*", "", text)
    for match in _CHECK_OPTION_RE.finditer(normalized):
        value = match.group("label")
        value = re.split(r"\s{3,}", value.strip())[0]
        value = value.strip(" .;:*")
        if value:
            options.append(value)
    return _dedupe(options)


def _label_before_options(line: str) -> str | None:
    marker = re.search(r"_{2,}|□|☐|☑|☒|\[\s*[xX ]?\s*\]", line)
    if not marker:
        return None
    label = line[: marker.start()]
    label = re.sub(r"^\s*\d+[\.\)]\s*", "", label)
    return label.strip(" :") or None


def _choice_type_from_label_options(label: str, options: Sequence[str]) -> str:
    lowered = label.lower()
    normalized_options = {option.strip().lower().rstrip("*") for option in options}
    if "select all" in lowered:
        return "checkbox"
    if "whether" in lowered or "check one" in lowered or normalized_options == {"yes", "no"}:
        return "radio"
    return "checkbox"


def _field_type_from_label(label: str) -> str:
    lowered = label.lower()
    if "signature" in lowered:
        return "signature"
    if "email" in lowered:
        return "email"
    if "phone" in lowered or "telephone" in lowered:
        return "tel"
    if "date" in lowered:
        return "date"
    if "how many" in lowered or "number" in lowered or "quantity" in lowered:
        return "number"
    return "text"


def _fill_fields_from_line(line: str) -> list[FormField]:
    fields: list[FormField] = []
    for match in _FILL_LINE_RE.finditer(line):
        label = _clean_field_label(match.group("label"))
        if not label or len(label.split()) > 8:
            continue
        fields.append(
            FormField(
                id=_slugify(label),
                label=label,
                field_type=_field_type_from_label(label),
            )
        )
    return fields


def _bracket_placeholder_fields_from_line(line: str) -> list[FormField]:
    normalized = re.sub(r"\\([\[\]])", r"\1", line.strip())
    matches = list(_BRACKET_FILL_RE.finditer(normalized))
    if not matches or matches[0].start() == 0:
        return []

    fields: list[FormField] = []
    cursor = 0
    for match in matches:
        label = _clean_field_label(normalized[cursor:match.start()])
        if not label or not _looks_like_static_field_label(label):
            return []
        fields.append(
            FormField(
                id=_slugify(label),
                label=label,
                field_type=_field_type_from_label(label),
            )
        )
        cursor = match.end()

    trailing = normalized[cursor:].strip(" :;.,-")
    if trailing:
        return []
    return fields


def _label_row_fields_from_line(line: str) -> list[FormField]:
    parts = [part.strip(" :") for part in re.split(r"\s{2,}", line.strip()) if part.strip(" :")]
    if len(parts) < 2:
        if _looks_like_standalone_static_field_label(line):
            parts = [line.strip(" :")]
        else:
            return []
    fields: list[FormField] = []
    for part in parts:
        label = _clean_field_label(part)
        if not _looks_like_static_field_label(label):
            return []
        fields.append(
            FormField(
                id=_slugify(label),
                label=label,
                field_type=_field_type_from_label(label),
            )
        )
    return fields


def _looks_like_standalone_static_field_label(line: str) -> bool:
    label = _clean_field_label(line)
    return _looks_like_static_field_label(label) and len(label.split()) <= 5


def _looks_like_static_field_label(label: str) -> bool:
    lowered = label.lower().strip()
    if not lowered or lowered.startswith("(") or lowered.endswith("."):
        return False
    if len(label.split()) > 6:
        return False
    return bool(
        re.search(
            r"\b(name|first|last|middle|mi|id|date|birth|email|phone|signature|address|city|state|zip)\b",
            lowered,
        )
    )


def _chunk_lines(page_text: str) -> list[list[str]]:
    chunks: list[list[str]] = []
    current: list[str] = []
    for raw_line in page_text.splitlines():
        line = raw_line.strip()
        if not line:
            if current:
                chunks.append(current)
                current = []
            continue
        current.append(line)
    if current:
        chunks.append(current)
    return chunks


def _split_title_and_remainder(line: str) -> tuple[str, str | None]:
    text = line.strip()
    if not text:
        return _fallback_title(""), None

    parts = [part.strip() for part in re.split(r"\s{3,}", text) if part.strip()]
    if len(parts) >= 2:
        return parts[0], re.sub(r"\s{2,}", " ", " ".join(parts[1:])).strip()

    if "___" in text:
        head = text.split("___", 1)[0].strip()
        if head:
            remainder = re.sub(r"\s{2,}", " ", text[len(head):]).strip()
            return head, remainder or None

    return text, None


def _clean_metadata_title(raw_title: str | None) -> str | None:
    if not raw_title:
        return None
    title = _normalize_title_text(raw_title.strip())
    if not title:
        return None
    title = re.sub(r"^Microsoft Word\s*-\s*", "", title, flags=re.IGNORECASE)
    title = re.sub(r"\.(pdf|docx|xlsx)$", "", title, flags=re.IGNORECASE)
    title = re.sub(r"\s{2,}", " ", title).strip(" -")
    return title or None


def _normalize_title_text(text: str) -> str:
    text = re.sub(r"([a-z])([A-Z])", r"\1 \2", text)
    text = re.sub(r"\b([A-Za-z]{4,}?)(of|for|and)\b", r"\1 \2", text)
    text = re.sub(r"\s{2,}", " ", text)
    return text.strip()


def _choose_document_title(
    page_texts: Sequence[tuple[int, str]],
    *,
    source_hint: str,
    metadata_title: str | None = None,
) -> str:
    first_page_lines: list[str] = []
    if page_texts:
        first_page_text = page_texts[0][1]
        for raw_line in first_page_text.splitlines():
            line = _normalize_title_text(raw_line.strip())
            if not line:
                continue
            first_page_lines.append(line)

    combined = _combined_short_title_lines(first_page_lines)
    first_candidate = _first_title_candidate(first_page_lines)
    form_candidate = _form_title_candidate(first_page_lines)
    if metadata_title:
        if first_candidate and _is_strong_visible_title(first_candidate):
            return first_candidate
        if form_candidate:
            return form_candidate
        if len(metadata_title.split()) == 1 and combined and len(combined.split()) > 1:
            return combined
        if len(metadata_title.split()) == 1 and first_candidate and len(first_candidate.split()) > 1:
            return first_candidate
        return metadata_title

    if form_candidate:
        return form_candidate
    if combined:
        return combined
    source_form_title = _form_title_from_source_hint(source_hint)
    if source_form_title and first_candidate and _should_prefer_source_form_title(
        source_form_title,
        first_candidate,
    ):
        return source_form_title
    if first_candidate:
        return first_candidate

    return _fallback_title(source_hint)


def _is_strong_visible_title(candidate: str) -> bool:
    letters = [char for char in candidate if char.isalpha()]
    return (
        len(candidate.split()) >= 2
        and bool(letters)
        and sum(1 for char in letters if char.isupper()) / len(letters) > 0.75
    )


def _combined_short_title_lines(lines: Sequence[str]) -> str | None:
    candidates: list[str] = []
    for line in lines[:4]:
        if any(pattern.match(line) for pattern in _TITLE_SKIP_PATTERNS):
            continue
        if _looks_like_non_title_line(line):
            break
        if len(line.split()) > 3:
            break
        if line.lower().startswith(("fonts by", "formatting by", "abstract")):
            break
        candidates.append(line)
        if len(candidates) >= 3:
            break
    if len(candidates) >= 2:
        return " ".join(candidates)
    return None


def _form_title_candidate(lines: Sequence[str]) -> str | None:
    for index, line in enumerate(lines[:8]):
        if (
            index + 1 < len(lines)
            and line.lower().endswith((" for", " of", " and"))
            and 0 < len(lines[index + 1].split()) <= 8
        ):
            line = f"{line} {lines[index + 1]}"
        elif (
            index + 1 < len(lines)
            and line.lower() == "credit for prior learning"
            and "course equivalency" in lines[index + 1].lower()
        ):
            line = f"{line} – {lines[index + 1].title()}"
        candidate, _ = _split_title_and_remainder(line)
        candidate = _canonicalize_title_candidate(candidate)
        lowered = candidate.lower()
        if any(pattern.match(candidate) for pattern in _TITLE_SKIP_PATTERNS):
            continue
        if _looks_like_non_title_line(candidate):
            continue
        if lowered.startswith(("this ", "these ", "the ")) or "..." in candidate:
            continue
        if len(candidate.split()) > 14:
            continue
        if (
            any(
                token in lowered
                for token in (
                    "application",
                    "certification",
                    "consent",
                    "petition",
                    "request",
                    "solicitud",
                    "credit for prior learning",
                    "waiver",
                    "form",
                )
            )
            and not lowered.startswith("note:")
        ):
            return candidate
    return None


def _first_title_candidate(lines: Sequence[str]) -> str | None:
    for line in lines:
        candidate, _ = _split_title_and_remainder(line)
        candidate = _canonicalize_title_candidate(candidate)
        if any(pattern.match(candidate) for pattern in _TITLE_SKIP_PATTERNS):
            continue
        if _looks_like_non_title_line(candidate):
            continue
        if "@" in candidate:
            continue
        if _looks_like_heading(candidate) or (0 < len(candidate.split()) <= 10):
            return candidate
    return None


def _canonicalize_title_candidate(candidate: str) -> str:
    candidate = re.sub(r"\s+For Office Use Only\b.*$", "", candidate, flags=re.IGNORECASE).strip()
    candidate = re.sub(r"\s+(?:Otoño|Invierno|Primavera|Verano|Fall|Winter|Spring|Summer)\b.*$", "", candidate, flags=re.IGNORECASE).strip()
    lowered = re.sub(r"[^a-z]", "", candidate.lower())
    if "invoice" in lowered or (lowered.startswith("in") and "voic" in lowered):
        return "Invoice"
    return candidate


def _looks_like_non_title_line(line: str) -> bool:
    lowered = line.lower()
    if "www." in lowered or lowered.startswith(("http://", "https://")):
        return True
    if re.match(r"^\d", line):
        return True
    if lowered in {
        "australia",
        "customer name",
        "student information:",
        "student information",
        "parent information:",
        "parent information",
        "los angeles community college district",
        "street",
        "country",
        "postcode city",
        "city",
    }:
        return True
    if any(token in lowered for token in ("pty. ltd.", "abn ", "zip code", "telephone number")):
        return True
    if re.search(r"\bvic \d{4}\b", lowered):
        return True
    if re.fullmatch(r"[$£€]?\d[\d,\.]*", line):
        return True
    if len(line.split()) > 3 and lowered.endswith((" for", " and", " of", " to")):
        return True
    return False


def _prune_title_from_lines(lines: Sequence[str], title: str) -> list[str]:
    pruned: list[str] = []
    removed = False
    for line in lines:
        if removed:
            pruned.append(line)
            continue
        candidate, remainder = _split_title_and_remainder(line)
        if candidate == title:
            removed = True
            if remainder:
                pruned.append(remainder)
            continue
        pruned.append(line)
    return pruned


def _looks_like_heading(line: str) -> bool:
    text = line.strip()
    if not text or len(text) > 80:
        return False
    if text.endswith((".", ";", "?", "!")):
        return False
    words = text.split()
    if len(words) > 10:
        return False
    if text.lower() in _REFERENCE_HEADINGS:
        return True
    uppercase_words = sum(1 for word in words if word[:1].isupper() or word.isupper())
    return uppercase_words >= max(1, len(words) - 1)


def _looks_like_callout_label(line: str) -> bool:
    return bool(_CALL_OUT_RE.match(line.strip()))


def _looks_like_list_block(lines: Sequence[str]) -> bool:
    return len(lines) >= 2 and all(_LIST_ITEM_RE.match(line) for line in lines)


def _strip_list_marker(line: str) -> str:
    match = _LIST_ITEM_RE.match(line)
    return match.group(1).strip() if match else line.strip()


def _looks_like_table_block(lines: Sequence[str]) -> bool:
    if len(lines) < 2:
        return False
    delimiter_rows = [line for line in lines if "|" in line or "\t" in line]
    if len(delimiter_rows) >= 2:
        return True
    spaced_rows = [_split_spaced_table_row(line) for line in lines]
    multi_cell_rows = [row for row in spaced_rows if len(row) >= 2]
    if len(multi_cell_rows) < 2:
        return False
    if len(multi_cell_rows) / len(lines) < 0.6:
        return False
    long_sentence_rows = [
        row
        for row in multi_cell_rows
        if sum(len(cell.split()) for cell in row) >= 12 and any(cell.endswith((".", ";")) for cell in row)
    ]
    if len(long_sentence_rows) >= len(multi_cell_rows) / 2:
        return False
    widths = {len(row) for row in multi_cell_rows}
    return len(widths) <= 3


def _split_table_row(line: str) -> list[str]:
    if "|" in line:
        return [cell.strip() for cell in line.split("|") if cell.strip()]
    if "\t" in line:
        return [cell.strip() for cell in line.split("\t") if cell.strip()]
    spaced = _split_spaced_table_row(line)
    if len(spaced) >= 2:
        return spaced
    return [line.strip()]


def _split_spaced_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in re.split(r"\s{2,}", line.strip()) if cell.strip()]


def _normalize_table_rows(rows: Sequence[Sequence[str]]) -> list[list[str]]:
    normalized = [list(row) for row in rows if any(str(cell).strip() for cell in row)]
    if not normalized:
        return [[""]]
    width = max(len(row) for row in normalized)
    return [row + [""] * (width - len(row)) for row in normalized]


def _render_input_from_parsed(
    parsed: ParsedDocument,
    *,
    image_assets: Sequence[ImageAssetLike] | None,
    sha256: str | None = None,
) -> dict[str, Any]:
    lead: list[str] = []
    sections: list[dict[str, Any]] = []
    references: list[str] = []
    footnotes: list[dict[str, Any]] = []
    pending_blocks_before_first_section: list[dict[str, Any]] = []
    seen_section = False

    for node in parsed.body:
        block_type = getattr(node, "type", "")
        if block_type == "section":
            section_payload, section_footnotes = _render_section_from_node(node)
            if pending_blocks_before_first_section:
                section_payload["blocks"] = [
                    *pending_blocks_before_first_section,
                    *section_payload["blocks"],
                ]
                pending_blocks_before_first_section = []
            sections.append(section_payload)
            footnotes.extend(section_footnotes)
            seen_section = True
            continue

        if block_type == "paragraph" and not seen_section:
            text = getattr(node, "text", "").strip()
            if text:
                lead.append(text)
            continue
        if block_type == "reference_list":
            references.extend(_entry_texts(getattr(node, "entries", [])))
            continue
        if block_type == "footnotes":
            footnotes.extend(
                _footnote_payloads(getattr(node, "entries", []), start_index=len(footnotes) + 1)
            )
            continue

        block_payload = _render_block_from_node(node)
        if block_payload is None:
            continue
        if sections:
            sections[-1]["blocks"].append(block_payload)
        else:
            pending_blocks_before_first_section.append(block_payload)

    if pending_blocks_before_first_section and sections:
        sections[0]["blocks"] = [
            *pending_blocks_before_first_section,
            *sections[0]["blocks"],
        ]
    elif pending_blocks_before_first_section:
        sections.append(
            {
                "id": "document-body",
                "heading": "Document body",
                "level": 2,
                "caption": None,
                "blocks": pending_blocks_before_first_section,
            }
        )

    heading_tree = _build_heading_tree(sections)
    for section in sections:
        section.pop("level", None)

    return {
        "title": parsed.title,
        "language": parsed.language,
        "page_count": parsed.page_count,
        "doc_kind": parsed.doc_kind,
        "subtitle": parsed.subtitle,
        "summary": parsed.summary,
        "lead": lead,
        "assets": [
            {
                **asset.model_dump(mode="json"),
                "url": _asset_url_for_render(asset, sha256=sha256),
            }
            for asset in normalize_document_assets(image_assets)
        ],
        "heading_tree": heading_tree,
        "sections": sections,
        "references": references,
        "footnotes": footnotes,
    }


def _asset_url_for_render(asset: Any, *, sha256: str | None) -> str:
    filename = getattr(asset, "filename", "")
    if sha256 and filename:
        return f"/images/{sha256}/{filename}"
    return getattr(asset, "id", "") or filename


def _render_section_from_node(node: Any) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    blocks: list[dict[str, Any]] = []
    footnotes: list[dict[str, Any]] = []

    for block in getattr(node, "blocks", []):
        block_type = getattr(block, "type", "")
        if block_type == "footnotes":
            footnotes.extend(
                _footnote_payloads(getattr(block, "entries", []), start_index=len(footnotes) + 1)
            )
            continue
        payload = _render_block_from_node(block)
        if payload is not None:
            blocks.append(payload)

    return {
        "id": node.id,
        "heading": node.heading,
        "level": getattr(node, "level", 2),
        "caption": getattr(node, "caption", None),
        "blocks": blocks,
    }, footnotes


def _render_block_from_node(node: Any) -> dict[str, Any] | None:
    block_type = getattr(node, "type", "")
    if block_type == "paragraph":
        return {"kind": "paragraph", "text": getattr(node, "text", "")}
    if block_type == "list":
        return {
            "kind": "list",
            "items": _list_item_texts(getattr(node, "items", [])),
            "ordered": getattr(node, "list_style", "unordered") == "ordered",
        }
    if block_type == "table":
        head = [
            [
                {
                    "text": cell.text,
                    "header": cell.kind == "header",
                    "scope": cell.scope,
                    "colspan": cell.colspan if cell.colspan > 1 else None,
                    "rowspan": cell.rowspan if cell.rowspan > 1 else None,
                }
                for cell in row.cells
            ]
            for row in getattr(node, "header_rows", [])
        ]
        body = [
            [
                {
                    "text": cell.text,
                    "header": cell.kind == "header",
                    "scope": cell.scope,
                    "colspan": cell.colspan if cell.colspan > 1 else None,
                    "rowspan": cell.rowspan if cell.rowspan > 1 else None,
                }
                for cell in row.cells
            ]
            for row in getattr(node, "body_rows", [])
        ]
        payload: dict[str, Any] = {
            "kind": "table",
            "caption": getattr(node, "caption", None),
        }
        if head:
            payload["head"] = head
        payload["rows"] = body
        return payload
    if block_type == "figure":
        return {
            "kind": "figure",
            "asset_id": getattr(node, "asset_id", ""),
            "caption": getattr(node, "caption", None),
            "alt_text": getattr(node, "alt_text_hint", None),
        }
    if block_type == "callout":
        return {
            "kind": "callout",
            "title": getattr(node, "title", None),
            "text": list(getattr(node, "body", [])),
        }
    if block_type == "quote":
        return {
            "kind": "quote",
            "text": getattr(node, "text", ""),
            "citation": getattr(node, "citation", None),
            "attribution": getattr(node, "attribution", None),
        }
    if block_type == "form_group":
        return {
            "kind": "form_group",
            "title": getattr(node, "legend", None),
            "description": getattr(node, "caption", None),
            "fields": [_render_field_payload(field) for field in getattr(node, "fields", [])],
        }
    if block_type == "reference_list":
        return {"kind": "references", "items": _entry_texts(getattr(node, "entries", []))}
    return None


def _render_field_payload(field: Any) -> dict[str, Any]:
    field_type = getattr(field, "control", "text")
    options = getattr(field, "options", []) or []
    payload: dict[str, Any] = {
        "id": field.id,
        "label": field.label,
        "type": field_type,
        "required": bool(getattr(field, "required", False)),
        "help_text": getattr(field, "help_text", None),
        "placeholder": getattr(field, "placeholder", None),
    }
    if options:
        payload["options"] = [
            {"value": option.value, "label": option.label, "selected": option.selected}
            for option in options
        ]
    return payload


def _build_heading_tree(sections: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    roots: list[dict[str, Any]] = []
    stack: list[dict[str, Any]] = []
    for section in sections:
        node = {
            "id": section["id"],
            "text": section["heading"],
            "children": [],
        }
        level = max(2, int(section.get("level", 2)))
        while stack and stack[-1]["level"] >= level:
            stack.pop()
        if stack:
            stack[-1]["node"]["children"].append(node)
        else:
            roots.append(node)
        stack.append({"level": level, "node": node})
    return roots


def _render_layout_from_v2_plan(layout_plan: LayoutPlan, parsed_document: ParsedDocument) -> dict[str, Any]:
    return {
        "template": _render_template_for_doc_kind(parsed_document.doc_kind),
        "hero_style": {
            "title-only": "compact",
            "summary": "article",
            "feature": "immersive",
            "compact": "compact",
        }.get(layout_plan.hero_treatment, "compact"),
        "toc_enabled": layout_plan.toc_policy != "hidden",
        "visual_density": {
            "comfortable": "comfortable",
            "balanced": "comfortable",
            "dense": "dense",
        }.get(layout_plan.density, "comfortable"),
        "section_layouts": [
            {
                "section_id": section.section_id,
                "layout": _render_section_layout(section),
            }
            for section in layout_plan.sections
        ],
    }


def _render_template_for_doc_kind(doc_kind: str) -> str:
    return {
        "scientific_article": "scientific",
        "memo": "memo",
        "brochure": "brochure",
        "reference_doc": "reference",
        "form_doc": "form",
        "report": "reference",
    }.get(doc_kind, "reference")


def _render_section_layout(section: Any) -> str:
    rail_blocks = [block for block in getattr(section, "blocks", []) if getattr(block, "placement", "") == "rail"]
    if getattr(section, "container", "") == "feature-stack":
        return "feature-figure"
    if getattr(section, "container", "") in {"table-stack", "appendix"}:
        return "dense-reference"
    if rail_blocks:
        rail_has_figure = any(getattr(block, "block_type", "") == "figure" for block in rail_blocks)
        rail_has_callout = any(getattr(block, "block_type", "") == "callout" for block in rail_blocks)
        if rail_has_callout and not rail_has_figure:
            return "sidebar-callout"
        return "split-figure-left" if getattr(section, "rail_side", "right") == "left" else "split-figure-right"
    return "single-column"


def _entry_texts(items: Sequence[Any]) -> list[str]:
    out: list[str] = []
    for item in items:
        text = getattr(item, "text", None)
        if isinstance(text, str) and text.strip():
            out.append(text.strip())
            continue
        if isinstance(item, str) and item.strip():
            out.append(item.strip())
    return out


def _footnote_payloads(items: Sequence[Any], *, start_index: int) -> list[dict[str, Any]]:
    payloads: list[dict[str, Any]] = []
    for offset, item in enumerate(items, start=start_index):
        label = getattr(item, "label", None) or str(offset)
        text = getattr(item, "text", None) or ""
        payloads.append(
            {
                "id": _slugify(f"footnote-{label}"),
                "label": label,
                "text": text,
            }
        )
    return payloads


def _list_item_texts(items: Sequence[Any]) -> list[str]:
    values: list[str] = []
    for item in items:
        text = getattr(item, "text", None)
        if isinstance(text, str) and text.strip():
            values.append(text.strip())
            continue
        term = getattr(item, "term", None)
        description = getattr(item, "description", None)
        if isinstance(term, str) and term.strip() and isinstance(description, str) and description.strip():
            values.append(f"{term.strip()}: {description.strip()}")
    return values


def _fallback_title(source_hint: str) -> str:
    stem = Path(source_hint.rstrip("/")).name
    if "." in stem:
        stem = stem.rsplit(".", 1)[0]
    stem = stem.replace("_", " ").replace("-", " ").strip()
    return stem.title() or "Document"


def _dedupe(values: Sequence[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for value in values:
        if not value or value in seen:
            continue
        seen.add(value)
        out.append(value)
    return out


def _text_chunks(text: str, *, max_length: int = _MAX_PARAGRAPH_SIGNAL_TEXT) -> list[str]:
    normalized = re.sub(r"\s+", " ", text).strip()
    if not normalized:
        return []
    chunks: list[str] = []
    remaining = normalized
    while len(remaining) > max_length:
        cut = max(
            remaining.rfind(". ", 0, max_length),
            remaining.rfind("; ", 0, max_length),
            remaining.rfind(" ", 0, max_length),
        )
        if cut < max_length // 2:
            cut = max_length
        chunk = remaining[:cut].strip()
        if chunk:
            chunks.append(chunk)
        remaining = remaining[cut:].strip()
    if remaining:
        chunks.append(remaining)
    return chunks


def _paragraph_signals(
    *,
    text: str,
    order: int,
    source: SignalSource,
    source_key: str,
    page_start: int,
    page_end: int,
) -> list[RawBlockSignal]:
    signals: list[RawBlockSignal] = []
    for offset, chunk in enumerate(_text_chunks(text)):
        suffix = f"-{offset + 1}" if offset else ""
        signals.append(
            RawBlockSignal(
                id=f"paragraph-{order + offset}",
                kind="paragraph",
                order=order + offset,
                source=source,
                source_key=f"{source_key}{suffix}",
                page_start=page_start,
                page_end=page_end,
                text=chunk,
            )
        )
    return signals


def _slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "item"


def _unique_slug(value: str, seen: set[str]) -> str:
    base = _slugify(value)
    candidate = base
    suffix = 2
    while candidate in seen:
        candidate = f"{base}-{suffix}"
        suffix += 1
    seen.add(candidate)
    return candidate


__all__ = ["run_structured_render_v2", "RenderedDocument"]
