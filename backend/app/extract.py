"""Server-side text extraction for DOCX and XLSX.

Gemini's File API only accepts PDF, plain text, and a few web formats —
DOCX and XLSX get "Unsupported MIME type" on upload. So for those
formats we extract the content server-side (python-docx / openpyxl) into
a structured text representation and feed *that* to Gemini as a text
prompt. The remediation prompts stay the same — the input just arrives
as a string instead of a file handle.

For XLSX we deliberately preserve formulas (not evaluated values) so the
calculator-path prompt can see the formula graph.
"""
from __future__ import annotations

import io
import json
from typing import Any


def extract_docx_as_markdown(doc_bytes: bytes) -> str:
    """Render a DOCX as Markdown-ish text preserving headings, paragraphs,
    lists, and simple tables. Not a faithful round-trip — good enough for
    Gemini to reason about structure."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(doc_bytes))
    lines: list[str] = []
    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                lines.append("")
                continue
            style = (block.style.name if block.style is not None else "").lower()
            if style.startswith("heading"):
                tail = style.replace("heading", "").strip()
                try:
                    level = max(1, min(6, int(tail)))
                except ValueError:
                    level = 2
                lines.append(f"{'#' * level} {text}")
            elif "list" in style:
                lines.append(f"- {text}")
            else:
                lines.append(text)
        elif isinstance(block, Table):
            lines.append("")
            lines.append(_render_docx_table(block))
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def _iter_block_items(doc: Any) -> Any:
    """Yield paragraphs and tables in document order (python-docx doesn't
    expose this natively at the top level)."""
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent_elm = doc.element.body
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _render_docx_table(table: Any) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
    if not rows:
        return ""
    widths = [max(len(row[i]) for row in rows) for i in range(len(rows[0]))]
    def fmt(row: list[str]) -> str:
        return "| " + " | ".join(c.ljust(widths[i]) for i, c in enumerate(row)) + " |"
    sep = "| " + " | ".join("-" * w for w in widths) + " |"
    out = [fmt(rows[0]), sep]
    out.extend(fmt(r) for r in rows[1:])
    return "\n".join(out)


def extract_xlsx_as_json(doc_bytes: bytes) -> str:
    """Emit a JSON dump of the workbook: one object per sheet with cells +
    formulas preserved. The calculator-path prompt reads formulas out of
    this JSON to decide whether the sheet is interactive."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(io.BytesIO(doc_bytes), data_only=False)
    try:
        sheets: list[dict[str, Any]] = []
        for name in wb.sheetnames:
            ws = wb[name]
            cells: list[dict[str, Any]] = []
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0
            for r in range(1, max_row + 1):
                for c in range(1, max_col + 1):
                    cell = ws.cell(row=r, column=c)
                    if cell.value is None:
                        continue
                    ref = f"{get_column_letter(c)}{r}"
                    entry: dict[str, Any] = {"ref": ref}
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        entry["formula"] = value
                    else:
                        entry["value"] = value
                    cells.append(entry)
            sheets.append({"name": name, "cells": cells})
        return json.dumps({"sheets": sheets}, default=str, indent=2)
    finally:
        wb.close()
