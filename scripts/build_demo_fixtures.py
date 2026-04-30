"""Generate the demo DOCX + XLSX fixtures used by demo/index.html.

Idempotent: overwrites existing files so repeated runs always reflect the
latest source. Not checked into demo/pdfs/ — re-run any time you want to
refresh them.

    python scripts/build_demo_fixtures.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

OUT = Path(__file__).resolve().parent.parent / "demo" / "pdfs"
OUT.mkdir(parents=True, exist_ok=True)


def build_memo() -> Path:
    path = OUT / "Sample_Memo.docx"
    doc = Document()
    doc.add_heading("Quarterly Accessibility Memo", level=1)
    doc.add_paragraph(
        "This short memo exercises the DOCX ingest path: a heading, a "
        "paragraph or two, a list, and a small data table."
    )
    doc.add_heading("Highlights", level=2)
    doc.add_paragraph("Three things shipped this quarter:", style="Normal")
    for item in (
        "Screen-reader testing for the order flow",
        "WCAG AA colour-contrast audit across marketing",
        "Keyboard-only regression tests in CI",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Open bugs by severity", level=2)
    table = doc.add_table(rows=4, cols=2)
    table.style = "Light List Accent 1"
    table.rows[0].cells[0].text = "Severity"
    table.rows[0].cells[1].text = "Count"
    for i, (sev, n) in enumerate(
        [("Blocker", "0"), ("Major", "3"), ("Minor", "12")], start=1
    ):
        table.rows[i].cells[0].text = sev
        table.rows[i].cells[1].text = n
    doc.save(path)
    return path


def build_gpa_calculator() -> Path:
    path = OUT / "GPA_Calculator.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "GPA"
    ws["A1"] = "Course"
    ws["B1"] = "Grade (0.0–4.0)"
    ws["C1"] = "Credits"
    courses = [
        ("English Composition", 3.7, 3),
        ("Calculus I", 4.0, 4),
        ("World History", 3.3, 3),
        ("Intro to Biology", 3.5, 4),
    ]
    for i, (course, grade, credits) in enumerate(courses, start=2):
        ws[f"A{i}"] = course
        ws[f"B{i}"] = grade
        ws[f"C{i}"] = credits
    ws["A7"] = "Total credits"
    ws["B7"] = "=SUM(C2:C5)"
    ws["A8"] = "Weighted GPA"
    ws["B8"] = "=SUMPRODUCT(B2:B5,C2:C5)/SUM(C2:C5)"
    # Widen the label column for readability in the original XlsxPane render.
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 10
    wb.save(path)
    return path


def build_static_table() -> Path:
    """A workbook that is NOT calculator-shaped — exercises the static-table
    fallback where Gemini should emit a plain `<table>` with no `<script>`."""
    path = OUT / "Rates_Reference.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = "Year"
    ws["B1"] = "Federal Funds Rate (%)"
    ws["C1"] = "Prime Rate (%)"
    rows = [
        (2019, 1.55, 4.75),
        (2020, 0.09, 3.25),
        (2021, 0.08, 3.25),
        (2022, 4.33, 7.50),
        (2023, 5.33, 8.50),
        (2024, 4.58, 7.75),
        (2025, 4.25, 7.50),
    ]
    for i, row in enumerate(rows, start=2):
        for j, value in enumerate(row):
            ws.cell(row=i, column=j + 1, value=value)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 20
    wb.save(path)
    return path


def main() -> None:
    for build in (build_memo, build_gpa_calculator, build_static_table):
        path = build()
        print(f"wrote {path.relative_to(OUT.parent.parent)}  ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
